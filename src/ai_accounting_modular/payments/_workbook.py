"""Resolve a bank's column layout and preserve its original workbook around payment cells."""

import re
from copy import copy
from io import BytesIO

from openpyxl import load_workbook

from ai_accounting_modular.contracts.public import ContractError, parse_fen

FIELDS = ("payee_name", "payee_account", "amount", "purpose")
HEADERS = {
    "payee_name": ("收款方名称", "收款人名称", "收款人姓名", "账户名称", "姓名"),
    "payee_account": ("收款方账号", "收款账号", "银行账号", "账号"),
    "amount": ("金额", "付款金额", "代发金额", "实发金额"),
    "purpose": ("附言/用途", "附言", "用途", "摘要"),
}


def _book(raw):
    try:
        return load_workbook(BytesIO(raw), data_only=False)
    except Exception as exc:
        raise ContractError(
            "payment_template_invalid", "银行模板须为可读取、未加密的xlsx文件。"
        ) from exc


def _check(book, layout):
    if layout["sheet_name"] not in book.sheetnames:
        raise ContractError("payment_template_invalid", "指定的代发表工作表不存在。")
    sheet = book[layout["sheet_name"]]
    columns, first = set(layout["columns"].values()), layout["first_data_row"]
    last = layout.get("last_data_row", 1048576)
    if last < first:
        raise ContractError("payment_template_invalid", "付款数据区域的末行不能早于首行。")
    if len(columns) != 4:
        raise ContractError(
            "payment_template_invalid", "收款名称、账号、金额和用途须使用四个不同列。"
        )
    if any(
        merged.max_row >= first
        and merged.min_row <= last
        and any(merged.min_col <= column <= merged.max_col for column in columns)
        for merged in sheet.merged_cells.ranges
    ):
        raise ContractError(
            "payment_template_unsupported", "付款数据列存在合并单元格，须先明确可写的数据区域。"
        )
    if "last_data_row" not in layout:
        occupied = {}
        for cell in sheet._cells.values():
            if cell.row >= first and cell.column in columns and cell.value not in (None, ""):
                occupied.setdefault(cell.row, {})[cell.column] = cell
        for row in occupied.values():
            amount = row.get(layout["columns"]["amount"])
            if (
                len(row) != 4
                or amount is None
                or (
                    amount.data_type != "f"
                    and re.fullmatch(r"[+-]?\d+(?:\.\d+)?", str(amount.value)) is None
                )
            ):
                raise ContractError(
                    "payment_template_layout_required",
                    "付款列下方含不能确定为付款示例的文字或公式。AI须查阅模板并明确last_data_row，保留同列尾注或合计。",
                )


def resolve_layout(raw, supplied=None):
    book = _book(raw)
    try:
        layout = supplied
        if layout is None:
            matches = []
            for sheet in book:
                for row in sheet.iter_rows(max_row=min(sheet.max_row, 20)):
                    found = {}
                    for cell in row:
                        text = "".join(str(cell.value or "").split()).replace("（", "(")
                        for field, aliases in HEADERS.items():
                            if any(text == name or text.startswith(name + "(") for name in aliases):
                                if field in found:
                                    found[field] = None
                                else:
                                    found[field] = cell.column
                    if set(found) == set(FIELDS) and all(found.values()):
                        matches.append(
                            {
                                "sheet_name": sheet.title,
                                "first_data_row": row[0].row + 1,
                                "columns": found,
                            }
                        )
            if len(matches) != 1:
                raise ContractError(
                    "payment_template_layout_required",
                    "未能唯一识别四类付款列，AI须先查阅模板再指定布局。",
                )
            layout = matches[0]
        _check(book, layout)
        return layout
    finally:
        book.close()


def text(value):
    if not isinstance(value, str) or not value.strip():
        raise ContractError("payment_recipient_invalid", "收款名称、账号和用途须为非空文本。")
    return value


def yuan(fen):
    whole, cents = divmod(parse_fen(fen), 100)
    return f"{whole}.{cents:02d}"


def _filled_book(raw, layout, rows):
    """Parse a template once, validate its region, and replace only payment cells."""
    book = _book(raw)
    try:
        _check(book, layout)
        book.template = False
        sheet = book[layout["sheet_name"]]
        first, columns = layout["first_data_row"], layout["columns"]
        last = layout.get("last_data_row", 1048576)
        if first + len(rows) - 1 > last:
            raise ContractError(
                "payment_template_capacity",
                "代发人数超过模板的可写区域，请先查阅银行模板的扩展规则。",
            )
        styles = {key: copy(sheet.cell(first, column)._style) for key, column in columns.items()}
        # A far-right instruction or preformatted row must not materialize a blank rectangle.
        for cell in tuple(sheet._cells.values()):
            if first <= cell.row <= last and cell.column in columns.values():
                cell.value = None
        for index, row in enumerate(rows, first):
            values = {
                "payee_name": text(row["name"]),
                "payee_account": text(row["account"]),
                "amount": yuan(row["amount_fen"]),
                "purpose": text(row["purpose"]),
            }
            for field, value in values.items():
                cell = sheet.cell(index, columns[field])
                cell._style = copy(styles[field])
                cell.value, cell.data_type, cell.number_format = value, "s", "@"
        return book
    except Exception:
        book.close()
        raise


def render(raw, layout, rows):
    book = _filled_book(raw, layout, rows)
    try:
        output = BytesIO()
        book.save(output)
        return output.getvalue()
    finally:
        book.close()


def verify(raw, template, layout, rows):
    """Check every payment cell and every retained workbook cell against the intended output."""
    actual = _book(raw)
    reference = None
    try:
        reference = _filled_book(template, layout, rows)
        if actual.template or actual.sheetnames != reference.sheetnames:
            raise ContractError("payment_export_invalid", "代发表的工作表与模板不一致。")
        for left, right in zip(actual, reference, strict=True):
            if (
                sorted(str(value) for value in left.merged_cells.ranges),
                {key: dict(value) for key, value in left.column_dimensions.items()},
            ) != (
                sorted(str(value) for value in right.merged_cells.ranges),
                {key: dict(value) for key, value in right.column_dimensions.items()},
            ):
                raise ContractError("payment_export_invalid", "代发表改变了模板布局。")
            for coordinate in left._cells.keys() | right._cells.keys():
                cell, required = left.cell(*coordinate), right.cell(*coordinate)
                if (cell.value, cell.data_type, tuple(cell._style) if cell.has_style else ()) != (
                    required.value,
                    required.data_type,
                    tuple(required._style) if required.has_style else (),
                ):
                    raise ContractError("payment_export_invalid", "代发表逐格回读核验失败。")
    finally:
        actual.close()
        if reference is not None:
            reference.close()
