"""Original spreadsheet cells, formulas, caches and hidden positions; no formula execution."""

import io
import mmap
import posixpath
from xml.etree import ElementTree
from zipfile import BadZipFile, ZipFile

import xlrd
from openpyxl.styles.numbers import BUILTIN_FORMATS, is_date_format, is_timedelta_format
from openpyxl.utils import column_index_from_string, get_column_letter
from openpyxl.utils.cell import coordinate_from_string
from openpyxl.utils.datetime import CALENDAR_MAC_1904, CALENDAR_WINDOWS_1900, from_excel

from ai_accounting_modular.contracts.public import ContractError, canonical_json


def _text(node):
    if node is None:
        return ""
    parts = []
    for child in node:
        if child.tag.endswith("}t"):
            parts.append(child.text or "")
        elif child.tag.endswith("}r"):
            parts.append(child.findtext("{*}t") or "")
    return "".join(parts)


def _strings(archive):
    if "xl/sharedStrings.xml" not in archive.namelist():
        return []
    values = []
    with archive.open("xl/sharedStrings.xml") as stream:
        events = ElementTree.iterparse(stream, events=("start", "end"))
        _, root = next(events)
        for event, node in events:
            if event == "end" and node.tag.endswith("}si"):
                values.append(_text(node))
                node.clear()
                root.remove(node)
    return values


def _formats(archive):
    if "xl/styles.xml" not in archive.namelist():
        return [(0, "General")]
    root = ElementTree.fromstring(archive.read("xl/styles.xml"))
    values = dict(BUILTIN_FORMATS)
    values.update(
        {
            int(node.attrib["numFmtId"]): node.attrib["formatCode"]
            for node in root.findall("{*}numFmts/{*}numFmt")
        }
    )
    return [
        (int(node.get("numFmtId", "0")), values.get(int(node.get("numFmtId", "0"))))
        for node in root.findall("{*}cellXfs/{*}xf")
    ]


def _indexed(values, index):
    number = int(index)
    if not 0 <= number < len(values):
        raise ValueError("spreadsheet index")
    return values[number]


def _cell(cell, namespace, strings, formats, epoch):
    value, kind = cell.findtext(namespace + "v"), cell.get("t", "n")
    if kind == "inlineStr":
        value = _text(cell.find(namespace + "is"))
    elif kind == "s" and value is not None:
        value = _indexed(strings, value)
    formula = cell.find(namespace + "f")
    if value == "" and kind == "n":
        value = None
    if value in (None, "") and formula is None:
        return None
    format_id, number_format = _indexed(formats, cell.get("s", "0"))
    payload = {"value": value, "type": kind, "number_format": number_format}
    if number_format is None:
        payload["number_format_id"] = format_id
    elif kind == "n" and value is not None and is_date_format(number_format):
        # The stored amount lexeme stays unchanged; only formatted dates gain a display value.
        duration = is_timedelta_format(number_format)
        converted = from_excel(float(value), epoch=epoch, timedelta=duration)
        payload["display_value"] = str(converted) if duration else converted.isoformat()
    if formula is not None:
        payload["formula"] = {"text": formula.text or "", **formula.attrib}
    return canonical_json(payload)


def _worksheet(xml, section, hidden, strings, formats, epoch):
    hidden_columns, stack = set(), []
    namespace, row_number = None, 0
    for event, node in ElementTree.iterparse(xml, events=("start", "end")):
        if event == "start":
            if not stack:
                if not node.tag.endswith("}worksheet"):
                    raise ValueError("worksheet type")
                namespace = node.tag.removesuffix("worksheet")
            stack.append(node)
            continue
        parent = stack[-2].tag if len(stack) > 1 else None
        if node.tag == namespace + "col" and parent == namespace + "cols":
            first, last = int(node.attrib["min"]), int(node.attrib["max"])
            if not 1 <= first <= last <= 16384:
                raise ValueError("worksheet columns")
            if node.get("hidden") in {"1", "true"}:
                hidden_columns.update(range(first, last + 1))
        elif node.tag == namespace + "row" and parent == namespace + "sheetData":
            row_number = int(node.get("r", row_number + 1))
            if not 1 <= row_number <= 1048576:
                raise ValueError("worksheet row")
            row_hidden = hidden or node.get("hidden") in {"1", "true"}
            column_number = 0
            for cell in node.findall(namespace + "c"):
                coordinate = cell.get("r")
                if coordinate is None:
                    column_number += 1
                    coordinate = f"{get_column_letter(column_number)}{row_number}"
                else:
                    column, actual_row = coordinate_from_string(coordinate)
                    column_number = column_index_from_string(column)
                    if actual_row != row_number:
                        raise ValueError("worksheet cell row")
                if not 1 <= column_number <= 16384:
                    raise ValueError("worksheet cell column")
                payload = _cell(cell, namespace, strings, formats, epoch)
                if payload is not None:
                    location = f"{section}!{coordinate}" + (
                        " [hidden]" if row_hidden or column_number in hidden_columns else ""
                    )
                    yield section, location, payload
            node.clear()
            stack[-2].remove(node)
        stack.pop()


def xlsx(stream, sections):
    try:
        with ZipFile(stream) as archive:
            workbook = ElementTree.fromstring(archive.read("xl/workbook.xml"))
            rels = ElementTree.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
            targets = {node.attrib["Id"]: node for node in rels}
            strings, formats = _strings(archive), _formats(archive)
            properties = workbook.find("{*}workbookPr")
            epoch = (
                CALENDAR_MAC_1904
                if properties is not None and properties.get("date1904") in {"1", "true"}
                else CALENDAR_WINDOWS_1900
            )
            for sheet in workbook.findall("{*}sheets/{*}sheet"):
                name = sheet.attrib["name"]
                section = "sheet:" + name
                sections.add(section)
                relationship = next(
                    value for key, value in sheet.attrib.items() if key.endswith("}id")
                )
                target = targets[relationship]
                if target.get("TargetMode") == "External":
                    raise ValueError("external worksheet")
                path = target.attrib["Target"]
                path = (
                    path.lstrip("/")
                    if path.startswith("/")
                    else posixpath.normpath(posixpath.join("xl", path))
                )
                if not path.startswith("xl/") or ".." in path.split("/"):
                    raise ValueError("worksheet path")
                with archive.open(path) as xml:
                    yield from _worksheet(
                        xml,
                        section,
                        sheet.get("state", "visible") != "visible",
                        strings,
                        formats,
                        epoch,
                    )
    except (
        BadZipFile,
        ElementTree.ParseError,
        ValueError,
        KeyError,
        IndexError,
        OverflowError,
        StopIteration,
    ) as error:
        raise ContractError(
            "source_unreadable", "Excel原件结构或单元格内容无法完整读取。"
        ) from error


def xls(stream, sections, unsupported):
    # xlrd accepts a read-only mmap, avoiding another full in-memory copy of the verified file.
    with (
        mmap.mmap(stream.fileno(), 0, access=mmap.ACCESS_READ) as data,
        io.StringIO() as diagnostic,
    ):
        book = xlrd.open_workbook(
            file_contents=data,
            formatting_info=True,
            on_demand=True,
            ragged_rows=True,
            logfile=diagnostic,
        )
        try:
            for index in range(book.nsheets):
                sheet = book.sheet_by_index(index)
                section = "sheet:" + sheet.name
                sections.add(section)
                unsupported.append(
                    section + "：XLS公式表达式不能提取，仅提供缓存值；空缓存不代表业务事实为空。"
                )
                for row in range(sheet.nrows):
                    for column in range(sheet.row_len(row)):
                        cell = sheet.cell(row, column)
                        if cell.ctype in {xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK}:
                            continue
                        value = cell.value
                        if cell.ctype == xlrd.XL_CELL_DATE:
                            value = xlrd.xldate_as_datetime(value, book.datemode).isoformat()
                        elif isinstance(value, float):
                            # Preserve the decoded Excel number; no accounting calculation occurs.
                            value = repr(value)
                        hidden = (
                            sheet.visibility
                            or bool(sheet.rowinfo_map.get(row) and sheet.rowinfo_map[row].hidden)
                            or bool(
                                sheet.colinfo_map.get(column) and sheet.colinfo_map[column].hidden
                            )
                        )
                        location = f"{section}!{get_column_letter(column + 1)}{row + 1}" + (
                            " [hidden]" if hidden else ""
                        )
                        yield (
                            section,
                            location,
                            canonical_json({"value": value, "type": cell.ctype}),
                        )
                book.unload_sheet(index)
        finally:
            book.release_resources()
