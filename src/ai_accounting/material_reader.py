"""Deterministic source locations and amount coverage, without accounting inference."""

from __future__ import annotations

import csv
import hashlib
import io
import posixpath
from decimal import Decimal, InvalidOperation
from pathlib import Path
from xml.etree import ElementTree
from zipfile import ZipFile

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from .config import get_settings
from .path_security import read_regular_file_in_root


def amount_fen(value) -> int:
    raw = (
        str(value)
        .strip()
        .replace(",", "")
        .replace("¥", "")
        .replace("￥", "")
        .replace("元", "")
        .replace(" ", "")
    )
    try:
        number = Decimal(raw)
        fen = number * 100
        if not number.is_finite() or fen != fen.to_integral_value():
            raise ValueError("MATERIAL_AMOUNT_PRECISION")
        return int(fen)
    except InvalidOperation as exc:
        raise ValueError("MATERIAL_AMOUNT_UNREADABLE") from exc


def verify_material_bytes(evidence) -> bytes:
    _, raw = read_regular_file_in_root(
        Path(evidence.storage_path),
        get_settings().finance_evidence_dir,
        max_bytes=get_settings().finance_max_evidence_bytes,
    )
    if hashlib.sha256(raw).hexdigest() != evidence.sha256:
        raise ValueError("MATERIAL_SOURCE_CHANGED")
    return raw


def _xlsx_numeric_text(raw):
    # openpyxl exposes decimals as binary floats. Read the original numeric XML
    # strings before conversion to integer fen, including cached formula values.
    result = {}
    with ZipFile(io.BytesIO(raw)) as archive:
        workbook = ElementTree.fromstring(archive.read("xl/workbook.xml"))
        relationships = ElementTree.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
        targets = {rel.attrib["Id"]: rel.attrib["Target"] for rel in relationships}
        for sheet in workbook.findall(".//{*}sheet"):
            rel_id = next(
                (value for key, value in sheet.attrib.items() if key.endswith("}id")), None
            )
            if rel_id is None:
                raise ValueError("MATERIAL_SHEET_REFERENCE_MISSING")
            target = targets[rel_id]
            path = (
                target.lstrip("/")
                if target.startswith("/")
                else posixpath.normpath(posixpath.join("xl", target))
            )
            document = ElementTree.fromstring(archive.read(path))
            for cell in document.findall(".//{*}c"):
                value = cell.findtext("{*}v")
                if cell.get("t", "n") == "n" and value is not None and value.strip():
                    result[sheet.attrib["name"], cell.attrib["r"]] = value
    return result


def inspect_material(evidence, spec) -> dict:
    raw = verify_material_bytes(evidence)
    source = {
        "evidence_id": str(evidence.id),
        "sha256": evidence.sha256,
        "name": evidence.original_name,
        "spec": spec.model_dump(mode="json"),
        "items": [],
        "issues": [],
        "coverage": [],
    }
    suffix = Path(evidence.original_name).suffix.lower()
    if suffix not in {".csv", ".xlsx"}:
        passages = spec.passages or {"全文": ""}
        for location, excerpt in passages.items():
            source["coverage"].append({"location": location, "excerpt": excerpt})
            source["items"].append(
                {
                    "key": f"{evidence.id}:{location}",
                    "location": location,
                    "excerpt": excerpt,
                    "amount_fen": None,
                    "source_id": str(evidence.id),
                }
            )
        if not spec.passages or any(not text.strip() for text in spec.passages.values()):
            source["issues"].append(
                {
                    "code": "MATERIAL_PASSAGES_UNREVIEWED",
                    "location": "全文",
                    "message": "请阅读全部原资料，登记页码或原文位置及核对内容。",
                }
            )
        return source
    cells = []
    if suffix == ".csv":
        decoded = None
        for encoding in ("utf-8-sig", "gb18030"):
            try:
                decoded = raw.decode(encoding)
                break
            except UnicodeDecodeError:
                pass
        if decoded is None:
            raise ValueError("MATERIAL_CSV_ENCODING")
        for row_no, row in enumerate(csv.reader(io.StringIO(decoded)), 1):
            for col_no, value in enumerate(row, 1):
                if value.strip():
                    cells.append(("CSV", row_no, get_column_letter(col_no), value, False, False))
    else:
        numeric_text = _xlsx_numeric_text(raw)
        formulas = load_workbook(io.BytesIO(raw), data_only=False, read_only=False)
        cached = load_workbook(io.BytesIO(raw), data_only=True, read_only=False)
        try:
            if spec.sheet and spec.sheet not in formulas.sheetnames:
                raise ValueError("MATERIAL_SHEET_NOT_FOUND")
            for sheet in formulas:
                if spec.sheet and sheet.title != spec.sheet:
                    if any(c.value is not None for row in sheet for c in row):
                        source["issues"].append(
                            {
                                "code": "MATERIAL_SHEET_UNREVIEWED",
                                "location": sheet.title,
                                "message": "原文件还有未纳入核对的工作表。省略sheet可核对全部表。",
                            }
                        )
                    continue
                for row in sheet:
                    for cell in row:
                        if cell.value is None or str(cell.value).strip() == "":
                            continue
                        value = (
                            cached[sheet.title][cell.coordinate].value
                            if cell.data_type == "f"
                            else cell.value
                        )
                        if type(value) in (int, float):
                            value = numeric_text[sheet.title, cell.coordinate]
                        cells.append(
                            (
                                sheet.title,
                                cell.row,
                                get_column_letter(cell.column),
                                value,
                                cell.data_type == "f" and value is None,
                                bool(
                                    sheet.row_dimensions[cell.row].hidden
                                    or sheet.column_dimensions[
                                        get_column_letter(cell.column)
                                    ].hidden
                                    or sheet.sheet_state != "visible"
                                ),
                            )
                        )
        finally:
            formulas.close()
            cached.close()
    if len(cells) > 100_000:
        raise ValueError("MATERIAL_TOO_MANY_CELLS")

    def sheet_columns(sheet):
        rules = spec.sheet_columns.get(sheet, spec.columns)
        columns = {c.column.upper(): c for c in rules}
        if len(columns) != len(rules):
            raise ValueError("MATERIAL_DUPLICATE_COLUMN")
        return columns

    present = {(sheet, row, col) for sheet, row, col, *_ in cells}
    occupied_rows = sorted({(sheet, row) for sheet, row, *_ in cells})
    for sheet, row in occupied_rows:
        if row <= spec.sheet_header_rows.get(sheet, spec.header_row):
            continue
        for column, rule in sheet_columns(sheet).items():
            if rule.role == "amount" and (sheet, row, column) not in present:
                cells.append((sheet, row, column, None, False, False))
    totals = {}
    declared_totals = []
    context = {}
    for sheet, row, column, value, *_ in cells:
        rule = sheet_columns(sheet).get(column)
        if rule and rule.role == "context" and value is not None:
            context.setdefault((sheet, row), []).append(str(value))
    for sheet, row_no, column, value, formula_missing, hidden in cells:
        location = f"{sheet}!{column}{row_no}"
        source["coverage"].append({"location": location, "value": str(value), "hidden": hidden})
        if row_no <= spec.sheet_header_rows.get(sheet, spec.header_row):
            continue
        rule = sheet_columns(sheet).get(column)
        is_total = row_no in spec.sheet_total_rows.get(sheet, spec.total_rows)

        pending_item = {
            "key": f"{evidence.id}:{location}",
            "location": location,
            "excerpt": " · ".join(
                [*context.get((sheet, row_no), []), (rule.label if rule else "") or column]
            ),
            "amount_fen": None,
            "source_id": str(evidence.id),
        }

        if formula_missing:
            source["issues"].append(
                {
                    "code": "MATERIAL_FORMULA_RESULT_MISSING",
                    "location": location,
                    "message": "公式没有可核对的计算结果，不能按零跳过。",
                }
            )
            if rule and rule.role == "amount" and not is_total:
                source["items"].append(pending_item)
            continue
        if rule is None:
            source["issues"].append(
                {
                    "code": "MATERIAL_COLUMN_UNMAPPED",
                    "location": location,
                    "message": "此列尚未说明是业务金额还是背景资料。",
                }
            )
            continue
        if rule.role == "context":
            continue
        if value is None:
            source["issues"].append(
                {
                    "code": "MATERIAL_AMOUNT_MISSING",
                    "location": location,
                    "message": "业务金额缺失，不能当作零或漏掉这一行。",
                }
            )
            if not is_total:
                source["items"].append(pending_item)
            continue
        try:
            amount = amount_fen(value)
        except ValueError as exc:
            source["issues"].append(
                {"code": str(exc), "location": location, "message": "该金额不能读取为整数分。"}
            )
            if not is_total:
                source["items"].append(pending_item)
            continue
        if is_total:
            declared_totals.append((sheet, column, location, amount))
            continue
        totals[sheet, column] = totals.get((sheet, column), 0) + amount
        source["items"].append(
            {
                "key": f"{evidence.id}:{location}",
                "location": location,
                "excerpt": " · ".join(
                    [*context.get((sheet, row_no), []), f"{rule.label or column}: {value}"]
                ),
                "amount_fen": amount,
                "source_id": str(evidence.id),
            }
        )
    for sheet, column, location, amount in declared_totals:
        if amount != totals.get((sheet, column), 0):
            source["issues"].append(
                {
                    "code": "MATERIAL_TOTAL_MISMATCH",
                    "location": location,
                    "expected_fen": amount,
                    "actual_fen": totals.get((sheet, column), 0),
                    "message": "原资料合计与明细不一致。",
                }
            )
    if not source["items"]:
        source["issues"].append(
            {
                "code": "MATERIAL_NO_BUSINESS_COLUMNS",
                "location": "全文",
                "message": "未登记业务金额列，不能将空清单当作核对完成。",
            }
        )
    source["detail_count"] = len(source["items"])
    source["control_totals"] = [
        {"location": location, "expected_fen": amount, "actual_fen": totals.get((sheet, column), 0)}
        for sheet, column, location, amount in declared_totals
    ]
    return source
