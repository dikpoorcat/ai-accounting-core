"""Source-row coverage against current typed facts, shared by close and exports.

Parsing establishes locations and amounts, never the business interpretation.
All confirmations are new typed revisions. A citation of a file is not coverage.
"""

from __future__ import annotations

import csv
import io
import posixpath
import re
from bisect import bisect_left, bisect_right
from decimal import Decimal, InvalidOperation
from types import SimpleNamespace
from typing import Annotated, ClassVar, Literal
from xml.etree import ElementTree
from zipfile import ZipFile

from openpyxl.utils import (
    column_index_from_string,
    get_column_letter,
    range_boundaries,
)
from openpyxl.utils.cell import coordinate_from_string
from pydantic import (
    BaseModel,
    ConfigDict,
    Field,
    StrictBool,
    TypeAdapter,
    model_serializer,
    model_validator,
)

from .build import calculator_build_id
from .contracts import Fact, FactVersion, KernelError, NeedsInformation
from .schema import table_name
from .storage import Store
from .types import ActualDate, Fen, YearMonth, canonical, checked, digest, sum_fen

Identifier = Annotated[str, Field(min_length=1, max_length=200)]
EvidenceDigest = Annotated[str, Field(pattern=r"^[0-9a-f]{64}$")]
Category = Literal["transactions", "payroll", "bank", "tax", "assets", "financing"]
PageLimit = Annotated[int, Field(strict=True, ge=1, le=500)]


class Detail(BaseModel):
    model_config = ConfigDict(extra="forbid", frozen=True, strict=True)


class Column(Detail):
    column: Annotated[str, Field(pattern=r"^[A-Z]+$")]
    role: Literal["amount", "context", "recognition_period"]
    label: str = ""
    funds_direction: Literal["signed_net", "inflow", "outflow"] | None = Field(
        default=None,
        description=(
            "原金额列的核算方向：signed_net为原正负净额，inflow/outflow为非负收入/支出列。"
            "仅金额列可声明，须依据原表头或明确来源，不能从label猜测。原金额及控制合计不改写；"
            "省略时兼容原有正负净额语义，旧来源序列化不新增字段。"
        ),
    )

    @model_validator(mode="after")
    def amount_direction_only(self):
        if self.funds_direction is not None and self.role != "amount":
            raise ValueError("funds direction may only describe an original amount column")
        return self

    @model_serializer(mode="wrap")
    def compatible_dump(self, handler):
        result = handler(self)
        if self.funds_direction is None:
            result.pop("funds_direction", None)
        return result


class Passage(Detail):
    location: Identifier
    page: Annotated[int, Field(ge=1)]
    excerpt: Annotated[str, Field(min_length=1)]
    amount_fen: Fen | None = None
    recognition_period: YearMonth | None = None


class MaterialRange(Detail):
    sheet: Identifier
    column: Annotated[str, Field(pattern=r"^[A-Z]+$")]
    first_row: Annotated[int, Field(ge=1)]
    last_row: Annotated[int, Field(ge=1)]

    @model_validator(mode="after")
    def ordered(self):
        if self.last_row < self.first_row:
            raise ValueError("material range must be increasing")
        return self


class ControlTotal(Detail):
    location: Identifier
    scope: Literal["specified", "all_detail_column"] = "specified"
    ranges: tuple[MaterialRange, ...] = ()
    locations: tuple[Identifier, ...] = ()

    @model_validator(mode="after")
    def explicit_members(self):
        if self.scope == "specified" and not (self.ranges or self.locations):
            raise ValueError("a grouped control needs explicit original members")
        if self.scope == "all_detail_column" and (self.ranges or self.locations):
            raise ValueError("whole-column control cannot also name partial members")
        return self


class Specification(Detail):
    format: Literal["csv", "xls", "xlsx", "pdf", "image", "text"]
    columns: tuple[Column, ...] = ()
    sheet_columns: dict[str, tuple[Column, ...]] = Field(default_factory=dict)
    header_rows: dict[str, int] = Field(default_factory=dict)
    total_rows: dict[str, tuple[int, ...]] = Field(default_factory=dict)
    passages: tuple[Passage, ...] = ()
    all_pages_reviewed: StrictBool | None = None
    controls: tuple[ControlTotal, ...] = ()

    @model_serializer(mode="wrap")
    def compatible_dump(self, handler):
        result = handler(self)
        if not self.controls:
            result.pop("controls", None)
        return result

    @model_validator(mode="after")
    def unique_columns(self):
        for columns in (self.columns, *self.sheet_columns.values()):
            if len({item.column for item in columns}) != len(columns):
                raise ValueError("duplicate material column")
            if sum(item.role == "recognition_period" for item in columns) > 1:
                raise ValueError("one declared recognition-period column per sheet")
        if any(type(row) is not int or row < 0 for row in self.header_rows.values()):
            raise ValueError("header rows must be nonnegative integers")
        if any(row <= 0 for rows in self.total_rows.values() for row in rows):
            raise ValueError("total rows must be positive")
        if len({item.location for item in self.passages}) != len(self.passages):
            raise ValueError("duplicate original passage location")
        if len({item.location for item in self.controls}) != len(self.controls):
            raise ValueError("duplicate original control location")
        if self.controls and self.format not in {"csv", "xls", "xlsx"}:
            raise ValueError("group controls require original spreadsheet locations")
        return self


class PeriodAllocationEntry(Detail):
    sheet: Identifier | None = None
    column: Annotated[str, Field(pattern=r"^[A-Z]+$")] | None = None
    first_row: Annotated[int, Field(ge=1)] | None = None
    last_row: Annotated[int, Field(ge=1)] | None = None
    location: Identifier | None = None
    recognition_period: YearMonth | None
    basis: Literal["original_period", "confirmed_period", "unknown", "document_period"]
    basis_evidence_digest: EvidenceDigest | None = None
    basis_location: Identifier | None = None
    basis_excerpt: Annotated[str, Field(min_length=1)] | None = None

    @model_validator(mode="after")
    def coherent(self):
        fields = (self.sheet, self.column, self.first_row, self.last_row)
        if self.location is not None:
            if any(value is not None for value in fields):
                raise ValueError("select a passage or a cell range, not both")
        elif any(value is None for value in fields) or self.first_row > self.last_row:
            raise ValueError("allocation requires one exact original range")
        if (self.basis == "unknown") != (self.recognition_period is None):
            raise ValueError("unknown periods must stay unknown")
        if self.basis == "confirmed_period" and not all(
            (self.basis_evidence_digest, self.basis_location, self.basis_excerpt)
        ):
            raise ValueError("confirmed periods need registered evidence and an exact basis")
        return self


class MaterialPeriodAllocation(Fact):
    kind: ClassVar[str] = "material_period_allocation"
    lane: ClassVar[str] = "material"
    registration_command: ClassVar[str] = "confirm_material_allocation"
    identity_fields: ClassVar[tuple[str, ...]] = ("source_id",)
    immutable: ClassVar[bool] = True
    immutable_fields: ClassVar[tuple[str, ...]] = ("source_id",)
    source_id: Identifier
    source_fact_id: Identifier
    entries: tuple[PeriodAllocationEntry, ...]

    def scopes(self):
        return (
            str(self.period),
            "material-source:" + self.source_id,
            *(
                "material-period:" + str(item.recognition_period)
                for item in self.entries
                if item.recognition_period is not None
            ),
            *(
                ("material-period-unknown",)
                if not self.entries or any(item.recognition_period is None for item in self.entries)
                else ()
            ),
        )


class MaterialSource(Fact):
    kind: ClassVar[str] = "material_source_v2"
    lane: ClassVar[str] = "material"
    identity_fields: ClassVar[tuple[str, ...]] = ("evidence_digest", "period")
    evidence_digest: EvidenceDigest
    category: Category
    purpose: Literal["business", "supporting"]
    supporting_purpose: str | None = None
    specification: Specification

    @model_validator(mode="after")
    def support_is_explicit(self):
        if self.purpose == "supporting" and not (self.supporting_purpose or "").strip():
            raise ValueError("supporting evidence needs an explicit purpose")
        return self

    def scopes(self):
        return (str(self.period), "material-evidence:" + self.evidence_digest)


class MaterialLink(Detail):
    subject_id: Identifier
    fact_kind: Identifier
    fact_id: Identifier
    calculation_id: Identifier
    amount_field: Annotated[str, Field(pattern=r"^(fact|result)\.[a-z][a-z0-9_]*_fen$")]
    amount_fen: Fen
    recognition_period: YearMonth


class MaterialResolution(Fact):
    kind: ClassVar[str] = "material_resolution_v2"
    lane: ClassVar[str] = "material"
    identity_fields: ClassVar[tuple[str, ...]] = ("source_id", "location")
    source_id: Identifier
    source_fact_id: Identifier
    location: Identifier
    treatment: Literal[
        "recognize", "duplicate", "supporting", "no_accounting", "other_period", "control_total"
    ]
    amount_fen: Fen | None = None
    recognition_period: YearMonth | None = None
    links: tuple[MaterialLink, ...] = ()
    duplicate_source_id: Identifier | None = None
    duplicate_location: Identifier | None = None
    reason: str | None = None
    non_accounting_reason: (
        Literal[
            "forecast",
            "balance_control",
            "not_company_business",
            "cancelled_before_recognition",
            "zero_amount",
        ]
        | None
    ) = None

    def scopes(self):
        return (
            str(self.period),
            "material-source:" + self.source_id,
            *(
                ("material-period:" + str(self.recognition_period),)
                if self.recognition_period
                else ()
            ),
            *("material-period:" + str(link.recognition_period) for link in self.links),
            *("material-business:" + link.subject_id for link in self.links),
        )


class MaterialGroupMember(Detail):
    location: Identifier
    amount_fen: Fen


class MaterialGroupResolution(Fact):
    """One approved pool, without asserting an unobserved row-to-result matrix."""

    kind: ClassVar[str] = "material_group_resolution"
    lane: ClassVar[str] = "material"
    registration_command: ClassVar[str] = "resolve_material_group"
    identity_fields: ClassVar[tuple[str, ...]] = ("source_id",)
    source_id: Identifier
    source_fact_id: Identifier
    members: tuple[MaterialGroupMember, ...] = Field(min_length=1, max_length=100_000)
    group_amount_fen: Fen
    links: tuple[MaterialLink, ...] = Field(min_length=1)
    joint_basis_confirmed: StrictBool | None = None
    basis_evidence_digest: EvidenceDigest
    basis_location: Identifier
    reason: str = Field(min_length=1, max_length=4000)

    @model_validator(mode="after")
    def conserved_pool(self):
        if len({item.location for item in self.members}) != len(self.members):
            raise ValueError("group members must be unique")
        if self.group_amount_fen == 0:
            raise ValueError("zero controls do not form an accounting allocation group")
        positive = self.group_amount_fen > 0
        if any(item.amount_fen and (item.amount_fen > 0) != positive for item in self.members):
            raise ValueError("group members must have the same direction")
        if any(not item.amount_fen or (item.amount_fen > 0) != positive for item in self.links):
            raise ValueError("group result allocations must have the same direction")
        if (
            sum_fen(item.amount_fen for item in self.members) != self.group_amount_fen
            or sum_fen(item.amount_fen for item in self.links) != self.group_amount_fen
        ):
            raise ValueError("original pool, group total and result allocations must agree")
        return self

    def scopes(self):
        return (
            "material-source:" + self.source_id,
            *("material-period:" + str(link.recognition_period) for link in self.links),
            *("material-business:" + link.subject_id for link in self.links),
        )


def register(registry):
    registry.register(MaterialSource)
    registry.register(MaterialResolution)
    registry.register(MaterialPeriodAllocation)
    registry.register(MaterialGroupResolution)


def _issue(code, message, *, location=None, **details):
    return {"field": "materials", "code": code, "message": message, "location": location, **details}


def _amount(value):
    text = str(value).strip()
    for char in (",", "¥", "￥", "元", " "):
        text = text.replace(char, "")
    try:
        number = Decimal(text)
        if not number.is_finite():
            raise ValueError("amount needs exact integer fen")
        parts = number.as_tuple()
        digits, scale = parts.digits, parts.exponent + 2
        if not any(digits):
            return 0
        if len(digits) + scale > 19:
            raise ValueError("amount exceeds signed 64-bit fen")
        if scale < 0:
            removed = -scale
            if removed >= len(digits) or any(digits[-removed:]):
                raise ValueError("amount needs exact integer fen")
            digits, scale = digits[:-removed], 0
        # Decimal arithmetic uses the ambient context. Exact digit shifting
        # avoids rounding and rejects enormous exponents before integer allocation.
        amount = int("".join(str(digit) for digit in digits)) * 10**scale
        return checked(-amount if parts.sign else amount)
    except (InvalidOperation, ValueError) as exc:
        raise ValueError("amount is missing, unreadable or not integer fen") from exc


def inspect_bytes(raw: bytes, specification: Specification) -> dict:
    """Normalize corrupt/unreadable original formats into a structured issue."""
    from zipfile import BadZipFile

    from openpyxl.utils.exceptions import InvalidFileException
    from pypdf.errors import PyPdfError
    from xlrd.biffh import XLRDError

    try:
        return _inspect_bytes(raw, specification)
    except KernelError:
        raise
    except (
        ValueError,
        IndexError,
        OverflowError,
        csv.Error,
        OSError,
        KeyError,
        InvalidFileException,
        PyPdfError,
        XLRDError,
        BadZipFile,
        ElementTree.ParseError,
    ) as exc:
        raise KernelError(
            "material_unreadable", "原件格式损坏或当前映射无法读取", detail=str(exc)
        ) from exc


MAX_MATERIAL_BYTES = 20 * 1024 * 1024
MAX_DATA_ROWS = 100_000
MAX_COLUMNS = 64
MAX_NONEMPTY_CELLS = 2_000_000
MAX_SHEETS = 32
MAX_HEADER_ROWS = 100
MAX_CONTROL_ROWS = 1000
MAX_EXPANDED_BYTES = 200 * 1024 * 1024


def _too_large(message):
    raise KernelError("material_too_large", message)


class _TableBudget:
    """Bound the original document before retaining its normalized output."""

    def __init__(self, spec):
        self.spec = spec
        self.data_rows = 0
        self.cells = 0
        self.sheets = set()
        self.control_locations = {item.location for item in spec.controls}
        if any(row > MAX_HEADER_ROWS for row in spec.header_rows.values()):
            _too_large("每表最多支持100行表头")
        if (
            sum(len(rows) for rows in spec.total_rows.values()) + len(spec.controls)
            > MAX_CONTROL_ROWS
        ):
            _too_large("控制合计行最多支持1000行")
        for columns in (spec.columns, *spec.sheet_columns.values()):
            if any(column_index_from_string(item.column) > MAX_COLUMNS for item in columns):
                _too_large("表格最多支持64列")

    def dimensions(self, sheet, row, column):
        self.sheets.add(sheet)
        if len(self.sheets) > MAX_SHEETS:
            _too_large("单个文件最多支持32张工作表")
        max_row = MAX_DATA_ROWS + self.spec.header_rows.get(sheet, 1) + MAX_CONTROL_ROWS
        if row > max_row or column > MAX_COLUMNS:
            _too_large("表格行号或列宽超出安全范围（最多64列）")

    def accept(self, sheet, row, cells):
        self.dimensions(sheet, row, max(cells, default=0))
        self.cells += len(cells)
        if self.cells > MAX_NONEMPTY_CELLS:
            _too_large("文件最多支持200万个非空单元格")
        amount_columns = [
            item.column
            for item in self.spec.sheet_columns.get(sheet, self.spec.columns)
            if item.role == "amount"
        ]
        grouped_control_row = bool(amount_columns) and all(
            f"{sheet}!{column}{row}" in self.control_locations for column in amount_columns
        )
        if (
            cells
            and row > self.spec.header_rows.get(sheet, 1)
            and row not in self.spec.total_rows.get(sheet, ())
            and not grouped_control_row
        ):
            self.data_rows += 1
            if self.data_rows > MAX_DATA_ROWS:
                _too_large("单个文件最多支持十万条数据行，表头和控制行另计")


def _csv_rows(raw, budget):
    try:
        raw.decode("utf-8-sig")
        encoding = "utf-8-sig"
    except UnicodeDecodeError:
        encoding = "gb18030"
    with io.TextIOWrapper(io.BytesIO(raw), encoding=encoding, newline="") as source:
        for number, row in enumerate(csv.reader(source), 1):
            budget.dimensions("CSV", number, len(row))
            cells = {
                index: (value, False, False) for index, value in enumerate(row, 1) if value.strip()
            }
            budget.accept("CSV", number, cells)
            if cells:
                yield "CSV", number, cells


def _xls_rows(raw, budget):
    import xlrd

    book = xlrd.open_workbook(file_contents=raw, formatting_info=True, on_demand=True)
    try:
        for sheet_index in range(book.nsheets):
            sheet = book.sheet_by_index(sheet_index)
            budget.dimensions(sheet.name, sheet.nrows, sheet.ncols)
            for row in range(sheet.nrows):
                cells = {}
                row_info = sheet.rowinfo_map.get(row)
                for column in range(sheet.ncols):
                    cell = sheet.cell(row, column)
                    if cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                        continue
                    value = str(cell.value)
                    if cell.ctype == xlrd.XL_CELL_DATE:
                        value = (
                            xlrd.xldate_as_datetime(cell.value, book.datemode).date().isoformat()
                        )
                    elif cell.ctype == xlrd.XL_CELL_ERROR:
                        value = None
                    column_info = sheet.colinfo_map.get(column)
                    cells[column + 1] = (
                        value,
                        cell.ctype == xlrd.XL_CELL_ERROR,
                        bool(
                            sheet.visibility
                            or (row_info and row_info.hidden)
                            or (column_info and column_info.hidden)
                        ),
                    )
                budget.accept(sheet.name, row + 1, cells)
                if cells:
                    yield sheet.name, row + 1, cells
            book.unload_sheet(sheet_index)
    finally:
        book.release_resources()


def _xlsx_display_text(node):
    """Spreadsheet text consists of direct text and rich runs, never phonetic guides."""
    parts = []
    for child in node:
        if child.tag.endswith("}t"):
            parts.append(child.text or "")
        elif child.tag.endswith("}r"):
            parts.append(child.findtext("{*}t") or "")
    return "".join(parts)


def _xlsx_shared_strings(archive):
    values = []
    if "xl/sharedStrings.xml" not in archive.namelist():
        return values
    with archive.open("xl/sharedStrings.xml") as source:
        events = ElementTree.iterparse(source, events=("start", "end"))
        _, root = next(events)
        for event, node in events:
            if event == "end" and node.tag.endswith("}si"):
                values.append(_xlsx_display_text(node))
                if len(values) > MAX_NONEMPTY_CELLS:
                    _too_large("共享文本数量超出单元格安全范围")
                node.clear()
                root.remove(node)
    return values


def _xlsx_dates(archive):
    if "xl/styles.xml" not in archive.namelist():
        return set()
    from openpyxl.styles.numbers import BUILTIN_FORMATS, is_date_format

    styles = ElementTree.fromstring(archive.read("xl/styles.xml"))
    formats = dict(BUILTIN_FORMATS)
    for item in styles.findall("{*}numFmts/{*}numFmt"):
        formats[int(item.attrib["numFmtId"])] = item.attrib["formatCode"]
    return {
        index
        for index, item in enumerate(styles.findall("{*}cellXfs/{*}xf"))
        if is_date_format(formats.get(int(item.attrib.get("numFmtId", "0")), "General"))
    }


def _xlsx_rows(raw, budget):
    """Stream exact XML values once, including hidden cells and formula caches."""
    from openpyxl.utils.datetime import CALENDAR_MAC_1904, CALENDAR_WINDOWS_1900, from_excel

    with ZipFile(io.BytesIO(raw)) as archive:
        if sum(item.file_size for item in archive.infolist()) > MAX_EXPANDED_BYTES:
            _too_large("Excel解压内容最多支持200 MiB")
        if len(archive.infolist()) > 10_000:
            _too_large("Excel压缩成员数量超出安全范围")
        book = ElementTree.fromstring(archive.read("xl/workbook.xml"))
        rels = ElementTree.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
        targets = {rel.attrib["Id"]: rel.attrib["Target"] for rel in rels}
        strings, date_styles = _xlsx_shared_strings(archive), _xlsx_dates(archive)
        props = book.find("{*}workbookPr")
        epoch = (
            CALENDAR_MAC_1904
            if props is not None and props.get("date1904") in {"1", "true"}
            else CALENDAR_WINDOWS_1900
        )
        for sheet in book.findall(".//{*}sheet"):
            name = sheet.attrib["name"]
            budget.dimensions(name, 0, 0)
            rel = next(value for key, value in sheet.attrib.items() if key.endswith("}id"))
            target = targets[rel]
            path = (
                target.lstrip("/")
                if target.startswith("/")
                else posixpath.normpath(posixpath.join("xl", target))
            )
            hidden_columns = set()
            last_row = 0
            sheet_hidden = sheet.get("state", "visible") != "visible"
            with archive.open(path) as source:
                events = ElementTree.iterparse(source, events=("start", "end"))
                stack = []
                main = None
                for event, node in events:
                    if event == "start":
                        if not stack:
                            if node.tag not in {
                                "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}worksheet",
                                "{http://purl.oclc.org/ooxml/spreadsheetml/main}worksheet",
                            }:
                                raise ValueError("unknown Excel worksheet namespace")
                            main = node.tag.removesuffix("worksheet")
                        stack.append(node)
                        continue
                    # Drawing anchors also contain row/col elements. Only the
                    # worksheet's direct dimension/cols/sheetData define cells.
                    parent = stack[-2].tag if len(stack) > 1 else None
                    if node.tag == main + "dimension" and len(stack) == 2:
                        # Excel may retain a full-sheet used range after formatting.
                        # Streaming actual cells, including formulas, sets the budget.
                        range_boundaries(node.attrib["ref"])
                    elif node.tag == main + "col" and len(stack) == 3 and parent == main + "cols":
                        first = int(node.attrib["min"])
                        last = int(node.attrib["max"])
                        if not 1 <= first <= last <= 16_384:
                            raise ValueError("invalid Excel column formatting range")
                        if node.get("hidden") in {"1", "true"}:
                            hidden_columns.update(range(first, min(last, MAX_COLUMNS) + 1))
                    elif (
                        node.tag == main + "row"
                        and len(stack) == 3
                        and parent == main + "sheetData"
                    ):
                        number = int(node.attrib["r"])
                        if number <= last_row:
                            raise ValueError("Excel row locations must be unique and increasing")
                        last_row = number
                        row_hidden = node.get("hidden") in {"1", "true"}
                        cells = {}
                        seen_columns = set()
                        for cell in node.findall(main + "c"):
                            column_text, cell_row = coordinate_from_string(cell.attrib["r"])
                            column = column_index_from_string(column_text)
                            if cell_row != number or column in seen_columns:
                                raise ValueError("inconsistent or duplicate Excel cell location")
                            seen_columns.add(column)
                            kind = cell.get("t", "n")
                            value = cell.findtext(main + "v")
                            formula = cell.find(main + "f") is not None
                            if kind == "inlineStr":
                                inline = cell.find(main + "is")
                                value = _xlsx_display_text(inline) if inline is not None else ""
                            elif kind == "s" and value is not None:
                                index = int(value)
                                if not 0 <= index < len(strings):
                                    raise ValueError("invalid Excel shared-string reference")
                                value = strings[index]
                            elif kind == "b" and value is not None:
                                value = value == "1"
                            elif kind == "e":
                                value = None
                            elif kind == "n" and value and int(cell.get("s", "0")) in date_styles:
                                # Only date serials use this conversion.
                                # Money stays in its original decimal XML text.
                                date_value = from_excel(float(value), epoch=epoch)
                                value = (
                                    date_value.date().isoformat()
                                    if hasattr(date_value, "date")
                                    else date_value.isoformat()
                                )
                            missing = (formula and (value is None or value == "")) or kind == "e"
                            if value is None or not str(value).strip():
                                if not missing:
                                    continue
                            cells[column] = (
                                value,
                                missing,
                                sheet_hidden or row_hidden or column in hidden_columns,
                            )
                        if cells:
                            budget.accept(name, number, cells)
                            yield name, number, cells
                        node.clear()
                        if len(stack) > 1:
                            stack[-2].remove(node)
                    stack.pop()


def _inspect_bytes(raw: bytes, specification: Specification) -> dict:
    """Normalize one row at a time, retaining only the public inspection result."""
    spec = specification
    if len(raw) > MAX_MATERIAL_BYTES:
        _too_large("资料原件最多支持20 MiB")
    if raw.startswith(b"%PDF-") and spec.format != "pdf":
        raise KernelError("material_format_mismatch", "PDF原件必须按全部实际页核对")
    if raw.startswith(b"PK\x03\x04") and spec.format != "xlsx":
        raise KernelError("material_format_mismatch", "表格压缩原件必须按实际单元格核对")
    if raw.startswith(bytes.fromhex("d0cf11e0a1b11ae1")) and spec.format != "xls":
        raise KernelError("material_format_mismatch", "旧式Excel原件必须按实际工作表核对")
    if spec.format not in {"csv", "xls", "xlsx"}:
        return _inspect_passages(raw, spec)
    budget = _TableBudget(spec)
    rows = {"csv": _csv_rows, "xls": _xls_rows, "xlsx": _xlsx_rows}[spec.format](raw, budget)
    mappings, issues, items, coverage, totals, sums = {}, [], [], [], [], {}
    grouped_locations = {item.location for item in spec.controls}
    for sheet, row, cells in rows:
        if sheet not in mappings:
            mappings[sheet] = {
                column_index_from_string(item.column): item
                for item in spec.sheet_columns.get(sheet, spec.columns)
            }
        mapping = mappings[sheet]
        if row <= spec.header_rows.get(sheet, 1):
            coverage.extend(
                {
                    "location": f"{sheet}!{get_column_letter(column)}{row}",
                    "value": str(value),
                    "hidden": hidden,
                }
                for column, (value, _, hidden) in cells.items()
            )
            continue
        for column, rule in mapping.items():
            if rule.role == "amount" and column not in cells:
                budget.cells += 1
                if budget.cells > MAX_NONEMPTY_CELLS:
                    _too_large("原单元格与待补金额字段合计最多支持200万个")
                cells[column] = None, False, False
        period, context = None, []
        for column, (value, _, _) in cells.items():
            rule = mapping.get(column)
            if rule is not None and rule.role == "context" and value is not None:
                context.append(str(value))
            if rule is not None and rule.role == "recognition_period" and value is not None:
                try:
                    period = (
                        ActualDate(str(value)).period
                        if len(str(value)) == 10
                        else YearMonth(str(value))
                    )
                except ValueError:
                    issues.append(
                        _issue(
                            "material_period_unreadable",
                            "原行的核算所属期不能读取",
                            location=f"{sheet}!{get_column_letter(column)}{row}",
                        )
                    )
        excerpt = " · ".join(context)
        for column, (value, formula_missing, hidden) in cells.items():
            letter = get_column_letter(column)
            location = f"{sheet}!{letter}{row}"
            coverage.append({"location": location, "value": str(value), "hidden": hidden})
            rule = mapping.get(column)
            if rule is None:
                issues.append(
                    _issue("material_column_unmapped", "非空列尚未说明业务含义", location=location)
                )
                continue
            if formula_missing:
                issues.append(
                    _issue(
                        "material_formula_result_missing",
                        "公式缺少已计算结果，不能跳过",
                        location=location,
                    )
                )
            if rule.role != "amount":
                continue
            amount = None
            try:
                amount = _amount(value)
            except ValueError:
                issues.append(
                    _issue(
                        "material_amount_missing",
                        "原行金额缺失或无法读取为整数分",
                        location=location,
                    )
                )
            if amount is not None and amount < 0 and rule.funds_direction in {"inflow", "outflow"}:
                issues.append(
                    _issue(
                        "material_unsigned_amount_negative",
                        "已声明非负收入/支出列出现负数，须复核原件方向或明确冲销语义，不能取绝对值",
                        location=location,
                    )
                )
            if row in spec.total_rows.get(sheet, ()) or location in grouped_locations:
                totals.append(
                    {"location": location, "sheet": sheet, "column": letter, "expected_fen": amount}
                )
            else:
                items.append(
                    {
                        "location": location,
                        "amount_fen": amount,
                        "recognition_period": period,
                        "excerpt": excerpt,
                        "sheet": sheet,
                        "column": letter,
                        "hidden": hidden,
                    }
                    | ({"funds_direction": rule.funds_direction} if rule.funds_direction else {})
                )
                key = sheet, letter
                old = sums.get(key, 0)
                sums[key] = (
                    checked(old + amount) if old is not None and amount is not None else None
                )
    _bind_controls(items, totals, spec, issues)
    if not items:
        issues.append(
            _issue("material_no_business_columns", "未识别业务金额列，不能以空集合确认处理完成")
        )
    return {"items": items, "issues": issues, "coverage": coverage, "control_totals": totals}


def _bind_controls(items, totals, spec, issues):
    """Controls reference original leaf amounts, never other controls."""
    by_location = {item["location"]: item for item in items}
    by_column = {}
    for item in items:
        by_column.setdefault((item["sheet"], item["column"]), []).append(item)
    range_index = {}
    for key, values in by_column.items():
        ordered = sorted(values, key=lambda item: _row_number(item["location"]))
        range_index[key] = ([_row_number(item["location"]) for item in ordered], ordered)
    definitions = {control.location: control for control in spec.controls}
    found = {control["location"] for control in totals}
    for location in definitions.keys() - found:
        issues.append(
            _issue(
                "material_control_position_invalid",
                "控制金额位置必须是原件中明确映射的金额单元格",
                location=location,
            )
        )
    memberships = 0
    for control in totals:
        definition = definitions.get(control["location"])
        members = []
        if definition is None or definition.scope == "all_detail_column":
            members = list(by_column.get((control["sheet"], control["column"]), ()))
        else:
            if memberships + len(definition.locations) > MAX_NONEMPTY_CELLS:
                _too_large("控制成员总数超过200万，需减少重复控制范围")
            for location in definition.locations:
                if location not in by_location:
                    issues.append(
                        _issue(
                            "material_control_member_invalid",
                            "控制成员必须是实际业务金额，不能是控制值或虚构位置",
                            location=control["location"],
                            member_location=location,
                        )
                    )
                else:
                    members.append(by_location[location])
            for selection in definition.ranges:
                rows, values = range_index.get((selection.sheet, selection.column), ((), ()))
                selected = values[
                    bisect_left(rows, selection.first_row) : bisect_right(rows, selection.last_row)
                ]
                if memberships + len(members) + len(selected) > MAX_NONEMPTY_CELLS:
                    _too_large("控制成员总数超过200万，需减少重复控制范围")
                if not selected:
                    issues.append(
                        _issue(
                            "material_control_empty_range",
                            "控制范围没有业务金额",
                            location=control["location"],
                        )
                    )
                members.extend(selected)
        locations = [item["location"] for item in members]
        if len(set(locations)) != len(locations):
            issues.append(
                _issue(
                    "material_control_duplicate_member",
                    "同一控制不能重复计入明细",
                    location=control["location"],
                )
            )
        memberships += len(locations)
        if memberships > MAX_NONEMPTY_CELLS:
            _too_large("控制成员总数超过200万，需减少重复控制范围")
        control["member_locations"] = locations
        control["member_count"] = len(locations)
        actual = (
            sum_fen(item["amount_fen"] for item in members)
            if all(item["amount_fen"] is not None for item in members)
            else None
        )
        control["actual_fen"] = actual
        if not members or control["expected_fen"] is None or actual != control["expected_fen"]:
            issues.append(
                _issue(
                    "material_total_mismatch",
                    "原资料控制合计与明确明细范围不一致",
                    location=control["location"],
                    expected_fen=control["expected_fen"],
                    actual_fen=actual,
                )
            )


def _row_number(location):
    return int(re.search(r"[0-9]+$", location).group())


def _allocation_id(source_id):
    return "material-periods:" + digest(source_id).hex()


def _typed_assignments(assignments):
    if assignments is None or (
        isinstance(assignments, tuple)
        and all(isinstance(item, PeriodAllocationEntry) for item in assignments)
    ):
        return assignments
    return TypeAdapter(tuple[PeriodAllocationEntry, ...]).validate_json(canonical(assignments))


def _entry_locations(entry, items, index):
    if entry.location is not None:
        return [entry.location] if entry.location in items else []
    rows, locations = index.get((entry.sheet, entry.column), ((), ()))
    return locations[bisect_left(rows, entry.first_row) : bisect_right(rows, entry.last_row)]


def _entry_for_item(item, period, basis):
    common = {"recognition_period": period, "basis": basis}
    if item.get("sheet") is None:
        return PeriodAllocationEntry(location=item["location"], **common)
    row = _row_number(item["location"])
    return PeriodAllocationEntry(
        sheet=item["sheet"], column=item["column"], first_row=row, last_row=row, **common
    )


def _automatic_period(source, item):
    if item.get("recognition_period") is not None:
        return item["recognition_period"], "original_period"
    if (
        source.purpose == "supporting"
        and item["amount_fen"] is None
        and source.specification.format not in {"csv", "xls", "xlsx"}
    ):
        return source.period, "document_period"
    return None, "unknown"


def _coalesce_entries(entries):
    result = []
    ordered = sorted(
        entries,
        key=lambda item: (
            item.sheet or "",
            item.column or "",
            item.first_row or 0,
            item.location or "",
        ),
    )
    for entry in ordered:
        if result and entry.location is None:
            previous = result[-1]
            ignored = {"first_row", "last_row"}
            if (
                previous.location is None
                and previous.last_row + 1 == entry.first_row
                and previous.model_dump(exclude=ignored) == entry.model_dump(exclude=ignored)
            ):
                result[-1] = previous.model_copy(update={"last_row": entry.last_row})
                continue
        result.append(entry)
    return tuple(result)


def _period_partition(source, inspection, entries, proofs, *, complete):
    """A complete exact partition; caller labels never override an original period."""
    items = {item["location"]: item for item in inspection["items"]}
    columns = {}
    for location, item in items.items():
        if item.get("sheet") is not None:
            columns.setdefault((item["sheet"], item["column"]), []).append(
                (_row_number(location), location)
            )
    index = {key: tuple(zip(*sorted(values), strict=True)) for key, values in columns.items()}
    assigned, issues = {}, []
    for entry in entries:
        locations = _entry_locations(entry, items, index)
        overlapping = next((location for location in locations if location in assigned), None)
        if overlapping is not None:
            issues.append(
                _issue(
                    "material_allocation_overlap", "同一业务位置不能重复归属", location=overlapping
                )
            )
            break  # One overlap invalidates the partition; do not expand hostile repeated ranges.
        if not locations:
            issues.append(
                _issue(
                    "material_allocation_location_invalid",
                    "归属范围没有实际业务位置",
                    location=entry.location,
                )
            )
        for location in locations:
            if location in assigned:
                issues.append(
                    _issue(
                        "material_allocation_overlap", "同一业务位置不能重复归属", location=location
                    )
                )
                continue
            item = items[location]
            original_period, original_basis = _automatic_period(source, item)
            if original_period is not None and entry.recognition_period != original_period:
                issues.append(
                    _issue(
                        "material_allocation_period_conflict",
                        "归属与原件明确期间冲突",
                        location=location,
                    )
                )
            if entry.basis in {"original_period", "document_period"} and (
                original_period is None or entry.basis != original_basis
            ):
                issues.append(
                    _issue(
                        "material_allocation_basis_invalid",
                        "原件并未提供所声称的期间依据",
                        location=location,
                    )
                )
            if entry.basis == "confirmed_period":
                if entry.basis_evidence_digest not in proofs:
                    issues.append(
                        _issue(
                            "material_allocation_evidence_missing",
                            "归属确认依据必须是本公司已经留存的不可变证据",
                            location=location,
                        )
                    )
            assigned[location] = entry.model_copy(
                update={
                    "location": location if item.get("sheet") is None else None,
                    "sheet": item.get("sheet"),
                    "column": item.get("column"),
                    "first_row": _row_number(location) if item.get("sheet") else None,
                    "last_row": _row_number(location) if item.get("sheet") else None,
                }
            )
    for location, item in items.items():
        if location not in assigned:
            period, basis = _automatic_period(source, item)
            if complete:
                issues.append(
                    _issue("material_allocation_gap", "归属版本漏掉原件业务位置", location=location)
                )
                period, basis = None, "unknown"
            assigned[location] = _entry_for_item(item, period, basis)
    periods = {location: entry.recognition_period for location, entry in assigned.items()}
    return _coalesce_entries(assigned.values()), periods, issues


def _allocation_proofs(connection, entries):
    evidence = {entry.basis_evidence_digest for entry in entries if entry.basis_evidence_digest}
    return {
        value
        for value in evidence
        if connection.execute(
            "SELECT 1 FROM evidence WHERE digest=?", (bytes.fromhex(value),)
        ).fetchone()
    }


def _allocation_versions(connection, registry, source_id):
    return _facts(
        connection, registry, MaterialPeriodAllocation.kind, "material-source:" + source_id
    )


def _inspect_passages(raw, spec):
    pages = 1
    if spec.format == "pdf":
        from pypdf import PdfReader

        pages = len(PdfReader(io.BytesIO(raw)).pages)
    if spec.format == "image" and not (
        raw.startswith(b"\x89PNG\r\n\x1a\n") or raw.startswith(b"\xff\xd8\xff")
    ):
        raise KernelError(
            "material_image_format", "图像核对目前接收PNG或JPEG原件；多页图像须转换为完整PDF后核对"
        )
    issues = []
    represented = {item.page for item in spec.passages}
    if spec.all_pages_reviewed is not True or represented != set(range(1, pages + 1)):
        issues.append(
            _issue(
                "material_pages_unreviewed",
                "须明确阅读全部页或图片，并逐页登记位置与摘录",
                pages=pages,
            )
        )
    items = []
    for item in spec.passages:
        if not item.excerpt.strip():
            issues.append(
                _issue("material_passage_empty", "原文摘录不能空白", location=item.location)
            )
        if spec.format == "text":
            try:
                text = raw.decode("utf-8-sig")
            except UnicodeDecodeError:
                text = raw.decode("gb18030")
            if item.excerpt not in text:
                issues.append(
                    _issue("material_excerpt_missing", "摘录不在原文中", location=item.location)
                )
        items.append(
            {
                "location": item.location,
                "page": item.page,
                "excerpt": item.excerpt,
                "amount_fen": item.amount_fen,
                "recognition_period": item.recognition_period,
            }
        )
    return {
        "items": items,
        "issues": issues,
        "coverage": [item.model_dump(mode="json") for item in spec.passages],
        "control_totals": [],
        "semantic_review": "人工或AI逐页阅读确认；程序只检查位置覆盖，不证明自由文本被完整理解",
    }


def _facts(connection, registry, kind, scope):
    if kind not in registry.models:
        return ()
    reader = SimpleNamespace(registry=registry)
    rows = connection.execute(
        "SELECT s.fact_id FROM fact_scope s JOIN fact_current c ON c.fact_id=s.fact_id "
        "WHERE s.kind=? AND s.scope_key=? ORDER BY c.subject_id",
        (kind, scope),
    ).fetchall()
    return tuple(Store.fact(reader, connection, row[0]) for row in rows)


def _amount_basis(fact, calculation, path):
    """Same field names and explicitly declared aliases share one economic capacity."""
    prefix, field = path.split(".")
    value = getattr(fact, field, None) if prefix == "fact" else calculation.values.get(field)
    if type(value) is not int:
        raise KernelError("material_amount_field_invalid", "关联字段不是明确的业务金额")
    aliases = {}
    for original, replacement in fact.material_amount_aliases.items():
        if not all(
            re.fullmatch(r"(?:fact|result)\.[a-z][a-z0-9_]*_fen", value)
            for value in (original, replacement)
        ):
            raise KernelError("material_amount_alias_invalid", "领域金额别名声明无效")
        original_field, replacement_field = original.split(".")[1], replacement.split(".")[1]
        if original_field == replacement_field:
            continue
        if original_field in aliases and aliases[original_field] != replacement_field:
            raise KernelError("material_amount_alias_invalid", "同一金额名称不能指向不同金额维度")
        aliases[original_field] = replacement_field

    def normalized(field):
        visited = set()
        while field in aliases:
            if field in visited:
                raise KernelError("material_amount_alias_invalid", "领域金额别名不能形成循环")
            visited.add(field)
            field = aliases[field]
        return field

    requested = normalized(path.split(".")[1])
    members = {}
    for prefix, values in (
        ("fact", {name: getattr(fact, name) for name in type(fact).model_fields}),
        ("result", calculation.values),
    ):
        for field, value in values.items():
            if field.endswith("_fen") and type(value) is int and normalized(field) == requested:
                members[f"{prefix}.{field}"] = value
    if len(set(members.values())) > 1:
        raise KernelError(
            "material_amount_basis_conflict",
            "同一业务金额的事实与结果表示不一致",
            amount_fields=sorted(members),
        )
    return requested


def _current_material_result(connection, registry, subject):
    row = connection.execute(
        "SELECT f.fact_id current_fact_id,c.* FROM fact_current f "
        "LEFT JOIN calculation_current a ON a.subject_id=f.subject_id "
        "LEFT JOIN calculation c ON c.id=a.calculation_id WHERE f.subject_id=?",
        (subject,),
    ).fetchone()
    if row is None:
        return None
    reader = SimpleNamespace(registry=registry)
    return Store.fact(reader, connection, row["current_fact_id"]), (
        Store.calculation(row) if row["id"] else None
    )


def _group_periods(version):
    return frozenset(link.recognition_period for link in version.fact.links)


def _groups_cover_unknown(allocation, groups):
    """Decide discovery from saved locations only; never read historical BLOBs."""
    if not allocation.fact.entries:
        return False
    locations = {member.location for group in groups for member in group.fact.members}
    columns = {}
    for location in locations:
        match = re.fullmatch(r"(.*)!([A-Z]+)([1-9][0-9]*)", location)
        if match:
            columns.setdefault((match[1], match[2]), []).append(int(match[3]))
    columns = {key: sorted(values) for key, values in columns.items()}
    for entry in allocation.fact.entries:
        if entry.recognition_period is not None:
            continue
        if entry.location is not None:
            if entry.location not in locations:
                return False
        else:
            rows = columns.get((entry.sheet, entry.column), ())
            if (
                bisect_right(rows, entry.last_row) - bisect_left(rows, entry.first_row)
                != entry.last_row - entry.first_row + 1
            ):
                return False
    return True


def _group_issues(connection, registry, version, source, inspection, *, current=None):
    """Validate a pool as a whole; its members never acquire invented individual links."""
    fact = version.fact
    issues = []

    def problem(code, message, **details):
        issues.append(_issue(code, message, group_id=version.subject_id, **details))

    if (
        source is None
        or source.fact.kind != MaterialSource.kind
        or source.id != fact.source_fact_id
    ):
        problem("material_source_changed", "联合组必须绑定当前原件的确切版本")
        return issues
    if (
        source.fact.purpose != "business"
        or source.fact.evidence_digest not in version.evidence
        or fact.basis_evidence_digest not in version.evidence
        or fact.joint_basis_confirmed is not True
        or not fact.reason.strip()
    ):
        problem("material_group_basis_missing", "联合组必须有完整原件和明确核准共同形成的依据")
    if fact.period != source.fact.period:
        problem("material_group_source_period", "组登记月份必须与原件登记月份相同")
    for proof in (source.fact.evidence_digest, fact.basis_evidence_digest):
        if not connection.execute(
            "SELECT 1 FROM evidence WHERE digest=?", (bytes.fromhex(proof),)
        ).fetchone():
            problem("material_group_basis_missing", "联合组依据尚未保全")
    items = {item["location"]: item for item in inspection["items"]}
    controls = {item["location"] for item in inspection["control_totals"]}
    selected = {member.location for member in fact.members}
    linked_periods = {link.recognition_period for link in fact.links}
    for member in fact.members:
        item = items.get(member.location)
        if item is None or member.location in controls:
            problem(
                "material_group_member_invalid",
                "组成员必须是真实业务明细，不能是控制合计",
                location=member.location,
            )
        elif item["amount_fen"] is not None and item["amount_fen"] != member.amount_fen:
            problem("material_amount_conflict", "组成员金额与原行不一致", location=member.location)
        elif (
            known_period := _automatic_period(source.fact, item)[0]
        ) is not None and linked_periods != {known_period}:
            problem(
                "material_group_period_already_known",
                "原行已有明确期间，联合组全部结果必须属于该期间",
                location=member.location,
            )
    allocations = _allocation_versions(connection, registry, fact.source_id)
    if len(allocations) != 1 or allocations[0].fact.source_fact_id != source.id:
        problem("material_allocation_required", "原件须先建立完整归属清单")
    else:
        # Known individual periods cannot be silently replaced by a collective pool.
        for entry in allocations[0].fact.entries:
            if entry.recognition_period is None or linked_periods == {entry.recognition_period}:
                continue
            if (
                entry.location in selected
                if entry.location
                else any(
                    f"{entry.sheet}!{entry.column}{row}" in selected
                    for row in range(entry.first_row, entry.last_row + 1)
                )
            ):
                problem(
                    "material_group_period_already_known", "联合组全部结果须遵守已有单项确认期间"
                )
                break
    for kind in (MaterialResolution.kind, MaterialGroupResolution.kind):
        for other in _facts(connection, registry, kind, "material-source:" + fact.source_id):
            if other.subject_id == version.subject_id or other.fact.source_fact_id != source.id:
                continue
            other_locations = (
                {other.fact.location}
                if kind == MaterialResolution.kind
                else {member.location for member in other.fact.members}
            )
            overlap = selected & other_locations
            if overlap:
                problem(
                    "material_group_overlap",
                    "同一原行不能同时属于逐行处置或另一联合组",
                    location=min(overlap),
                )
    current = current or (lambda subject: _current_material_result(connection, registry, subject))
    totals, targets = {}, {}
    for link in fact.links:
        target = current(link.subject_id)
        if target is None or target[1] is None:
            problem(
                "material_result_missing",
                "联合组关联业务没有当前正式结果",
                subject_id=link.subject_id,
            )
            continue
        original, calculation = target
        if (
            original.id != link.fact_id
            or original.fact.kind != link.fact_kind
            or calculation.id != link.calculation_id
            or calculation.fact_id != original.id
            or connection.execute(
                "SELECT 1 FROM pending WHERE subject_id=? LIMIT 1", (link.subject_id,)
            ).fetchone()
        ):
            problem(
                "material_result_stale",
                "联合组关联事实或正式结果已经变化",
                subject_id=link.subject_id,
            )
            continue
        if calculation.period != link.recognition_period:
            problem(
                "material_period_mismatch",
                "联合组月份必须来自所绑定正式结果",
                subject_id=link.subject_id,
            )
        try:
            original.fact.validate_material_amount(
                link.amount_field,
                link.amount_fen,
                source_amounts=tuple(member.amount_fen for member in fact.members),
                source_directions=tuple(
                    items.get(member.location, {}).get("funds_direction") for member in fact.members
                ),
            )
            basis = _amount_basis(original.fact, calculation, link.amount_field)
        except KernelError as error:
            problem(error.code, str(error), subject_id=link.subject_id)
            continue
        prefix, field = link.amount_field.split(".")
        capacity = getattr(original.fact, field) if prefix == "fact" else calculation.values[field]
        key = link.subject_id, basis
        totals[key] = sum_fen((totals.get(key, 0), abs(link.amount_fen)))
        targets[key] = original, calculation, capacity
    for (subject_id, basis), (original, calculation, capacity) in targets.items():
        allocated = totals[subject_id, basis]
        for kind in (MaterialResolution.kind, MaterialGroupResolution.kind):
            for other in _facts(connection, registry, kind, "material-business:" + subject_id):
                if other.subject_id == version.subject_id:
                    continue
                active_source = connection.execute(
                    "SELECT fact_id FROM fact_current WHERE subject_id=?", (other.fact.source_id,)
                ).fetchone()
                if active_source is None or active_source[0] != other.fact.source_fact_id:
                    continue
                for used in other.fact.links:
                    if (used.subject_id, used.fact_id, used.fact_kind, used.calculation_id) != (
                        subject_id,
                        original.id,
                        original.fact.kind,
                        calculation.id,
                    ):
                        continue
                    try:
                        if _amount_basis(original.fact, calculation, used.amount_field) == basis:
                            allocated = sum_fen((allocated, abs(used.amount_fen)))
                    except KernelError:
                        pass  # Its own active row/group reports the stale invalid association.
        if allocated > abs(capacity):
            problem(
                "material_business_overallocated",
                "联合组与其他原资料合计超过同一业务金额",
                subject_id=subject_id,
            )
    return issues


def check_completeness(connection, month: int, registry) -> dict:
    """Read the bound snapshot; the returned digest belongs in the close manifest."""
    period = YearMonth.from_ordinal(month)
    by_id = {
        item.subject_id: item
        for item in _facts(connection, registry, MaterialSource.kind, str(period))
    }
    current_sources = dict(
        connection.execute(
            "SELECT c.subject_id,c.fact_id FROM fact_current c JOIN subject s ON s.id=c.subject_id "
            "WHERE s.kind=?",
            (MaterialSource.kind,),
        ).fetchall()
    )
    group_cache = {}

    def source_groups(source_id):
        if source_id not in group_cache:
            group_cache[source_id] = _facts(
                connection, registry, MaterialGroupResolution.kind, "material-source:" + source_id
            )
        return group_cache[source_id]

    group_versions = {
        item.id: item
        for item in _facts(
            connection, registry, MaterialGroupResolution.kind, "material-period:" + str(period)
        )
        if current_sources.get(item.fact.source_id) == item.fact.source_fact_id
    }
    allocation_versions = {
        item.id: item
        for scope in ("material-period:" + str(period), "material-period-unknown")
        for item in _facts(connection, registry, MaterialPeriodAllocation.kind, scope)
        if current_sources.get(item.fact.source_id) == item.fact.source_fact_id
    }
    # A v2 source has no partition. A changed source invalidates its old partition.
    # Collective pools own finite months without inventing individual row periods.
    for ident, allocation in list(allocation_versions.items()):
        if any(entry.recognition_period == period for entry in allocation.fact.entries):
            continue
        groups = [
            item
            for item in source_groups(allocation.fact.source_id)
            if item.fact.source_fact_id == allocation.fact.source_fact_id
            and item.fact.joint_basis_confirmed is True
        ]
        if _groups_cover_unknown(allocation, groups):
            del allocation_versions[ident]
    # A v2 source has no partition. A changed source invalidates its old partition.
    # Discover these using current identity metadata, without opening historical BLOBs.
    mapped_sources = set()
    if MaterialPeriodAllocation.kind in registry.models:
        mapped_sources = {
            row[0]
            for row in connection.execute(
                f"SELECT a.source_id FROM {table_name(MaterialPeriodAllocation.kind)} a "
                "JOIN fact_current m ON m.fact_id=a.revision_id "
                "JOIN fact_current s ON s.subject_id=a.source_id AND s.fact_id=a.source_fact_id"
            )
        }
    resolution_versions = {
        item.id: item
        for scope in (str(period), "material-period:" + str(period))
        for item in _facts(connection, registry, MaterialResolution.kind, scope)
        if item.fact.source_id in current_sources
    }
    unallocated = current_sources.keys() - mapped_sources
    relevant = (
        set(by_id)
        | {item.fact.source_id for item in resolution_versions.values()}
        | {item.fact.source_id for item in allocation_versions.values()}
        | {item.fact.source_id for item in group_versions.values()}
    )
    by_item, issues, parsed, parsed_items = {}, [], {}, {}
    for source_id in sorted(unallocated):
        issues.append(
            _issue(
                "material_allocation_required",
                "原件尚未建立与当前来源一致的归属版本",
                source_id=source_id,
            )
        )
    relevant.difference_update(unallocated)
    legacy = {
        row[0].hex()
        for row in connection.execute(
            "SELECT DISTINCT i.evidence_digest FROM material_item i JOIN material_revision m "
            "ON m.id=i.inventory_id WHERE m.period=?",
            (month,),
        )
    }
    for evidence in sorted(legacy):
        matches = _facts(connection, registry, MaterialSource.kind, "material-evidence:" + evidence)
        if not matches:
            issues.append(
                _issue(
                    "material_source_not_registered",
                    "已接收原件尚未登记逐项核对来源",
                    evidence_digest=evidence,
                )
            )
        by_id.update((item.subject_id, item) for item in matches)
        relevant.update(item.subject_id for item in matches)
    relevant.difference_update(unallocated)
    queue, loaded = list(sorted(relevant)), set()
    while queue:
        source_id = queue.pop()
        if source_id in unallocated:
            continue
        if source_id in loaded:
            continue
        loaded.add(source_id)
        if source_id not in by_id:
            matches = _facts(connection, registry, MaterialSource.kind, "@" + source_id)
            if not matches:
                issues.append(
                    _issue(
                        "material_source_missing",
                        "逐项处置的原件来源不存在或已撤去",
                        source_id=source_id,
                    )
                )
                continue
            by_id[source_id] = matches[0]
        version = by_id[source_id]
        fact = version.fact
        if (
            len(
                _facts(
                    connection,
                    registry,
                    MaterialSource.kind,
                    "material-evidence:" + fact.evidence_digest,
                )
            )
            != 1
        ):
            issues.append(
                _issue(
                    "material_duplicate_source", "同一原件只能有一个来源身份", source_id=source_id
                )
            )
        for item in _facts(
            connection, registry, MaterialResolution.kind, "material-source:" + source_id
        ):
            resolution_versions[item.id] = item
            if item.fact.treatment == "duplicate" and item.fact.duplicate_source_id:
                queue.append(item.fact.duplicate_source_id)
        group_versions.update((item.id, item) for item in source_groups(source_id))
        row = connection.execute(
            "SELECT content FROM evidence WHERE digest=?", (bytes.fromhex(fact.evidence_digest),)
        ).fetchone()
        if row is None or fact.evidence_digest not in version.evidence:
            issues.append(
                _issue(
                    "material_source_evidence_missing",
                    "来源必须引用留存的原件证据",
                    source_id=source_id,
                )
            )
            continue
        try:
            parsed[source_id] = inspect_bytes(row[0], fact.specification)
            parsed_items[source_id] = {
                item["location"]: item for item in parsed[source_id]["items"]
            }
        except (ValueError, OSError, KeyError) as exc:
            issues.append(
                _issue(
                    "material_unreadable",
                    "原件或解析映射无效",
                    source_id=source_id,
                    detail=str(exc),
                )
            )
    for item in resolution_versions.values():
        by_item.setdefault((item.fact.source_id, item.fact.location), []).append(item)
    by_group_member = {}
    group_amounts = {}
    for group in group_versions.values():
        if current_sources.get(group.fact.source_id) != group.fact.source_fact_id:
            continue
        for member in group.fact.members:
            key = group.fact.source_id, member.location
            by_group_member.setdefault(key, []).append(group)
            group_amounts[group.id, member.location] = member.amount_fen
    reader = SimpleNamespace(registry=registry)
    current_cache, capacities, active_capacities, coverage = {}, {}, set(), []
    item_periods, partition_issues = {}, {}
    for source_id, inspection in parsed.items():
        source = by_id[source_id]
        versions = _allocation_versions(connection, registry, source_id)
        allocation_versions.update((item.id, item) for item in versions)
        if len(versions) == 1 and versions[0].fact.source_fact_id == source.id:
            allocation = versions[0]
            proofs = _allocation_proofs(connection, allocation.fact.entries)
            proofs.intersection_update(allocation.evidence)
            _, periods, errors = _period_partition(
                source.fact, inspection, allocation.fact.entries, proofs, complete=True
            )
            if source.fact.evidence_digest not in allocation.evidence:
                errors.append(
                    _issue("material_allocation_evidence_missing", "归属版本必须引用实际原件")
                )
            if allocation.fact.period != source.fact.period:
                errors.append(
                    _issue("material_allocation_source_changed", "归属版本接收月与来源不一致")
                )
        else:
            _, periods, errors = _period_partition(
                source.fact, inspection, (), set(), complete=False
            )
            errors.append(
                _issue(
                    "material_allocation_source_changed",
                    "归属版本与当前原件版本不一致，须重新核对",
                )
            )
            periods = {key: None for key in periods}
        item_periods[source_id], partition_issues[source_id] = periods, errors

    def recovered_amount(source_id, location):
        matches = by_item.get((source_id, location), ())
        if len(matches) == 1:
            resolution = matches[0].fact
            if (
                resolution.source_fact_id == by_id[source_id].id
                and resolution.amount_fen is not None
                and (resolution.reason or "").strip()
            ):
                return resolution.amount_fen
        groups = by_group_member.get((source_id, location), ())
        if len(groups) == 1 and groups[0].fact.source_fact_id == by_id[source_id].id:
            return group_amounts[groups[0].id, location]
        return None

    def inspection_issues(source_id, inspection):
        repaired = {
            item["location"]
            for item in inspection["items"]
            if item["amount_fen"] is None
            and recovered_amount(source_id, item["location"]) is not None
        }
        controls = []
        for control in inspection["control_totals"]:
            expected = control["expected_fen"]
            matches = by_item.get((source_id, control["location"]), ())
            if (
                expected is None
                and len(matches) == 1
                and matches[0].fact.treatment == "control_total"
            ):
                expected = recovered_amount(source_id, control["location"])
                if expected is not None:
                    repaired.add(control["location"])
            amounts = [
                item["amount_fen"]
                if item["amount_fen"] is not None
                else recovered_amount(source_id, item["location"])
                for location in control["member_locations"]
                for item in (parsed_items[source_id][location],)
            ]
            actual = sum_fen(amounts) if all(value is not None for value in amounts) else None
            if expected is None or actual != expected:
                controls.append(
                    _issue(
                        "material_total_mismatch",
                        "控制合计与已核定明细不一致",
                        location=control["location"],
                        expected_fen=expected,
                        actual_fen=actual,
                    )
                )
            if matches and (
                len(matches) != 1
                or matches[0].fact.treatment != "control_total"
                or matches[0].fact.links
                or matches[0].fact.amount_fen != expected
                or matches[0].fact.source_fact_id != by_id[source_id].id
            ):
                controls.append(
                    _issue(
                        "material_control_invalid",
                        "控制行只能明确确认控制合计，不能作为新增业务或覆盖原值",
                        location=control["location"],
                    )
                )
        return [
            item
            for item in inspection["issues"]
            if item["code"] != "material_total_mismatch"
            and not (
                item["location"] in repaired
                and item["code"] in {"material_amount_missing", "material_formula_result_missing"}
            )
        ] + controls

    def current(subject):
        if subject not in current_cache:
            row = connection.execute(
                "SELECT f.fact_id current_fact_id,c.* FROM fact_current f "
                "LEFT JOIN calculation_current a ON a.subject_id=f.subject_id "
                "LEFT JOIN calculation c ON c.id=a.calculation_id WHERE f.subject_id=?",
                (subject,),
            ).fetchone()
            if row is None:
                current_cache[subject] = None
            else:
                fact_version = Store.fact(reader, connection, row["current_fact_id"])
                current_cache[subject] = fact_version, Store.calculation(row) if row["id"] else None
        return current_cache[subject]

    group_errors = {}

    def collective(key):
        matches = by_group_member.get(key, ())
        if not matches:
            return None
        if len(matches) != 1 or key in by_item:
            return [_issue("material_group_overlap", "原行同时存在多个处置归属", location=key[1])]
        group = matches[0]
        if group.id not in group_errors:
            group_errors[group.id] = _group_issues(
                connection,
                registry,
                group,
                by_id.get(key[0]),
                parsed.get(key[0], {}),
                current=current,
            )
        return group_errors[group.id]

    def resolve(key, trail=(), *, active=True):
        if key in trail:
            return [_issue("material_duplicate_cycle", "重复来源引用形成循环", location=key[1])]
        grouped = collective(key)
        if grouped is not None:
            return grouped
        source = by_id.get(key[0])
        item = parsed_items.get(key[0], {}).get(key[1])
        matches = by_item.get(key, ())
        if source is None or item is None or len(matches) != 1:
            return [
                _issue(
                    "material_item_unresolved",
                    "原行需要唯一的明确处置",
                    source_id=key[0],
                    location=key[1],
                )
            ]
        version, source_fact = matches[0], source.fact
        fact = version.fact
        result = []
        if fact.source_fact_id != source.id:
            return [
                _issue(
                    "material_source_changed", "解析来源版本变化，须复核原行处置", location=key[1]
                )
            ]
        amount = item["amount_fen"] if item["amount_fen"] is not None else fact.amount_fen
        if (
            item["amount_fen"] is not None
            and fact.amount_fen is not None
            and fact.amount_fen != item["amount_fen"]
        ):
            result.append(
                _issue("material_amount_conflict", "处置金额与原行不一致", location=key[1])
            )
        if fact.treatment in {"supporting", "no_accounting"}:
            if not (fact.reason or "").strip() or (
                fact.treatment == "no_accounting" and fact.non_accounting_reason is None
            ):
                result.append(
                    _issue(
                        "material_non_accounting_basis",
                        "不入账处置需要明确理由和依据",
                        location=key[1],
                    )
                )
            if fact.links:
                result.append(
                    _issue(
                        "material_unexpected_links",
                        "不入账处置不能同时分配核算金额",
                        location=key[1],
                    )
                )
            if fact.treatment == "supporting" and (
                source_fact.purpose != "supporting" or amount is not None
            ):
                result.append(
                    _issue(
                        "material_supporting_conflict",
                        "业务原件须明确业务处置，不能以支持证据跳过",
                        location=key[1],
                    )
                )
            if fact.non_accounting_reason == "zero_amount" and amount != 0:
                result.append(
                    _issue("material_zero_conflict", "非零业务不能按零金额跳过", location=key[1])
                )
            return result
        if fact.treatment == "control_total":
            return [
                _issue("material_control_invalid", "业务明细不能当作控制合计跳过", location=key[1])
            ]
        if amount is None:
            return [_issue("material_amount_unconfirmed", "原文金额尚未明确确认", location=key[1])]
        if fact.treatment == "duplicate":
            target = fact.duplicate_source_id, fact.duplicate_location
            if not all(target) or fact.links or not (fact.reason or "").strip():
                return [
                    _issue(
                        "material_duplicate_basis",
                        "重复资料须明确原业务位置和依据",
                        location=key[1],
                    )
                ]
            target_item = parsed_items.get(target[0], {}).get(target[1])
            target_amount = (
                target_item["amount_fen"]
                if target_item is not None and target_item["amount_fen"] is not None
                else recovered_amount(*target)
                if target[0] in by_id
                else None
            )
            if target_item is None or target_amount != amount:
                return [
                    _issue(
                        "material_duplicate_amount", "重复资料金额与原业务不一致", location=key[1]
                    )
                ]
            if item_periods.get(target[0], {}).get(target[1]) != item_periods.get(key[0], {}).get(
                key[1]
            ):
                return [
                    _issue(
                        "material_duplicate_period", "重复资料须属于同一业务期间", location=key[1]
                    )
                ]
            return resolve(target, (*trail, key), active=active)
        if not fact.links:
            return [_issue("material_links_missing", "业务原行尚未关联正式处理", location=key[1])]
        if sum_fen(link.amount_fen for link in fact.links) != amount:
            result.append(
                _issue(
                    "material_split_mismatch", "拆分或关联金额之和必须等于原行金额", location=key[1]
                )
            )
        for link in fact.links:
            target = current(link.subject_id)
            expected_period = item_periods.get(key[0], {}).get(key[1])
            if target is None or target[1] is None:
                result.append(
                    _issue(
                        "material_result_missing",
                        "关联业务没有当前正式结果",
                        location=key[1],
                        subject_id=link.subject_id,
                    )
                )
                continue
            fact_version, calculation = target
            if (
                fact_version.id != link.fact_id
                or fact_version.fact.kind != link.fact_kind
                or calculation.id != link.calculation_id
                or calculation.fact_id != fact_version.id
                or connection.execute(
                    "SELECT 1 FROM pending WHERE subject_id=? LIMIT 1", (link.subject_id,)
                ).fetchone()
            ):
                result.append(
                    _issue(
                        "material_result_stale",
                        "关联事实、结果已变化或已失效",
                        location=key[1],
                        subject_id=link.subject_id,
                    )
                )
                continue
            if calculation.period != link.recognition_period or (
                expected_period is not None and expected_period != link.recognition_period
            ):
                result.append(
                    _issue(
                        "material_period_mismatch",
                        "原行明确归属期与关联业务不一致",
                        location=key[1],
                    )
                )
            if fact.treatment == "other_period" and link.recognition_period == source_fact.period:
                result.append(
                    _issue(
                        "material_transfer_period", "转入其他期间须明确不同归属月", location=key[1]
                    )
                )
            prefix, field = link.amount_field.split(".")
            value = (
                getattr(fact_version.fact, field, None)
                if prefix == "fact"
                else calculation.values.get(field)
            )
            if type(value) is not int or not field.endswith("_fen"):
                result.append(
                    _issue(
                        "material_amount_field_invalid",
                        "关联字段不是明确的业务金额",
                        location=key[1],
                    )
                )
                continue
            try:
                fact_version.fact.validate_material_amount(
                    link.amount_field,
                    link.amount_fen,
                    source_amounts=(amount,) if type(amount) is int else (),
                    source_directions=(item.get("funds_direction"),),
                )
                basis = _amount_basis(fact_version.fact, calculation, link.amount_field)
            except KernelError as error:
                result.append(_issue(error.code, str(error), location=key[1], **error.details))
                continue
            capacity_key = link.subject_id, basis
            capacities[capacity_key] = abs(value)
            if active:
                active_capacities.add(capacity_key)
        return result

    file_summaries, file_diagnostics = [], {}
    for source_id in sorted(relevant):
        source = by_id.get(source_id)
        if source is None or source_id not in parsed:
            continue
        inspection = parsed[source_id]
        periods = item_periods[source_id]
        source_errors = partition_issues[source_id]
        issues.extend(dict(item, source_id=source_id) for item in source_errors)
        control_members = {
            control["location"]: control["member_locations"]
            for control in inspection["control_totals"]
        }
        inspected_errors = inspection_issues(source_id, inspection)
        file_diagnostics[source_id] = [*source_errors, *inspected_errors]
        for error in inspected_errors:
            location = error.get("location")
            if error["code"] in {
                "material_total_mismatch",
                "material_control_invalid",
                "material_amount_missing",
                "material_formula_result_missing",
            }:
                members = control_members.get(location, [location])
                related = {periods.get(member) for member in members}
                if related and period not in related and None not in related:
                    continue
            issues.append(dict(error, source_id=source_id))
        file_unprocessed = 0
        unknown_count = 0
        for item in inspection["items"]:
            allocated_period = periods[item["location"]]
            groups = by_group_member.get((source_id, item["location"]), ())
            shared_periods = (
                frozenset().union(*(_group_periods(group) for group in groups))
                if groups
                else frozenset()
            )
            active = (
                period in shared_periods
                if groups
                else allocated_period is None or allocated_period == period
            )
            item_issues = [
                dict(issue, source_id=source_id)
                for issue in resolve((source_id, item["location"]), active=active)
            ]
            if allocated_period is None and not groups:
                unknown_count += 1
                item_issues.append(
                    _issue(
                        "material_period_unknown",
                        "原行的公司核算所属期尚无依据",
                        source_id=source_id,
                        location=item["location"],
                    )
                )
            file_unprocessed += bool(item_issues)
            file_diagnostics[source_id].extend(item_issues)
            if active:
                issues.extend(item_issues)
                coverage.append(
                    {
                        "source_id": source_id,
                        "source_fact_id": source.id,
                        "location": item["location"],
                        "amount_fen": item["amount_fen"],
                        "recognition_period": allocated_period,
                        **(
                            {
                                "joint_periods": sorted(shared_periods),
                                "group_ids": sorted(group.subject_id for group in groups),
                            }
                            if groups
                            else {}
                        ),
                        "complete": not item_issues,
                    }
                )
        file_summaries.append(
            {
                "source_id": source_id,
                "source_fact_id": source.id,
                "item_count": len(inspection["items"]),
                "unprocessed_count": file_unprocessed,
                "unknown_period_count": unknown_count,
                "issue_count": len(inspected_errors) + len(source_errors),
                "status": "complete"
                if not (file_unprocessed or inspected_errors or source_errors)
                else "needs_information",
            }
        )
        locations = {
            item["location"] for item in (*inspection["items"], *inspection["control_totals"])
        }
        for key in by_item:
            if key[0] == source_id and key[1] not in locations:
                location_issue = _issue(
                    "material_location_unknown",
                    "处置位置不在原件业务行中",
                    source_id=source_id,
                    location=key[1],
                )
                issues.append(location_issue)
                file_diagnostics[source_id].append(location_issue)
    # Lookup other allocations by the touched business identity. Historical
    # documents are not opened merely to enforce a numeric allocation capacity.
    for key, capacity in capacities.items():
        competitors = tuple(
            item
            for kind in (MaterialResolution.kind, MaterialGroupResolution.kind)
            for item in _facts(connection, registry, kind, "material-business:" + key[0])
            if current_sources.get(item.fact.source_id) == item.fact.source_fact_id
        )
        resolution_versions.update((item.id, item) for item in competitors)
        fact_version, calculation = current(key[0])
        amount = 0
        for item in competitors:
            if isinstance(item.fact, MaterialResolution) and item.fact.treatment not in {
                "recognize",
                "other_period",
            }:
                continue
            for link in item.fact.links:
                if (
                    link.subject_id != key[0]
                    or link.fact_id != fact_version.id
                    or link.fact_kind != fact_version.fact.kind
                    or link.calculation_id != calculation.id
                ):
                    continue
                try:
                    basis = _amount_basis(fact_version.fact, calculation, link.amount_field)
                except KernelError:
                    continue  # Its own row reports the inconsistent business result.
                if basis == key[1]:
                    amount += abs(link.amount_fen)
        if amount > capacity:
            capacity_issue = _issue(
                "material_business_overallocated",
                "多条原资料重复占用了同一业务金额，须明确重复关系",
                subject_id=key[0],
                amount_field=key[1],
            )
            if key in active_capacities:
                issues.append(capacity_issue)
            for source_id in {item.fact.source_id for item in competitors}:
                if source_id in file_diagnostics:
                    file_diagnostics[source_id].append(capacity_issue)
    for summary in file_summaries:
        summary["issue_count"] = len(file_diagnostics[summary["source_id"]])
        summary["status"] = "needs_information" if summary["issue_count"] else "complete"
    for source_id in sorted(unallocated):
        file_summaries.append(
            {
                "source_id": source_id,
                "source_fact_id": current_sources[source_id],
                "status": "needs_information",
                "allocation_required": True,
            }
        )
    for issue in issues:
        source = by_id.get(issue.get("source_id"))
        if source is not None:
            issue["category"] = source.fact.category
            issue["field"] = "materials." + source.fact.category
    result = {
        "period": period,
        "issues": issues,
        "coverage": coverage,
        "source_versions": sorted(item.id for item in by_id.values()),
        "resolution_versions": sorted(resolution_versions),
        "group_versions": sorted(group_versions),
        "allocation_versions": sorted(allocation_versions),
        "file_summaries": file_summaries,
        "file_status": "complete"
        if not issues and all(item["status"] == "complete" for item in file_summaries)
        else "needs_information",
    }
    result["fact_ids"] = sorted(
        set(result["source_versions"])
        | set(result["resolution_versions"])
        | set(result["allocation_versions"])
        | set(result["group_versions"])
    )
    return {
        **result,
        "status": "complete" if not issues else "needs_information",
        "coverage_digest": digest(result).hex(),
    }


class Materials:
    def __init__(self, engine):
        self.engine, self.store = engine, engine.store

    def _inspection(self, evidence_digest, spec):
        with self.store.connection(read_only=True) as connection:
            row = connection.execute(
                "SELECT content FROM evidence WHERE digest=?", (bytes.fromhex(evidence_digest),)
            ).fetchone()
            if row is None:
                raise NeedsInformation("evidence_digest", "先留存原件")
        # All source bytes have been copied; parsing does not hold a SQLite read transaction.
        return inspect_bytes(row[0], spec)

    def inspect(
        self,
        evidence_digest: str,
        specification: dict,
        *,
        after: str | None = None,
        limit: PageLimit = 100,
    ):
        """Bounded public rows; the cursor binds original bytes, mapping and parser build."""
        _page_limit(limit)
        spec = Specification.model_validate_json(canonical(specification))
        spec_digest = digest(spec.model_dump(mode="json")).hex()
        inspection_digest = digest(
            {
                "evidence_digest": evidence_digest,
                "specification_digest": spec_digest,
                "parser_build": calculator_build_id(),
            }
        ).hex()
        offset = 0
        if after is not None:
            if not isinstance(after, str) or not re.fullmatch(r"[0-9a-f]{64}:[0-9]{1,8}", after):
                raise KernelError("material_cursor_invalid", "资料游标无效")
            binding, position = after.split(":")
            if binding != inspection_digest:
                raise KernelError(
                    "material_cursor_stale", "原件、解析映射或程序已变化，请重新读取首页"
                )
            offset = int(position)
        result = self._inspection(evidence_digest, spec)
        if offset > len(result["items"]):
            raise KernelError("material_cursor_invalid", "资料游标超出明细范围")
        page = result["items"][offset : offset + limit]
        end = offset + len(page)
        return {
            "evidence_digest": evidence_digest,
            "specification_digest": spec_digest,
            "inspection_digest": inspection_digest,
            "status": "needs_information" if result["issues"] else "ready",
            "summary": {
                "item_count": len(result["items"]),
                "cell_count": len(result["coverage"]),
                "issue_count": len(result["issues"]),
                "control_total_count": len(result["control_totals"]),
            },
            "items": page,
            "coverage": [
                {"location": item["location"], "hidden": item.get("hidden", False)} for item in page
            ],
            "issues": result["issues"][:limit],
            "issues_truncated": len(result["issues"]) > limit,
            "control_totals": [
                {key: value for key, value in control.items() if key != "member_locations"}
                for control in result["control_totals"][:limit]
            ],
            "control_totals_truncated": len(result["control_totals"]) > limit,
            "next_cursor": f"{inspection_digest}:{end}" if end < len(result["items"]) else None,
        }

    def receive(
        self,
        subject_id: str,
        data: dict,
        *,
        evidence: tuple[str, ...],
        expected_revision: int,
        request_id: str,
    ):
        fact = MaterialSource.model_validate_json(canonical(data))
        inspection = self._inspection(fact.evidence_digest, fact.specification)
        if fact.evidence_digest not in evidence:
            raise NeedsInformation("evidence", "来源事实必须引用实际原件")
        entries, _, _ = _period_partition(fact, inspection, (), set(), complete=False)
        request_hash, save_source = self.engine._registration(
            False,
            fact.kind,
            subject_id,
            data,
            evidence=evidence,
            expected_revision=expected_revision,
        )

        def operation(connection):
            source = save_source(connection)
            previous = _allocation_versions(connection, self.store.registry, subject_id)
            if len(previous) > 1:
                raise KernelError("material_allocation_duplicate", "来源存在多个归属身份")
            allocation = MaterialPeriodAllocation(
                period=fact.period,
                source_id=subject_id,
                source_fact_id=source["fact_id"],
                entries=entries,
            )
            _, save_allocation = self.engine._registration(
                False,
                allocation.kind,
                _allocation_id(subject_id),
                allocation.model_dump(mode="json"),
                evidence=evidence,
                expected_revision=previous[0].revision if previous else 0,
            )
            stored = save_allocation(connection)
            return {
                **source,
                "allocation_fact_id": stored["fact_id"],
                "allocation_revision": stored["revision"],
            }

        return self.engine._write(
            request_id,
            digest(["receive_material", request_hash.hex()]),
            None,
            ("material",),
            "receive_material",
            operation,
        )

    def _prepare_allocation(self, source_id, assignments):
        if assignments is not None and (
            not isinstance(assignments, tuple)
            or any(not isinstance(item, PeriodAllocationEntry) for item in assignments)
        ):
            raise TypeError("assignments must contain typed PeriodAllocationEntry values")
        with self.store.connection(read_only=True) as connection:
            connection.execute("BEGIN")
            source = self.store.current_fact(connection, source_id)
            if source.fact.kind != MaterialSource.kind:
                raise KernelError("material_source_invalid", "归属确认必须指向资料来源")
            versions = _allocation_versions(connection, self.store.registry, source_id)
            if len(versions) > 1:
                raise KernelError("material_allocation_duplicate", "来源存在多个归属身份")
            if assignments is None and versions and versions[0].fact.source_fact_id == source.id:
                assignments = versions[0].fact.entries
            raw = connection.execute(
                "SELECT content FROM evidence WHERE digest=?",
                (bytes.fromhex(source.fact.evidence_digest),),
            ).fetchone()
            if raw is None or source.fact.evidence_digest not in source.evidence:
                raise NeedsInformation("evidence", "来源未留存实际原件")
            proofs = _allocation_proofs(connection, assignments or ())
            prior_resolutions = (
                _facts(
                    connection,
                    self.store.registry,
                    MaterialResolution.kind,
                    "material-source:" + source_id,
                )
                if assignments is None
                else ()
            )
            prior_proofs = {
                proof
                for item in prior_resolutions
                for proof in item.evidence
                if connection.execute(
                    "SELECT 1 FROM evidence WHERE digest=?", (bytes.fromhex(proof),)
                ).fetchone()
            }
            epochs = self.store.epochs(connection)
        inspection = inspect_bytes(raw[0], source.fact.specification)
        if assignments is None:
            by_location = {}
            for version in prior_resolutions:
                by_location.setdefault(version.fact.location, []).append(version)
            reused = []
            for item in inspection["items"]:
                matches = by_location.get(item["location"], ())
                if _automatic_period(source.fact, item)[0] is not None or len(matches) != 1:
                    continue
                old = matches[0]
                if old.fact.source_fact_id != source.id or old.fact.recognition_period is None:
                    continue
                candidates = sorted(set(old.evidence) & prior_proofs)
                if not candidates:
                    continue
                reused.append(
                    PeriodAllocationEntry(
                        location=item["location"],
                        recognition_period=old.fact.recognition_period,
                        basis="confirmed_period",
                        basis_evidence_digest=candidates[0],
                        basis_location=f"material_resolution:{old.id}:recognition_period",
                        basis_excerpt=old.fact.reason
                        or f"已确认处置核算所属期：{old.fact.recognition_period}",
                    )
                )
                proofs.add(candidates[0])
            assignments = tuple(reused)
        entries, periods, issues = _period_partition(
            source.fact, inspection, assignments or (), proofs, complete=False
        )
        fact = MaterialPeriodAllocation(
            period=source.fact.period,
            source_id=source_id,
            source_fact_id=source.id,
            entries=entries,
        )
        evidence = tuple(sorted({source.fact.evidence_digest, *proofs}))
        counts = {}
        for value in periods.values():
            key = str(value) if value is not None else "unknown"
            counts[key] = counts.get(key, 0) + 1
        result = {
            "source_id": source_id,
            "source_fact_id": source.id,
            "expected_revision": versions[0].revision if versions else 0,
            "epochs": epochs,
            "period_counts": counts,
            "item_count": len(periods),
            "entry_count": len(entries),
            "issues": issues[:100],
            "issue_count": len(issues),
            "issues_truncated": len(issues) > 100,
            "status": "needs_information" if issues else "ready",
        }
        result["digest"] = digest(
            [result, fact.model_dump(mode="json"), evidence, calculator_build_id()]
        ).hex()
        return result, fact, evidence

    def preview_period_allocation(
        self,
        source_id: str,
        assignments: tuple[PeriodAllocationEntry, ...] | None = None,
    ):
        return self._prepare_allocation(source_id, _typed_assignments(assignments))[0]

    def confirm_period_allocation(
        self,
        source_id: str,
        assignments: tuple[PeriodAllocationEntry, ...] | None = None,
        *,
        preview_digest: str,
        epochs: dict,
        expected_revision: int,
        request_id: str,
    ):
        assignments = _typed_assignments(assignments)
        request_hash = digest(
            [
                "material_allocation",
                source_id,
                [item.model_dump(mode="json") for item in assignments or ()],
                preview_digest,
                epochs,
                expected_revision,
            ]
        )
        cached = self.engine._cached(request_id, request_hash)
        if cached is not None:
            return cached
        preview, fact, evidence = self._prepare_allocation(source_id, assignments)
        if preview["digest"] != preview_digest or preview["expected_revision"] != expected_revision:
            raise KernelError("preview_expired", "原件或归属确认已变化，请重新预览")
        if preview["status"] != "ready":
            raise KernelError(
                "needs_information", "归属依据未通过核对", fact_issues=preview["issues"]
            )
        _, operation = self.engine._registration(
            False,
            fact.kind,
            _allocation_id(source_id),
            fact.model_dump(mode="json"),
            evidence=evidence,
            expected_revision=expected_revision,
        )
        return self.engine._write(
            request_id, request_hash, epochs, ("material",), "material_allocation", operation
        )

    def resolve(
        self,
        subject_id: str,
        data: dict,
        *,
        evidence: tuple[str, ...],
        expected_revision: int,
        request_id: str,
    ):
        return self.engine.save_fact(
            MaterialResolution.kind,
            subject_id,
            data,
            evidence=evidence,
            expected_revision=expected_revision,
            request_id=request_id,
        )

    def resolve_group(
        self,
        subject_id: str,
        data: dict,
        *,
        evidence: tuple[str, ...],
        expected_revision: int,
        request_id: str,
    ):
        request_hash, operation = self.engine._registration(
            False,
            MaterialGroupResolution.kind,
            subject_id,
            data,
            evidence=evidence,
            expected_revision=expected_revision,
        )
        request_hash = digest(["resolve_material_group", request_hash.hex()])
        replay = self.engine._cached(request_id, request_hash)
        if replay is not None:
            return replay
        fact = MaterialGroupResolution.model_validate_json(canonical(data))
        version = FactVersion("pending-group", subject_id, expected_revision + 1, fact, evidence)
        # Validate on a read snapshot; the writer only checks epochs and saves.
        with self.store.connection(read_only=True) as connection:
            connection.execute("BEGIN")
            epochs = self.store.epochs(connection)
            try:
                source = self.store.current_fact(connection, fact.source_id)
            except KernelError:
                source = None
            if source is None or source.fact.kind != MaterialSource.kind:
                raise KernelError("material_source_changed", "联合组必须引用实际资料来源")
            raw = connection.execute(
                "SELECT content FROM evidence WHERE digest=?",
                (bytes.fromhex(source.fact.evidence_digest),),
            ).fetchone()
            if raw is None:
                raise NeedsInformation("evidence", "先保全完整资料原件")
            inspection = inspect_bytes(raw[0], source.fact.specification)
            issues = _group_issues(connection, self.store.registry, version, source, inspection)
        if issues:
            raise KernelError(issues[0]["code"], issues[0]["message"], fact_issues=issues)
        return self.engine._write(
            request_id,
            request_hash,
            epochs,
            ("material",),
            "resolve_material_group",
            operation,
            checked_lanes=("accounting", "material"),
        )

    def check(self, period: str, *, limit: PageLimit = 100):
        _page_limit(limit)
        with self.store.connection(read_only=True) as connection:
            connection.execute("BEGIN")
            result = check_completeness(connection, YearMonth(period).ordinal, self.store.registry)
        return {
            **result,
            "issue_count": len(result["issues"]),
            "issues": result["issues"][:limit],
            "issues_truncated": len(result["issues"]) > limit,
            "coverage_count": len(result["coverage"]),
            "coverage": result["coverage"][:limit],
            "coverage_truncated": len(result["coverage"]) > limit,
            **{
                key + suffix: value
                for key in (
                    "source_versions",
                    "resolution_versions",
                    "group_versions",
                    "allocation_versions",
                    "fact_ids",
                    "file_summaries",
                )
                for suffix, value in (
                    ("", result[key][:limit]),
                    ("_count", len(result[key])),
                    ("_truncated", len(result[key]) > limit),
                )
            },
        }


def _page_limit(value):
    if type(value) is not int or not 1 <= value <= 500:
        raise ValueError("material page limit must be an integer between 1 and 500")
