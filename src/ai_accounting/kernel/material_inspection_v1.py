"""Frozen v1 read-only decoder for original material positions.

Historical duplicate checks require exact saved source locations to remain
readable after the current material inspection rules change.
"""

from __future__ import annotations

import csv
import io
import posixpath
import re
from bisect import bisect_left, bisect_right
from decimal import Decimal, InvalidOperation
from typing import Annotated, Literal
from xml.etree import ElementTree
from zipfile import ZipFile

from openpyxl.utils import column_index_from_string, get_column_letter, range_boundaries
from openpyxl.utils.cell import coordinate_from_string
from pydantic import (
    BaseModel,
    ConfigDict,
    Field,
    StrictBool,
    StrictInt,
    model_serializer,
    model_validator,
)

from .content_v1 import _V1ActualDate as ActualDate
from .content_v1 import _V1YearMonth as YearMonth
from .contracts import KernelError
from .history_encoding_v1 import digest


def inventory_reference(row, item_digests):
    content = {
        "inventory_id": row["id"],
        "period": str(YearMonth.from_ordinal(row["period"])),
        "category": row["category"],
        "expected": row["expected"],
        "received": row["received"],
        "no_business": bool(row["no_business"]),
        "evidence_digest": row["evidence_digest"].hex(),
        "item_digests": sorted(item_digests),
    }
    return {
        "inventory_id": content["inventory_id"],
        "period": content["period"],
        "category": content["category"],
        "content_digest": digest(content).hex(),
    }


def checked(value: int) -> int:
    if type(value) is not int or not -(2**63) <= value <= 2**63 - 1:
        raise ValueError("amount must be a signed 64-bit integer number of fen")
    return value


def sum_fen(values) -> int:
    result = 0
    for value in values:
        result = checked(result + checked(value))
    return result

Fen = Annotated[StrictInt, Field(ge=-(2**63), le=2**63 - 1)]


Identifier = Annotated[str, Field(min_length=1, max_length=200)]


class Detail(BaseModel):
    model_config = ConfigDict(extra="forbid", frozen=True, strict=True)


class Column(Detail):
    column: Annotated[str, Field(pattern=r"^[A-Z]+$")]
    role: Literal["amount", "context", "recognition_period"]
    label: str = ""
    funds_direction: Literal["signed_net", "inflow", "outflow"] | None = Field(
        default=None,
        description=(
            "原金额列的核算方向：signed_net为原正负净额，inflow/outflow为非负收入/支出列。"
            "仅金额列可声明，须依据原表头或明确来源，不能从label猜测。原金额及控制合计不改写；"
            "省略时兼容原有正负净额语义，旧来源序列化不新增字段。"
        ),
    )

    @model_validator(mode="after")
    def amount_direction_only(self):
        if self.funds_direction is not None and self.role != "amount":
            raise ValueError("funds direction may only describe an original amount column")
        return self

    @model_serializer(mode="wrap")
    def compatible_dump(self, handler):
        result = handler(self)
        if self.funds_direction is None:
            result.pop("funds_direction", None)
        return result


class Passage(Detail):
    location: Identifier
    page: Annotated[int, Field(ge=1)]
    excerpt: Annotated[str, Field(min_length=1)]
    amount_fen: Fen | None = None
    recognition_period: YearMonth | None = None


class MaterialRange(Detail):
    sheet: Identifier
    column: Annotated[str, Field(pattern=r"^[A-Z]+$")]
    first_row: Annotated[int, Field(ge=1)]
    last_row: Annotated[int, Field(ge=1)]

    @model_validator(mode="after")
    def ordered(self):
        if self.last_row < self.first_row:
            raise ValueError("material range must be increasing")
        return self


class ControlTotal(Detail):
    location: Identifier
    scope: Literal["specified", "all_detail_column"] = "specified"
    ranges: tuple[MaterialRange, ...] = ()
    locations: tuple[Identifier, ...] = ()

    @model_validator(mode="after")
    def explicit_members(self):
        if self.scope == "specified" and not (self.ranges or self.locations):
            raise ValueError("a grouped control needs explicit original members")
        if self.scope == "all_detail_column" and (self.ranges or self.locations):
            raise ValueError("whole-column control cannot also name partial members")
        return self


class Specification(Detail):
    format: Literal["csv", "xls", "xlsx", "pdf", "image", "text"]
    columns: tuple[Column, ...] = ()
    sheet_columns: dict[str, tuple[Column, ...]] = Field(default_factory=dict)
    header_rows: dict[str, int] = Field(default_factory=dict)
    total_rows: dict[str, tuple[int, ...]] = Field(default_factory=dict)
    passages: tuple[Passage, ...] = ()
    all_pages_reviewed: StrictBool | None = None
    controls: tuple[ControlTotal, ...] = ()

    @model_serializer(mode="wrap")
    def compatible_dump(self, handler):
        result = handler(self)
        if not self.controls:
            result.pop("controls", None)
        return result

    @model_validator(mode="after")
    def unique_columns(self):
        for columns in (self.columns, *self.sheet_columns.values()):
            if len({item.column for item in columns}) != len(columns):
                raise ValueError("duplicate material column")
            if sum(item.role == "recognition_period" for item in columns) > 1:
                raise ValueError("one declared recognition-period column per sheet")
        if any(type(row) is not int or row < 0 for row in self.header_rows.values()):
            raise ValueError("header rows must be nonnegative integers")
        if any(row <= 0 for rows in self.total_rows.values() for row in rows):
            raise ValueError("total rows must be positive")
        if len({item.location for item in self.passages}) != len(self.passages):
            raise ValueError("duplicate original passage location")
        if len({item.location for item in self.controls}) != len(self.controls):
            raise ValueError("duplicate original control location")
        if self.controls and self.format not in {"csv", "xls", "xlsx"}:
            raise ValueError("group controls require original spreadsheet locations")
        return self


def _issue(code, message, *, location=None, **details):
    return {"field": "materials", "code": code, "message": message, "location": location, **details}


def _amount(value):
    text = str(value).strip()
    for char in (",", "¥", "￥", "元", " "):
        text = text.replace(char, "")
    try:
        number = Decimal(text)
        if not number.is_finite():
            raise ValueError("amount needs exact integer fen")
        parts = number.as_tuple()
        digits, scale = parts.digits, parts.exponent + 2
        if not any(digits):
            return 0
        if len(digits) + scale > 19:
            raise ValueError("amount exceeds signed 64-bit fen")
        if scale < 0:
            removed = -scale
            if removed >= len(digits) or any(digits[-removed:]):
                raise ValueError("amount needs exact integer fen")
            digits, scale = digits[:-removed], 0
        # Decimal arithmetic uses the ambient context. Exact digit shifting
        # avoids rounding and rejects enormous exponents before integer allocation.
        amount = int("".join(str(digit) for digit in digits)) * 10**scale
        return checked(-amount if parts.sign else amount)
    except (InvalidOperation, ValueError) as exc:
        raise ValueError("amount is missing, unreadable or not integer fen") from exc


def inspect_bytes(raw: bytes, specification: Specification) -> dict:
    """Normalize corrupt/unreadable original formats into a structured issue."""
    from zipfile import BadZipFile

    from openpyxl.utils.exceptions import InvalidFileException
    from pypdf.errors import PyPdfError
    from xlrd.biffh import XLRDError

    try:
        return _inspect_bytes(raw, specification)
    except KernelError:
        raise
    except (
        ValueError,
        IndexError,
        OverflowError,
        csv.Error,
        OSError,
        KeyError,
        InvalidFileException,
        PyPdfError,
        XLRDError,
        BadZipFile,
        ElementTree.ParseError,
    ) as exc:
        raise KernelError(
            "material_unreadable", "原件格式损坏或当前映射无法读取", detail=str(exc)
        ) from exc


MAX_MATERIAL_BYTES = 20 * 1024 * 1024


MAX_DATA_ROWS = 100_000


MAX_COLUMNS = 64


MAX_NONEMPTY_CELLS = 2_000_000


MAX_SHEETS = 32


MAX_HEADER_ROWS = 100


MAX_CONTROL_ROWS = 1000


MAX_EXPANDED_BYTES = 200 * 1024 * 1024


def _too_large(message):
    raise KernelError("material_too_large", message)


class _TableBudget:
    """Bound the original document before retaining its normalized output."""

    def __init__(self, spec):
        self.spec = spec
        self.data_rows = 0
        self.cells = 0
        self.sheets = set()
        self.control_locations = {item.location for item in spec.controls}
        self._amount_columns = {}
        if any(row > MAX_HEADER_ROWS for row in spec.header_rows.values()):
            _too_large("每表最多支持100行表头")
        if (
            sum(len(rows) for rows in spec.total_rows.values()) + len(spec.controls)
            > MAX_CONTROL_ROWS
        ):
            _too_large("控制合计行最多支持1000行")
        for columns in (spec.columns, *spec.sheet_columns.values()):
            if any(column_index_from_string(item.column) > MAX_COLUMNS for item in columns):
                _too_large("表格最多支持64列")

    def dimensions(self, sheet, row, column):
        self.sheets.add(sheet)
        if len(self.sheets) > MAX_SHEETS:
            _too_large("单个文件最多支持32张工作表")
        max_row = MAX_DATA_ROWS + self.spec.header_rows.get(sheet, 1) + MAX_CONTROL_ROWS
        if row > max_row or column > MAX_COLUMNS:
            _too_large("表格行号或列宽超出安全范围（最多64列）")

    def accept(self, sheet, row, cells):
        self.dimensions(sheet, row, max(cells, default=0))
        self.cells += len(cells)
        if self.cells > MAX_NONEMPTY_CELLS:
            _too_large("文件最多支持200万个非空单元格")
        if sheet not in self._amount_columns:
            self._amount_columns[sheet] = tuple(
                item.column
                for item in self.spec.sheet_columns.get(sheet, self.spec.columns)
                if item.role == "amount"
            )
        amount_columns = self._amount_columns[sheet]
        grouped_control_row = bool(amount_columns) and all(
            f"{sheet}!{column}{row}" in self.control_locations for column in amount_columns
        )
        if (
            cells
            and row > self.spec.header_rows.get(sheet, 1)
            and row not in self.spec.total_rows.get(sheet, ())
            and not grouped_control_row
        ):
            self.data_rows += 1
            if self.data_rows > MAX_DATA_ROWS:
                _too_large("单个文件最多支持十万条数据行，表头和控制行另计")


def _csv_rows(raw, budget):
    try:
        raw.decode("utf-8-sig")
        encoding = "utf-8-sig"
    except UnicodeDecodeError:
        encoding = "gb18030"
    with io.TextIOWrapper(io.BytesIO(raw), encoding=encoding, newline="") as source:
        for number, row in enumerate(csv.reader(source), 1):
            budget.dimensions("CSV", number, len(row))
            cells = {
                index: (value, False, False) for index, value in enumerate(row, 1) if value.strip()
            }
            budget.accept("CSV", number, cells)
            if cells:
                yield "CSV", number, cells


def _xls_rows(raw, budget):
    import xlrd

    book = xlrd.open_workbook(file_contents=raw, formatting_info=True, on_demand=True)
    try:
        for sheet_index in range(book.nsheets):
            sheet = book.sheet_by_index(sheet_index)
            budget.dimensions(sheet.name, sheet.nrows, sheet.ncols)
            for row in range(sheet.nrows):
                cells = {}
                row_info = sheet.rowinfo_map.get(row)
                for column in range(sheet.ncols):
                    cell = sheet.cell(row, column)
                    if cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                        continue
                    value = str(cell.value)
                    if cell.ctype == xlrd.XL_CELL_DATE:
                        value = (
                            xlrd.xldate_as_datetime(cell.value, book.datemode).date().isoformat()
                        )
                    elif cell.ctype == xlrd.XL_CELL_ERROR:
                        value = None
                    column_info = sheet.colinfo_map.get(column)
                    cells[column + 1] = (
                        value,
                        cell.ctype == xlrd.XL_CELL_ERROR,
                        bool(
                            sheet.visibility
                            or (row_info and row_info.hidden)
                            or (column_info and column_info.hidden)
                        ),
                    )
                budget.accept(sheet.name, row + 1, cells)
                if cells:
                    yield sheet.name, row + 1, cells
            book.unload_sheet(sheet_index)
    finally:
        book.release_resources()


def _xlsx_display_text(node):
    """Spreadsheet text consists of direct text and rich runs, never phonetic guides."""
    parts = []
    for child in node:
        if child.tag.endswith("}t"):
            parts.append(child.text or "")
        elif child.tag.endswith("}r"):
            parts.append(child.findtext("{*}t") or "")
    return "".join(parts)


def _xlsx_shared_strings(archive):
    values = []
    if "xl/sharedStrings.xml" not in archive.namelist():
        return values
    with archive.open("xl/sharedStrings.xml") as source:
        events = ElementTree.iterparse(source, events=("start", "end"))
        _, root = next(events)
        for event, node in events:
            if event == "end" and node.tag.endswith("}si"):
                values.append(_xlsx_display_text(node))
                if len(values) > MAX_NONEMPTY_CELLS:
                    _too_large("共享文本数量超出单元格安全范围")
                node.clear()
                root.remove(node)
    return values


def _xlsx_dates(archive):
    if "xl/styles.xml" not in archive.namelist():
        return set()
    from openpyxl.styles.numbers import BUILTIN_FORMATS, is_date_format

    styles = ElementTree.fromstring(archive.read("xl/styles.xml"))
    formats = dict(BUILTIN_FORMATS)
    for item in styles.findall("{*}numFmts/{*}numFmt"):
        formats[int(item.attrib["numFmtId"])] = item.attrib["formatCode"]
    return {
        index
        for index, item in enumerate(styles.findall("{*}cellXfs/{*}xf"))
        if is_date_format(formats.get(int(item.attrib.get("numFmtId", "0")), "General"))
    }


def _xlsx_rows(raw, budget):
    """Stream exact XML values once, including hidden cells and formula caches."""
    from openpyxl.utils.datetime import CALENDAR_MAC_1904, CALENDAR_WINDOWS_1900, from_excel

    with ZipFile(io.BytesIO(raw)) as archive:
        if sum(item.file_size for item in archive.infolist()) > MAX_EXPANDED_BYTES:
            _too_large("Excel解压内容最多支持200 MiB")
        if len(archive.infolist()) > 10_000:
            _too_large("Excel压缩成员数量超出安全范围")
        book = ElementTree.fromstring(archive.read("xl/workbook.xml"))
        rels = ElementTree.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
        targets = {rel.attrib["Id"]: rel.attrib["Target"] for rel in rels}
        strings, date_styles = _xlsx_shared_strings(archive), _xlsx_dates(archive)
        props = book.find("{*}workbookPr")
        epoch = (
            CALENDAR_MAC_1904
            if props is not None and props.get("date1904") in {"1", "true"}
            else CALENDAR_WINDOWS_1900
        )
        for sheet in book.findall(".//{*}sheet"):
            name = sheet.attrib["name"]
            budget.dimensions(name, 0, 0)
            rel = next(value for key, value in sheet.attrib.items() if key.endswith("}id"))
            target = targets[rel]
            path = (
                target.lstrip("/")
                if target.startswith("/")
                else posixpath.normpath(posixpath.join("xl", target))
            )
            hidden_columns = set()
            last_row = 0
            sheet_hidden = sheet.get("state", "visible") != "visible"
            with archive.open(path) as source:
                events = ElementTree.iterparse(source, events=("start", "end"))
                stack = []
                main = None
                for event, node in events:
                    if event == "start":
                        if not stack:
                            if node.tag not in {
                                "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}worksheet",
                                "{http://purl.oclc.org/ooxml/spreadsheetml/main}worksheet",
                            }:
                                raise ValueError("unknown Excel worksheet namespace")
                            main = node.tag.removesuffix("worksheet")
                        stack.append(node)
                        continue
                    # Drawing anchors also contain row/col elements. Only the
                    # worksheet's direct dimension/cols/sheetData define cells.
                    parent = stack[-2].tag if len(stack) > 1 else None
                    if node.tag == main + "dimension" and len(stack) == 2:
                        # Excel may retain a full-sheet used range after formatting.
                        # Streaming actual cells, including formulas, sets the budget.
                        range_boundaries(node.attrib["ref"])
                    elif node.tag == main + "col" and len(stack) == 3 and parent == main + "cols":
                        first = int(node.attrib["min"])
                        last = int(node.attrib["max"])
                        if not 1 <= first <= last <= 16_384:
                            raise ValueError("invalid Excel column formatting range")
                        if node.get("hidden") in {"1", "true"}:
                            hidden_columns.update(range(first, min(last, MAX_COLUMNS) + 1))
                    elif (
                        node.tag == main + "row"
                        and len(stack) == 3
                        and parent == main + "sheetData"
                    ):
                        number = int(node.attrib["r"])
                        if number <= last_row:
                            raise ValueError("Excel row locations must be unique and increasing")
                        last_row = number
                        row_hidden = node.get("hidden") in {"1", "true"}
                        cells = {}
                        seen_columns = set()
                        for cell in node.findall(main + "c"):
                            column_text, cell_row = coordinate_from_string(cell.attrib["r"])
                            column = column_index_from_string(column_text)
                            if cell_row != number or column in seen_columns:
                                raise ValueError("inconsistent or duplicate Excel cell location")
                            seen_columns.add(column)
                            kind = cell.get("t", "n")
                            value = cell.findtext(main + "v")
                            formula = cell.find(main + "f") is not None
                            if kind == "inlineStr":
                                inline = cell.find(main + "is")
                                value = _xlsx_display_text(inline) if inline is not None else ""
                            elif kind == "s" and value is not None:
                                index = int(value)
                                if not 0 <= index < len(strings):
                                    raise ValueError("invalid Excel shared-string reference")
                                value = strings[index]
                            elif kind == "b" and value is not None:
                                value = value == "1"
                            elif kind == "e":
                                value = None
                            elif kind == "n" and value and int(cell.get("s", "0")) in date_styles:
                                # Only date serials use this conversion.
                                # Money stays in its original decimal XML text.
                                date_value = from_excel(float(value), epoch=epoch)
                                value = (
                                    date_value.date().isoformat()
                                    if hasattr(date_value, "date")
                                    else date_value.isoformat()
                                )
                            missing = (formula and (value is None or value == "")) or kind == "e"
                            if value is None or not str(value).strip():
                                if not missing:
                                    continue
                            cells[column] = (
                                value,
                                missing,
                                sheet_hidden or row_hidden or column in hidden_columns,
                            )
                        if cells:
                            budget.accept(name, number, cells)
                            yield name, number, cells
                        node.clear()
                        if len(stack) > 1:
                            stack[-2].remove(node)
                    stack.pop()


def _inspect_bytes(raw: bytes, specification: Specification) -> dict:
    """Normalize one row at a time, retaining only the public inspection result."""
    spec = specification
    if len(raw) > MAX_MATERIAL_BYTES:
        _too_large("资料原件最多支持20 MiB")
    if raw.startswith(b"%PDF-") and spec.format != "pdf":
        raise KernelError("material_format_mismatch", "PDF原件必须按全部实际页核对")
    if raw.startswith(b"PK\x03\x04") and spec.format != "xlsx":
        raise KernelError("material_format_mismatch", "表格压缩原件必须按实际单元格核对")
    if raw.startswith(bytes.fromhex("d0cf11e0a1b11ae1")) and spec.format != "xls":
        raise KernelError("material_format_mismatch", "旧式Excel原件必须按实际工作表核对")
    if spec.format not in {"csv", "xls", "xlsx"}:
        return _inspect_passages(raw, spec)
    budget = _TableBudget(spec)
    rows = {"csv": _csv_rows, "xls": _xls_rows, "xlsx": _xlsx_rows}[spec.format](raw, budget)
    mappings, issues, items, coverage, totals, sums = {}, [], [], [], [], {}
    known_periods = {}
    grouped_locations = {item.location for item in spec.controls}
    for sheet, row, cells in rows:
        if sheet not in mappings:
            mappings[sheet] = {
                column_index_from_string(item.column): item
                for item in spec.sheet_columns.get(sheet, spec.columns)
            }
        mapping = mappings[sheet]
        if row <= spec.header_rows.get(sheet, 1):
            coverage.extend(
                {
                    "location": f"{sheet}!{get_column_letter(column)}{row}",
                    "value": str(value),
                    "hidden": hidden,
                }
                for column, (value, _, hidden) in cells.items()
            )
            continue
        for column, rule in mapping.items():
            if rule.role == "amount" and column not in cells:
                budget.cells += 1
                if budget.cells > MAX_NONEMPTY_CELLS:
                    _too_large("原单元格与待补金额字段合计最多支持200万个")
                cells[column] = None, False, False
        period, context = None, []
        for column, (value, _, _) in cells.items():
            rule = mapping.get(column)
            if rule is not None and rule.role == "context" and value is not None:
                context.append(str(value))
            if rule is not None and rule.role == "recognition_period" and value is not None:
                try:
                    text = str(value)
                    if text not in known_periods:
                        known_periods[text] = (
                            ActualDate(text).period if len(text) == 10 else YearMonth(text)
                        )
                    period = known_periods[text]
                except ValueError:
                    issues.append(
                        _issue(
                            "material_period_unreadable",
                            "原行的核算所属期不能读取",
                            location=f"{sheet}!{get_column_letter(column)}{row}",
                        )
                    )
        excerpt = " · ".join(context)
        for column, (value, formula_missing, hidden) in cells.items():
            letter = get_column_letter(column)
            location = f"{sheet}!{letter}{row}"
            coverage.append({"location": location, "value": str(value), "hidden": hidden})
            rule = mapping.get(column)
            if rule is None:
                issues.append(
                    _issue("material_column_unmapped", "非空列尚未说明业务含义", location=location)
                )
                continue
            if formula_missing:
                issues.append(
                    _issue(
                        "material_formula_result_missing",
                        "公式缺少已计算结果，不能跳过",
                        location=location,
                    )
                )
            if rule.role != "amount":
                continue
            amount = None
            try:
                amount = _amount(value)
            except ValueError:
                issues.append(
                    _issue(
                        "material_amount_missing",
                        "原行金额缺失或无法读取为整数分",
                        location=location,
                    )
                )
            if amount is not None and amount < 0 and rule.funds_direction in {"inflow", "outflow"}:
                issues.append(
                    _issue(
                        "material_unsigned_amount_negative",
                        "已声明非负收入/支出列出现负数，须复核原件方向或明确冲销语义，不能取绝对值",
                        location=location,
                    )
                )
            if row in spec.total_rows.get(sheet, ()) or location in grouped_locations:
                totals.append(
                    {"location": location, "sheet": sheet, "column": letter, "expected_fen": amount}
                )
            else:
                items.append(
                    {
                        "location": location,
                        "amount_fen": amount,
                        "recognition_period": period,
                        "excerpt": excerpt,
                        "sheet": sheet,
                        "column": letter,
                        "hidden": hidden,
                    }
                    | ({"funds_direction": rule.funds_direction} if rule.funds_direction else {})
                )
                key = sheet, letter
                old = sums.get(key, 0)
                sums[key] = (
                    checked(old + amount) if old is not None and amount is not None else None
                )
    _bind_controls(items, totals, spec, issues)
    if not items:
        issues.append(
            _issue("material_no_business_columns", "未识别业务金额列，不能以空集合确认处理完成")
        )
    return {"items": items, "issues": issues, "coverage": coverage, "control_totals": totals}


def _bind_controls(items, totals, spec, issues):
    """Controls reference original leaf amounts, never other controls."""
    if not totals and not spec.controls:
        return
    by_location = {item["location"]: item for item in items}
    by_column = {}
    for item in items:
        by_column.setdefault((item["sheet"], item["column"]), []).append(item)
    range_index = {}
    for key, values in by_column.items():
        ordered = sorted(values, key=lambda item: _row_number(item["location"]))
        range_index[key] = ([_row_number(item["location"]) for item in ordered], ordered)
    definitions = {control.location: control for control in spec.controls}
    found = {control["location"] for control in totals}
    for location in definitions.keys() - found:
        issues.append(
            _issue(
                "material_control_position_invalid",
                "控制金额位置必须是原件中明确映射的金额单元格",
                location=location,
            )
        )
    memberships = 0
    for control in totals:
        definition = definitions.get(control["location"])
        members = []
        if definition is None or definition.scope == "all_detail_column":
            members = list(by_column.get((control["sheet"], control["column"]), ()))
        else:
            if memberships + len(definition.locations) > MAX_NONEMPTY_CELLS:
                _too_large("控制成员总数超过200万，需减少重复控制范围")
            for location in definition.locations:
                if location not in by_location:
                    issues.append(
                        _issue(
                            "material_control_member_invalid",
                            "控制成员必须是实际业务金额，不能是控制值或虚构位置",
                            location=control["location"],
                            member_location=location,
                        )
                    )
                else:
                    members.append(by_location[location])
            for selection in definition.ranges:
                rows, values = range_index.get((selection.sheet, selection.column), ((), ()))
                selected = values[
                    bisect_left(rows, selection.first_row) : bisect_right(rows, selection.last_row)
                ]
                if memberships + len(members) + len(selected) > MAX_NONEMPTY_CELLS:
                    _too_large("控制成员总数超过200万，需减少重复控制范围")
                if not selected:
                    issues.append(
                        _issue(
                            "material_control_empty_range",
                            "控制范围没有业务金额",
                            location=control["location"],
                        )
                    )
                members.extend(selected)
        locations = [item["location"] for item in members]
        if len(set(locations)) != len(locations):
            issues.append(
                _issue(
                    "material_control_duplicate_member",
                    "同一控制不能重复计入明细",
                    location=control["location"],
                )
            )
        memberships += len(locations)
        if memberships > MAX_NONEMPTY_CELLS:
            _too_large("控制成员总数超过200万，需减少重复控制范围")
        control["member_locations"] = locations
        control["member_count"] = len(locations)
        actual = (
            sum_fen(item["amount_fen"] for item in members)
            if all(item["amount_fen"] is not None for item in members)
            else None
        )
        control["actual_fen"] = actual
        if not members or control["expected_fen"] is None or actual != control["expected_fen"]:
            issues.append(
                _issue(
                    "material_total_mismatch",
                    "原资料控制合计与明确明细范围不一致",
                    location=control["location"],
                    expected_fen=control["expected_fen"],
                    actual_fen=actual,
                )
            )


def _row_number(location):
    return int(re.search(r"[0-9]+$", location).group())


def _inspect_passages(raw, spec):
    pages = 1
    if spec.format == "pdf":
        from pypdf import PdfReader

        pages = len(PdfReader(io.BytesIO(raw)).pages)
    if spec.format == "image" and not (
        raw.startswith(b"\x89PNG\r\n\x1a\n") or raw.startswith(b"\xff\xd8\xff")
    ):
        raise KernelError(
            "material_image_format", "图像核对目前接收PNG或JPEG原件；多页图像须转换为完整PDF后核对"
        )
    issues = []
    represented = {item.page for item in spec.passages}
    if spec.all_pages_reviewed is not True or represented != set(range(1, pages + 1)):
        issues.append(
            _issue(
                "material_pages_unreviewed",
                "须明确阅读全部页或图片，并逐页登记位置与摘录",
                pages=pages,
            )
        )
    items = []
    for item in spec.passages:
        if not item.excerpt.strip():
            issues.append(
                _issue("material_passage_empty", "原文摘录不能空白", location=item.location)
            )
        if spec.format == "text":
            try:
                text = raw.decode("utf-8-sig")
            except UnicodeDecodeError:
                text = raw.decode("gb18030")
            if item.excerpt not in text:
                issues.append(
                    _issue("material_excerpt_missing", "摘录不在原文中", location=item.location)
                )
        items.append(
            {
                "location": item.location,
                "page": item.page,
                "excerpt": item.excerpt,
                "amount_fen": item.amount_fen,
                "recognition_period": item.recognition_period,
            }
        )
    return {
        "items": items,
        "issues": issues,
        "coverage": [item.model_dump(mode="json") for item in spec.passages],
        "control_totals": [],
        "semantic_review": "人工或AI逐页阅读确认；程序只检查位置覆盖，不证明自由文本被完整理解",
    }
