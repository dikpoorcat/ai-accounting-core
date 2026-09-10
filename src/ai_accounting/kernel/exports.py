"""Frozen payment instructions; generating a file never records an actual payment.

The requested period is the management payment month. An explicitly saved
payment_period groups a source there; otherwise its accounting month is used.
An omitted source_ids means every eligible unpaid salary, bonus, labor and
employee reimbursement in that scope. Explicit source_ids request a subset.
Both modes use the same material and unresolved-accounting checks as month close.
"""

from __future__ import annotations

import hashlib
import json
import os
import re
import tempfile
import uuid
from io import BytesIO
from pathlib import Path

from ai_accounting.mybank_export import render_import, validate_template, yuan

from .backup import _worker_lock
from .contracts import KernelError, NeedsInformation
from .periods import Periods
from .types import YearMonth, canonical, digest, sum_fen

WORKBOOK_NAME = "银行批量代发.xlsx"
MANIFEST_NAME = "代发核对.json"
EXPORT_KINDS = ("payroll", "annual_bonus", "labor", "expense", "employee_advance")
DEFAULT_CATEGORIES = {
    "payroll": "工资",
    "annual_bonus": "奖金",
    "labor": "劳务",
    "expense": "报销",
    "employee_advance": "报销",
}


def _evidence_digest(value: str) -> bytes:
    if not isinstance(value, str) or not re.fullmatch(r"[0-9a-f]{64}", value):
        raise ValueError("evidence digest must be a lowercase SHA-256 hexadecimal value")
    return bytes.fromhex(value)


def _safe_text(value: str, field: str, *, maximum: int) -> str:
    if (
        not isinstance(value, str)
        or not value.strip()
        or value != value.strip()
        or len(value) > maximum
        or value.startswith(("=", "+", "-", "@"))
        or any(ord(character) < 32 for character in value)
    ):
        raise ValueError(f"invalid {field}")
    return value


class Exports:
    def __init__(self, engine):
        self.engine, self.store = engine, engine.store

    def save_payee(
        self,
        party_id: str,
        *,
        name: str,
        account: str,
        evidence_digest: str,
        expected_revision: int,
        request_id: str,
    ):
        _safe_text(party_id, "party_id", maximum=200)
        _safe_text(name, "payee name", maximum=100)
        _safe_text(account, "payee account", maximum=64)
        if not account.isascii() or not account.isdecimal():
            raise ValueError("payee account must contain ASCII digits and preserve leading zeroes")
        if type(expected_revision) is not int or expected_revision < 0:
            raise ValueError("invalid expected payee revision")
        evidence = _evidence_digest(evidence_digest)
        request_hash = digest(
            ["payee", party_id, name, account, evidence_digest, expected_revision]
        )

        def operation(connection):
            if not connection.execute(
                "SELECT 1 FROM evidence WHERE digest=?", (evidence,)
            ).fetchone():
                raise NeedsInformation("payee.evidence", "需要已留存的收款人账户依据")
            revision = connection.execute(
                "SELECT coalesce(max(revision),0) FROM payee_revision WHERE party_id=?",
                (party_id,),
            ).fetchone()[0]
            if revision != expected_revision:
                raise KernelError("payee_version_conflict", "收款资料已经变化")
            identifier = uuid.uuid4().hex
            connection.execute(
                "INSERT INTO payee_revision VALUES(?,?,?,?,?,?)",
                (
                    identifier,
                    party_id,
                    revision + 1,
                    name,
                    account,
                    evidence,
                ),
            )
            return {
                "status": "saved",
                "payee_revision_id": identifier,
                "party_id": party_id,
                "revision": revision + 1,
            }

        return self.engine._write(
            request_id, request_hash, None, ("management",), "payee", operation
        )

    def preview(
        self, period: str, *, template_evidence_digest: str, source_ids: list[str] | None = None
    ):
        month = YearMonth(period)
        if source_ids is not None and not isinstance(source_ids, list):
            raise ValueError("source_ids must be a list of stable source identities")
        requested = None if source_ids is None else sorted(set(source_ids))
        if requested is not None and (
            not requested or any(not isinstance(item, str) or not item for item in requested)
        ):
            raise ValueError("source_ids must contain at least one stable source identity")
        template_digest = _evidence_digest(template_evidence_digest)
        with self.store.connection(read_only=True) as connection:
            connection.execute("BEGIN")
            epochs = self.store.epochs(connection)
            template = connection.execute(
                "SELECT content FROM evidence WHERE digest=?",
                (template_digest,),
            ).fetchone()
            if template is None:
                raise NeedsInformation("template_evidence_digest", "需要留存银行四列原始模板")
            template_content = template[0]
            calculations = connection.execute(
                "SELECT c.*,m.id management_id,m.payment_period,m.payment_category "
                "FROM calculation_current a JOIN calculation c ON c.id=a.calculation_id "
                "LEFT JOIN management_revision m ON m.id=(SELECT n.id FROM management_revision n "
                "WHERE n.subject_id=c.subject_id ORDER BY n.revision DESC LIMIT 1) "
                "WHERE c.kind IN ('payroll','annual_bonus','labor','expense','employee_advance') "
                "AND coalesce(m.payment_period,c.period)=? "
                "AND (? IS NULL OR c.subject_id IN (SELECT value FROM json_each(?))) "
                "ORDER BY c.subject_id",
                (
                    month.ordinal,
                    None if requested is None else canonical(requested),
                    canonical(requested or []),
                ),
            ).fetchall()
            outcomes = {row["id"]: json.loads(row["outcome"]) for row in calculations}
            obligations = [
                item
                for outcome in outcomes.values()
                for item in outcome["values"].get("obligations", ())
            ]
            party_ids = sorted(
                {item["counterparty_id"] for item in obligations if item.get("counterparty_id")}
            )
            payees = {
                row["party_id"]: dict(row)
                for row in connection.execute(
                    "SELECT p.* FROM payee_revision p "
                    "WHERE p.party_id IN (SELECT value FROM json_each(?)) "
                    "AND p.revision=(SELECT max(n.revision) "
                    "FROM payee_revision n WHERE n.party_id=p.party_id)",
                    (canonical(party_ids),),
                )
            }
            balances = {
                (row["category"], row["balance_key"]): row["amount"]
                for row in connection.execute(
                    "SELECT * FROM balance WHERE category='payable' "
                    "AND balance_key IN (SELECT value FROM json_each(?))",
                    (canonical(sorted({item["key"] for item in obligations})),),
                )
            }
            sources, periods, eligible = [], {str(month)}, set()
            for calculation in calculations:
                subject = calculation["subject_id"]
                outcome = outcomes[calculation["id"]]
                if (
                    calculation["kind"] == "expense"
                    and outcome["values"].get("creditor_kind") != "employee"
                ):
                    continue
                if (
                    calculation["kind"] == "employee_advance"
                    and outcome["values"].get("payer_kind") != "employee"
                ):
                    continue
                name = (
                    "primary" if calculation["kind"] in {"expense", "employee_advance"} else "net"
                )
                obligations = [
                    item
                    for item in outcome["values"].get("obligations", ())
                    if item["name"] == name and item["normal"] == "credit"
                ]
                if len(obligations) != 1:
                    raise KernelError("invalid_export_obligation", "代发来源没有唯一明确的应付款")
                eligible.add(subject)
                item = obligations[0]
                key = (item["category"], item["key"])
                # The projection canonically omits settled zero balances.
                remaining = balances.get(key, 0)
                if remaining < 0 or remaining > item["amount_fen"]:
                    raise KernelError("invalid_payable_balance", "应付款余额与已确认来源不一致")
                source_period = str(YearMonth.from_ordinal(calculation["period"]))
                periods.add(source_period)
                if remaining == 0:
                    continue
                party_id = item.get("counterparty_id")
                if not party_id:
                    raise NeedsInformation(
                        "counterparty_id", "代发来源缺少明确收款人", sources=(subject,)
                    )
                payee = payees.get(party_id)
                if payee is None:
                    raise NeedsInformation(
                        "payee", "需要已确认的收款姓名和账号", sources=(party_id,)
                    )
                category = (
                    calculation["payment_category"] or DEFAULT_CATEGORIES[calculation["kind"]]
                )
                _safe_text(category, "payment category", maximum=32)
                if len(f"{period} {category}") > 40:
                    raise KernelError("payment_memo_too_long", "代发月份和用途合计不能超过40字")
                sources.append(
                    {
                        "subject_id": subject,
                        "kind": calculation["kind"],
                        "source_period": source_period,
                        "calculation_id": calculation["id"],
                        "obligation": item["key"],
                        "amount_fen": remaining,
                        "party_id": party_id,
                        "payee_revision_id": payee["id"],
                        "name": payee["name"],
                        "account": payee["account"],
                        "category": category,
                        "management_revision_id": calculation["management_id"],
                    }
                )
            if requested is not None and set(requested) != eligible:
                raise KernelError(
                    "invalid_export_sources",
                    "指定来源不属于本次可代发范围",
                    source_ids=sorted(set(requested) - eligible),
                )
            inventory_versions = {}
            for source_period in sorted(periods):
                inventories, issues, unpublished = Periods.completeness(
                    connection,
                    YearMonth(source_period).ordinal,
                    self.store.registry,
                )
                issues = list(issues)
                issues.extend(
                    {"field": row["id"], "message": "业务事实尚未正式核算"}
                    for row in unpublished
                    if row["kind"] in self.store.registry.evaluators
                )
                pending = connection.execute(
                    "SELECT p.subject_id FROM pending p CROSS JOIN fact_current c "
                    "CROSS JOIN fact_revision f WHERE c.subject_id=p.subject_id "
                    "AND f.id=c.fact_id AND f.period<=? LIMIT 1",
                    (YearMonth(source_period).ordinal,),
                ).fetchone()
                if pending:
                    issues.append({"field": pending[0], "message": "存在尚未完成的会计更正"})
                if issues:
                    raise KernelError(
                        "materials_incomplete",
                        "资料或会计处理尚未完整，不能生成代发",
                        period=source_period,
                        fact_issues=issues,
                    )
                inventory_versions[source_period] = {
                    key: row["id"] for key, row in inventories.items()
                }
            connection.commit()
        validate_template(template_content)
        if not sources:
            raise KernelError("no_payables", "该范围没有可代发的未付余额")
        grouped = {}
        for source in sources:
            key = (source["party_id"], source["category"])
            if key not in grouped:
                grouped[key] = {
                    name: source[name]
                    for name in (
                        "party_id",
                        "payee_revision_id",
                        "name",
                        "account",
                        "category",
                    )
                } | {"amount_fen": 0, "sources": []}
            grouped[key]["amount_fen"] = sum_fen((grouped[key]["amount_fen"], source["amount_fen"]))
            grouped[key]["sources"].append(source)
        rows = [grouped[key] for key in sorted(grouped)]
        if len(rows) > 2000:
            raise KernelError("payment_export_too_large", "银行每个文件最多支持2000条代发记录")
        result = {
            "status": "preview",
            "company_id": self.store.company_id,
            "database_id": self.store.database_id,
            "period": str(month),
            "scope": "complete" if requested is None else "selected",
            "source_ids": requested,
            "template_evidence_digest": template_evidence_digest,
            "epochs": epochs,
            "inventories": inventory_versions,
            "rows": rows,
            "total_fen": sum_fen(row["amount_fen"] for row in rows),
        }
        result["digest"] = digest(result).hex()
        return result

    def confirm(
        self,
        period: str,
        *,
        template_evidence_digest: str,
        preview_digest: str,
        epochs: dict,
        output_directory: str,
        request_id: str,
        source_ids: list[str] | None = None,
    ):
        directory = str(Path(output_directory).resolve())
        requested = None if source_ids is None else sorted(set(source_ids))
        request_hash = digest(
            [
                "payment_export",
                period,
                template_evidence_digest,
                preview_digest,
                epochs,
                directory,
                requested,
            ]
        )
        cached = self.engine._cached(request_id, request_hash)
        if cached is not None:
            return cached
        preview = self.preview(
            period, template_evidence_digest=template_evidence_digest, source_ids=requested
        )
        if preview["digest"] != preview_digest or any(
            preview["epochs"].get(lane) != epochs.get(lane)
            for lane in ("accounting", "material", "management")
        ):
            raise KernelError("preview_expired", "代发金额、资料或收款信息已变化，请重新预览")
        payload = {
            "plan": preview,
            "output_directory": directory,
            "template_evidence_digest": template_evidence_digest,
        }

        def operation(connection):
            job_id = uuid.uuid4().hex
            connection.execute(
                "INSERT INTO jobs(id,kind,payload,status) "
                "VALUES(?, 'payment_export', ?, 'pending')",
                (job_id, canonical(payload)),
            )
            return {
                "status": "queued",
                "job_id": job_id,
                "preview_digest": preview_digest,
                "row_count": len(preview["rows"]),
                "total_fen": preview["total_fen"],
            }

        return self.engine._write(
            request_id,
            request_hash,
            epochs,
            (),
            "payment_export",
            operation,
            checked_lanes=("accounting", "material", "management"),
        )


def _verify_workbook(content: bytes, plan: dict) -> None:
    from openpyxl import load_workbook

    validate_template(content)
    book = load_workbook(BytesIO(content), data_only=False)
    try:
        sheet = book.active
        for index, row in enumerate(plan["rows"], 2):
            cells = [sheet.cell(index, column) for column in range(1, 5)]
            expected = [
                row["name"],
                row["account"],
                yuan(row["amount_fen"]),
                f"{plan['period']} {row['category']}",
            ]
            if [cell.value for cell in cells] != expected or any(
                cell.data_type != "s" for cell in cells
            ):
                raise KernelError("export_verification_failed", "已生成代发内容与冻结计划不一致")
        if any(
            cell.value is not None
            for row in sheet.iter_rows(
                min_row=len(plan["rows"]) + 2,
                max_row=sheet.max_row,
                min_col=1,
                max_col=4,
            )
            for cell in row
        ):
            raise KernelError("export_verification_failed", "代发文件包含冻结计划之外的付款记录")
    finally:
        book.close()


def _published_result(target: Path, job_id: str, plan: dict) -> dict:
    try:
        manifest = json.loads((target / MANIFEST_NAME).read_text(encoding="utf-8"))
        content = (target / WORKBOOK_NAME).read_bytes()
    except (OSError, ValueError) as exc:
        raise KernelError("export_target_conflict", "目标目录已经存在且不属于此代发任务") from exc
    checksum = hashlib.sha256(content).hexdigest()
    if (
        manifest.get("job_id") != job_id
        or manifest.get("preview_digest") != plan["digest"]
        or manifest.get("sha256") != checksum
        or manifest.get("plan") != plan
    ):
        raise KernelError("export_target_conflict", "已存在代发文件与任务的冻结计划不一致")
    _verify_workbook(content, plan)
    return {
        "directory": str(target),
        "path": str(target / WORKBOOK_NAME),
        "sha256": checksum,
        "row_count": len(plan["rows"]),
        "total_fen": plan["total_fen"],
    }


def _render_job(target: Path, job_id: str, plan: dict, template: bytes) -> dict:
    if target.exists():
        return _published_result(target, job_id, plan)
    content = render_import(template, plan["rows"], plan["period"])
    _verify_workbook(content, plan)
    manifest = {
        "job_id": job_id,
        "preview_digest": plan["digest"],
        "sha256": hashlib.sha256(content).hexdigest(),
        "plan": plan,
        "notice": "本文件仅为代发指令；生成文件不代表已付款，不产生银行流水或付款凭证。",
    }
    target.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix=".payment-export-", dir=target.parent) as temporary:
        staging = Path(temporary) / "batch"
        staging.mkdir()
        for filename, data in (
            (WORKBOOK_NAME, content),
            (MANIFEST_NAME, canonical(manifest).encode()),
        ):
            with (staging / filename).open("xb") as handle:
                handle.write(data)
                handle.flush()
                os.fsync(handle.fileno())
        if target.exists():
            return _published_result(target, job_id, plan)
        staging.rename(target)
    return _published_result(target, job_id, plan)


def run_export_jobs(engine, *, limit: int = 10, fault=None) -> list[dict]:
    """Retry frozen jobs, including a crash after file publication but before SQL completion."""
    if type(limit) is not int or not 1 <= limit <= 100:
        raise ValueError("job limit must be 1..100")
    fault = fault or (lambda stage, job_id: None)
    outcomes, attempted = [], []
    with _worker_lock(engine.store.path) as acquired:
        if not acquired:
            return outcomes
        for _ in range(limit):
            with engine.store.connection() as connection:
                connection.execute("BEGIN IMMEDIATE")
                exclude = f"AND id NOT IN ({','.join('?' for _ in attempted)})" if attempted else ""
                job = connection.execute(
                    "SELECT id,payload FROM jobs WHERE kind='payment_export' "
                    f"AND status IN ('pending','running','failed') {exclude} "
                    "ORDER BY attempts,id LIMIT 1",
                    attempted,
                ).fetchone()
                if job is None:
                    connection.rollback()
                    break
                connection.execute(
                    "UPDATE jobs SET status='running',attempts=attempts+1,last_error=NULL "
                    "WHERE id=?",
                    (job["id"],),
                )
                connection.commit()
            attempted.append(job["id"])
            try:
                payload = json.loads(job["payload"])
                plan = payload["plan"]
                if (
                    digest({key: value for key, value in plan.items() if key != "digest"}).hex()
                    != plan["digest"]
                ):
                    raise KernelError("invalid_export_plan", "冻结代发计划摘要不一致")
                if payload["template_evidence_digest"] != plan["template_evidence_digest"]:
                    raise KernelError("invalid_export_plan", "冻结代发模板引用不一致")
                if (
                    plan["company_id"] != engine.store.company_id
                    or plan["database_id"] != engine.store.database_id
                ):
                    raise KernelError("company_mismatch", "代发任务与当前公司不一致")
                with engine.store.connection(read_only=True) as connection:
                    template = connection.execute(
                        "SELECT content FROM evidence WHERE digest=?",
                        (_evidence_digest(payload["template_evidence_digest"]),),
                    ).fetchone()
                if template is None:
                    raise KernelError("missing_template", "冻结计划的原始模板不可用")
                fault("before_files", job["id"])
                result = _render_job(
                    Path(payload["output_directory"]), job["id"], plan, template[0]
                )
                fault("files_published", job["id"])
                status, error = "succeeded", None
            except Exception as exc:
                result, status = None, "failed"
                error = f"{type(exc).__name__}: {exc}"[:500]
            with engine.store.connection() as connection:
                connection.execute("BEGIN IMMEDIATE")
                connection.execute(
                    "UPDATE jobs SET status=?,last_error=?,result=? WHERE id=?",
                    (status, error, canonical(result) if result else None, job["id"]),
                )
                connection.commit()
            outcomes.append(
                {"job_id": job["id"], "status": status, "result": result, "error": error}
            )
    return outcomes
