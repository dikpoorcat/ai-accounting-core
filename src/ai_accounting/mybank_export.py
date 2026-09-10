"""MYbank workbook format and authenticated CLI; no external monetary input."""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import sys
from copy import copy
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory

from openpyxl import load_workbook


class MybankExportError(ValueError):
    """An actionable export error without account numbers."""


def yuan(fen: int) -> str:
    if type(fen) is not int:
        raise MybankExportError("金额必须来自内核整数分")
    sign = "-" if fen < 0 else ""
    whole, cents = divmod(abs(fen), 100)
    return f"{sign}{whole}.{cents:02d}"


def canonical_hash(value: object) -> str:
    return hashlib.sha256(
        json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode()
    ).hexdigest()


def read_workbook_file(path: str) -> bytes:
    try:
        return Path(path).read_bytes()
    except OSError as exc:
        raise MybankExportError(f"无法读取 {Path(path).name}，请检查文件或占用状态。") from exc


def _book(content: bytes):
    try:
        return load_workbook(BytesIO(content), data_only=False)
    except Exception as exc:
        raise MybankExportError("无法读取未加密的 Excel 工作簿") from exc


def validate_template(content: bytes) -> None:
    book = _book(content)
    try:
        if len(book.worksheets) != 1:
            raise MybankExportError("银行模板必须只有一个工作表")
        headers = [re.sub(r"\s+", "", str(c.value or "")) for c in book.active[1][:4]]
        prefixes = ("收款方名称", "收款方账号", "金额", "附言/用途")
        if len(headers) != 4 or any(
            not h.startswith(p) for h, p in zip(headers, prefixes, strict=True)
        ):
            raise MybankExportError("请使用收款方名称、收款方账号、金额、附言/用途四列银行模板")
        if any(c.data_type == "f" for row in book.active for c in row):
            raise MybankExportError("银行模板不能含公式")
    finally:
        book.close()


def read_recipients(content: bytes) -> dict[str, str]:
    """Read only names/accounts; amounts and remarks never enter the snapshot."""
    book = _book(content)
    try:
        sheet = book.active
        headers = [re.sub(r"\s+", "", str(c.value or "")) for c in sheet[1]]
        if headers[:3] == ["编号(*)", "账户名称(*)", "账号(*)"]:
            name_col, account_col = 2, 3
        elif (
            len(headers) >= 2
            and headers[0].startswith("收款方名称")
            and headers[1].startswith("收款方账号")
        ):
            name_col, account_col = 1, 2
        elif headers[:2] == ["姓名", "账号"]:
            name_col, account_col = 1, 2
        else:
            raise MybankExportError("收款资料须含姓名、账号两列，或使用已有银行收款模板")
        result: dict[str, str] = {}
        accounts: set[str] = set()
        for row in range(2, sheet.max_row + 1):
            name_cell, account_cell = sheet.cell(row, name_col), sheet.cell(row, account_col)
            if name_cell.value is None and account_cell.value is None:
                continue
            if any(
                c.data_type != "s" or not isinstance(c.value, str)
                for c in (name_cell, account_cell)
            ):
                raise MybankExportError(
                    f"收款资料第 {row} 行姓名和账号必须为文本，不能是公式或数字"
                )
            name, account = name_cell.value.strip(), account_cell.value.strip()
            if not name or not account or any(ch.isspace() for ch in account):
                raise MybankExportError(f"收款资料第 {row} 行姓名或账号无效")
            if name.startswith(("=", "+", "-", "@")) or account.startswith(("=", "+", "-", "@")):
                raise MybankExportError(f"收款资料第 {row} 行含公式标记")
            if name in result or account in accounts:
                raise MybankExportError(f"收款资料第 {row} 行姓名或账号重复，请核对")
            if re.fullmatch(r"\d+[.eE].*", account):
                raise MybankExportError(f"收款资料第 {row} 行账号疑似小数或科学计数法")
            result[name] = account
            accounts.add(account)
        return result
    finally:
        book.close()


def render_import(
    template: bytes, rows: list[dict], period: str, category: str | None = None
) -> bytes:
    validate_template(template)
    if not rows or len(rows) > 2000:
        raise MybankExportError("每个银行文件须有 1 至 2000 条付款记录")

    def memo(row):
        label = row.get("category", category)
        if not label:
            raise MybankExportError("付款类别缺失")
        value = f"{period} {label}"
        if len(value) > 40:
            raise MybankExportError("附言/用途不能超过 40 字")
        return value

    book = _book(template)
    try:
        book.template = False
        sheet = book.active
        styles = [copy(sheet.cell(2, col)._style) for col in range(1, 5)]
        # Clear only payment cells. Instructions/merged cells to the right remain intact.
        for cells in sheet.iter_rows(
            min_row=2, max_row=max(sheet.max_row, len(rows) + 1), min_col=1, max_col=4
        ):
            for cell in cells:
                cell.value = None
        for index, row in enumerate(rows, 2):
            if type(row["amount_fen"]) is not int or row["amount_fen"] <= 0:
                raise MybankExportError("仅能导出内核正数整数分金额")
            values = (row["name"], row["account"], yuan(row["amount_fen"]), memo(row))
            for col, value in enumerate(values, 1):
                cell = sheet.cell(index, col)
                cell._style = copy(styles[col - 1])
                cell.value = value
                cell.data_type = "s"
                cell.number_format = "@"
        output = BytesIO()
        book.save(output)
    finally:
        book.close()
    content = output.getvalue()
    verify = _book(content)
    try:
        if verify.template:
            raise MybankExportError("输出不是普通 XLSX 文件")
        for index, row in enumerate(rows, 2):
            cells = [verify.active.cell(index, col) for col in range(1, 5)]
            expected = [row["name"], row["account"], yuan(row["amount_fen"]), memo(row)]
            if [c.value for c in cells] != expected or any(c.data_type != "s" for c in cells):
                raise MybankExportError("输出回读校验失败")
    finally:
        verify.close()
    return content


def publish_export(output_dir: str, files: dict[str, bytes], summary: dict) -> dict:
    target = Path(output_dir).resolve()
    if target.exists():
        raise MybankExportError("输出目录已存在；请使用新的批次目录，避免覆盖旧付款文件")
    target.parent.mkdir(parents=True, exist_ok=True)
    with TemporaryDirectory(prefix=".mybank-", dir=target.parent) as temporary:
        staging = Path(temporary) / "batch"
        staging.mkdir()
        for name, content in files.items():
            (staging / name).write_bytes(content)
        summary = {
            **summary,
            "status": "generated",
            "files": [
                {"path": str(target / name), "sha256": hashlib.sha256(content).hexdigest()}
                for name, content in files.items()
            ],
        }
        (staging / "代发核对.json").write_text(
            json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8"
        )
        staging.rename(target)
    return summary


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="从已登录公司的内核事实预览/生成网商银行代发文件")
    parser.add_argument(
        "--request", required=True, help="仅含公司、期间、来源 ID 和文件路径的 JSON"
    )
    parser.add_argument("--generate", action="store_true", help="按预览哈希生成文件，默认仅预览")
    parser.add_argument("--expected-source-hash")
    parser.add_argument("--output-dir")
    parser.add_argument("--report", help="将本次内核核对结果另存为新的 JSON 文件")
    args = parser.parse_args(argv)
    from . import mcp_server
    from .config import get_settings
    from .mybank_schemas import GenerateMybankExportRequest, PreviewMybankExportRequest

    try:
        if args.report and Path(args.report).exists():
            raise MybankExportError("核对报告已存在，请使用新文件名")
        payload = json.loads(Path(args.request).read_text(encoding="utf-8-sig"))
        request = PreviewMybankExportRequest.model_validate(payload)
        if args.generate:
            request = GenerateMybankExportRequest.model_validate(
                {
                    **request.model_dump(),
                    "expected_source_hash": args.expected_source_hash,
                    "output_dir": args.output_dir,
                }
            )
        elif args.expected_source_hash or args.output_dir:
            parser.error("生成参数必须与 --generate 一起使用")
        mcp_server._initialize_mcp_credential_store(environment=get_settings().finance_environment)
        try:
            tool_name = (
                "finance_generate_mybank_export"
                if args.generate
                else "finance_preview_mybank_export"
            )
            # Registered tools carry owner authentication and company database routing.
            # Module-level handler functions deliberately do not carry those wrappers.
            tool = mcp_server.mcp._tool_manager.get_tool(tool_name)
            result = tool.fn(request=request)
        finally:
            mcp_server._clear_mcp_credential_store()
        text = json.dumps(result, ensure_ascii=False, indent=2)
        if args.report:
            with Path(args.report).open("x", encoding="utf-8") as handle:
                handle.write(text)
        print(text)
        return 0 if result.get("status") in {"ready", "generated"} else 2
    except (OSError, ValueError) as exc:
        from pydantic import ValidationError

        message = (
            "请求结构无效；不接受外部金额，请检查参数"
            if isinstance(exc, ValidationError)
            else str(exc)
        )
        print(
            json.dumps({"status": "rejected", "errors": [message]}, ensure_ascii=False),
            file=sys.stderr,
        )
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
