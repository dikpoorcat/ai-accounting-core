"""Preserved pure regressions from tests/test_mybank_export.py; no legacy service fixtures."""

from __future__ import annotations

import hashlib
import io
import re
import zipfile
from datetime import date
from io import BytesIO

import pytest
from openpyxl import Workbook, load_workbook

from ai_accounting.financial_statement_template import (
    TEMPLATE_SHA256,
    _template_bytes,
    render_quarterly_template,
)
from ai_accounting.mybank_export import MybankExportError, read_recipients


def workbook(rows, *, template=False):
    book = Workbook()
    book.template = template
    for row in rows:
        book.active.append(row)
    stream = BytesIO()
    book.save(stream)
    return stream.getvalue()


@pytest.mark.parametrize("account", [123456789012345678, "1.234567E+18", "=1+1"])
def test_unsafe_account_encoding_rejected(account):
    with pytest.raises(MybankExportError):
        read_recipients(workbook([["姓名", "账号"], ["张三", account]]))


# This preserves the legacy workbook XML/formula compatibility regression,
# replacing only its ORM business fixture with explicit renderer input.
def test_tax_template_preserves_structure_and_cached_values():
    data = {
        "organization": {
            "name": "合成 & 核验公司",
            "taxpayer_identification_number": "91330106MA1234567T",
        },
        "period": {"quarter_start": "2026-01-01", "quarter_end": "2026-03-31"},
        "statements": {
            "balance_sheet": {
                str(line): {"ending_fen": 0, "beginning_fen": 0} for line in range(1, 54)
            },
            "profit_statement": {
                str(line): {"current_fen": 0, "year_to_date_fen": 0} for line in range(1, 33)
            },
            "cash_flow_statement": {
                str(line): {"current_fen": 0, "year_to_date_fen": 0} for line in range(1, 23)
            },
        },
    }
    balance = data["statements"]["balance_sheet"]
    for line in (1, 15, 30, 52, 53):
        balance[str(line)]["ending_fen"] = 140000
    balance["48"]["ending_fen"], balance["51"]["ending_fen"] = 100000, 40000
    for line in (1, 21, 28, 32):
        data["statements"]["profit_statement"][str(line)] = {
            "current_fen": 40000,
            "year_to_date_fen": 40000,
        }
    for line, amount in (
        (1, 40000),
        (7, 40000),
        (15, 100000),
        (19, 100000),
        (20, 140000),
        (22, 140000),
    ):
        data["statements"]["cash_flow_statement"][str(line)] = {
            "current_fen": amount,
            "year_to_date_fen": amount,
        }
    generated = render_quarterly_template(data)
    template = _template_bytes()
    assert hashlib.sha256(template).hexdigest().upper() == TEMPLATE_SHA256

    source = load_workbook(io.BytesIO(template), data_only=False)
    output = load_workbook(io.BytesIO(generated), data_only=False)
    cached = load_workbook(io.BytesIO(generated), data_only=True)
    assert (
        output.sheetnames
        == source.sheetnames
        == ["资产负债表", "利润表_月季报", "现金流量表_月季报"]
    )
    for source_sheet, output_sheet in zip(source.worksheets, output.worksheets, strict=True):
        assert source_sheet.protection.sheet == output_sheet.protection.sheet
        assert tuple(source_sheet.merged_cells.ranges) == tuple(output_sheet.merged_cells.ranges)
        assert len(source_sheet.data_validations.dataValidation) == len(
            output_sheet.data_validations.dataValidation
        )
        for row in source_sheet.iter_rows():
            for source_cell in row:
                if source_cell.data_type == "f":
                    assert output_sheet[source_cell.coordinate].value == source_cell.value
    assert cached["资产负债表"]["D37"].value == 1400
    assert cached["资产负债表"]["I37"].value == 0
    assert cached["利润表_月季报"]["D3"].value == "91330106MA1234567T"
    assert cached["利润表_月季报"]["F3"].value == data["organization"]["name"]
    assert cached["利润表_月季报"]["D4"].value.date() == date(2026, 1, 1)
    assert cached["现金流量表_月季报"]["F4"].value.date() == date(2026, 3, 31)
    assert cached["利润表_月季报"]["D37"].value == 400
    assert cached["现金流量表_月季报"]["D30"].value == 1400

    with (
        zipfile.ZipFile(io.BytesIO(template)) as source_zip,
        zipfile.ZipFile(io.BytesIO(generated)) as output_zip,
    ):
        assert source_zip.namelist() == output_zip.namelist()
        for part_name, root_name in (
            ("xl/worksheets/sheet1.xml", "worksheet"),
            ("xl/worksheets/sheet2.xml", "worksheet"),
            ("xl/worksheets/sheet3.xml", "worksheet"),
            ("xl/workbook.xml", "workbook"),
        ):
            source_xml = source_zip.read(part_name).decode("utf-8")
            output_xml = output_zip.read(part_name).decode("utf-8")
            root_pattern = rf"<{root_name}\b[^>]*>"
            source_root = re.search(root_pattern, source_xml)
            output_root = re.search(root_pattern, output_xml)
            assert source_root is not None and output_root is not None
            assert output_root.group(0) == source_root.group(0)
            ignorable = re.search(r'mc:Ignorable="([^"]+)"', output_root.group(0))
            if ignorable is not None:
                for prefix in ignorable.group(1).split():
                    assert f"xmlns:{prefix}=" in output_root.group(0)
