"""Selected remuneration exports retain exact original and payment validation."""

from __future__ import annotations

import json
from io import BytesIO

import pytest
from openpyxl import Workbook, load_workbook
from test_exports import evidence, inventory, queue
from test_exports import setup as setup  # noqa: F401
from test_materials import csv_spec

from ai_accounting.kernel.contracts import KernelError
from ai_accounting.kernel.domains.adjustments import EmployeeAdvance
from ai_accounting.kernel.domains.transactions import Allocation, Expense, Payment
from ai_accounting.kernel.exports import MANIFEST_NAME, WORKBOOK_NAME, run_export_jobs
from ai_accounting.kernel.materials import Materials
from ai_accounting.kernel.periods import Periods
from ai_accounting.kernel.types import YearMonth


def selected(export, template, sources=None):
    return export.preview("2026-01", template_evidence_digest=template,
                          source_ids=sources or ["january"])


def receive(company, subject, raw, *, category="payroll", period="2026-01", spec=None):
    proof = evidence(company, raw, subject)
    source = Materials(company.engine).receive(
        subject,
        {"period": period, "category": category, "purpose": "business",
         "evidence_digest": proof, "specification": spec or csv_spec()},
        evidence=(proof,), expected_revision=0, request_id=company.request(),
    )
    return dict(source, period=period), proof


def resolve(company, source, proof, location="CSV!B2", *, subject="january", amount=1000000,
            field="result.gross_fen", **changes):
    with company.engine.store.connection(read_only=True) as connection:
        calculation_id = connection.execute(
            "SELECT calculation_id FROM calculation_current WHERE subject_id=?", (subject,),
        ).fetchone()[0]
    data = {
        "period": source["period"], "source_id": source["subject_id"],
        "source_fact_id": source["fact_id"], "location": location, "treatment": "recognize",
        "links": [{"calculation_id": calculation_id, "amount_field": field, "amount_fen": amount}],
        **changes,
    }
    return Materials(company.engine).resolve(
        "resolution-" + source["subject_id"] + "-" + location, data,
        evidence=(proof,), expected_revision=0, request_id=company.request(),
    )


def assert_incomplete(export, template, *, sources=None, code=None):
    with pytest.raises(KernelError) as error:
        selected(export, template, sources)
    assert error.value.code == "materials_incomplete"
    issues = error.value.details["fact_issues"]
    if code:
        assert any(issue.get("code") == code for issue in issues), issues
    return issues


def test_selected_salary_can_deliver_file_while_complete_export_and_close_keep_month_checks(
    setup, tmp_path,
):
    company, export, template = setup
    company.save(Expense(period="2026-01", counterparty_id="supplier", amount_fen=12345,
                         expense_class="administration", creditor_kind="supplier"), "other-expense")
    Periods(company.engine).inventory(
        "2026-01", "transactions", evidence=[], expected=1, no_business=False,
        confirmation_evidence=company.owner_confirmation, request_id=company.request(),
    )
    before = {table: company.count(table)
              for table in ("fact_revision", "calculation", "voucher_version", "voucher_line")}
    with company.engine.store.connection(read_only=True) as connection:
        balances = [tuple(row) for row in connection.execute(
            "SELECT category,balance_key,amount FROM balance ORDER BY category,balance_key"
        )]
    directory = tmp_path / "selected-salary-file"
    preview, queued, _ = queue(company, export, template, directory, source_ids=["january"])
    assert preview["total_fen"] == 907400 and preview["scope"] == "selected"
    assert preview["inventories"]["2026-01"] == {}
    assert not directory.exists()
    completed = run_export_jobs(company.engine)
    assert completed[0]["job_id"] == queued["job_id"] and completed[0]["status"] == "succeeded"
    book = load_workbook(directory / WORKBOOK_NAME, data_only=False)
    try:
        assert [cell.value for cell in book.active[2]] == [
            "张三", "001234567890", "9074.00", "2026-01 工资",
        ]
        assert all(cell.data_type == "s" for cell in book.active[2])
    finally:
        book.close()
    manifest = json.loads((directory / MANIFEST_NAME).read_text(encoding="utf-8"))
    assert manifest["plan"] == preview
    assert {table: company.count(table) for table in before} == before
    with company.engine.store.connection(read_only=True) as connection:
        assert [tuple(row) for row in connection.execute(
            "SELECT category,balance_key,amount FROM balance ORDER BY category,balance_key"
        )] == balances
    with pytest.raises(KernelError) as complete:
        export.preview("2026-01", template_evidence_digest=template)
    assert complete.value.code == "materials_incomplete"
    with company.engine.store.connection(read_only=True) as connection:
        _, issues, unpublished = Periods.completeness(
            connection, YearMonth("2026-01").ordinal, company.engine.store.registry,
        )
    assert issues and any(row["id"] == "other-expense" for row in unpublished)


@pytest.mark.parametrize("other_label", ["未选择工资", "普通业务混合行"])
def test_related_original_requires_current_month_unselected_and_mixed_rows(setup, other_label):
    company, export, template = setup
    source, proof = receive(
        company, "shared-salary",
        f"name,amount,period\n已选择工资,10000.00,2026-01\n{other_label},10.00,2026-01\n".encode(),
    )
    resolve(company, source, proof)
    issues = assert_incomplete(export, template, code="material_item_unresolved")
    assert any(issue.get("source_id") == "shared-salary" and issue.get("location") == "CSV!B3"
               for issue in issues)


def test_unlinked_salary_original_is_in_scope_and_unrelated_original_is_not(setup):
    company, export, template = setup
    receive(company, "unrelated-office", b"name,amount,period\noffice,10.00,2026-01\n",
            category="transactions")
    assert selected(export, template)["total_fen"] == 907400
    receive(company, "unlinked-salary", b"name,amount,period\nwage,10.00,2026-01\n")
    issues = assert_incomplete(export, template, code="material_item_unresolved")
    assert any(issue.get("source_id") == "unlinked-salary" for issue in issues)


def test_received_salary_inventory_original_without_material_source_still_blocks(setup):
    company, export, template = setup
    proof = evidence(company, b"name,amount,period\nunregistered,200.00,2026-01\n",
                     "received-but-not-registered")
    originals = sorted({item for period, item in company.materials["payroll"]
                        if period == "2026-01"} | {proof})
    Periods(company.engine).inventory(
        "2026-01", "payroll", evidence=originals, expected=len(originals), no_business=False,
        confirmation_evidence=company.owner_confirmation, request_id=company.request(),
    )
    assert_incomplete(export, template, code="material_source_not_registered")


def test_explicit_future_row_keeps_existing_period_semantics(setup):
    company, export, template = setup
    source, proof = receive(
        company, "salary-with-future",
        b"name,amount,period\njanuary,10000.00,2026-01\nfuture,10.00,2026-02\n",
    )
    resolve(company, source, proof)
    assert selected(export, template)["total_fen"] == 907400


@pytest.mark.parametrize("receipt_period", ["2026-01", "2026-02"])
def test_future_attribution_cannot_hide_current_salary_result_in_any_checker(setup, receipt_period):
    company, export, template = setup
    source, proof = receive(company, "misdated-current-salary",
                            b"name,amount,period\nwage,10000.00,2026-02\n",
                            period=receipt_period)
    resolve(company, source, proof)
    assert_incomplete(export, template, code="material_period_mismatch")
    with pytest.raises(KernelError) as complete:
        export.preview("2026-01", template_evidence_digest=template)
    assert complete.value.code == "materials_incomplete"
    assert any(issue.get("code") == "material_period_mismatch"
               for issue in complete.value.details["fact_issues"])
    with company.engine.store.connection(read_only=True) as connection:
        _, issues, _ = Periods.completeness(
            connection, YearMonth("2026-01").ordinal, company.engine.store.registry,
        )
    assert any(issue.get("code") == "material_period_mismatch" for issue in issues)


def test_reverse_duplicate_original_expands_related_source_closure(setup):
    company, export, template = setup
    original, proof = receive(company, "salary-original",
                              b"name,amount,period\nwage,10000.00,2026-01\n")
    resolve(company, original, proof)
    copy, copy_proof = receive(
        company, "later-duplicate",
        b"name,amount,period\ncopy,10000.00,2026-01\nunfinished,10.00,2026-01\n",
        category="transactions", period="2026-02",
    )
    Materials(company.engine).resolve(
        "duplicate-resolution", {"period": "2026-02", "source_id": copy["subject_id"],
                                 "source_fact_id": copy["fact_id"], "location": "CSV!B2",
                                 "treatment": "duplicate", "duplicate_source_id": "salary-original",
                                 "duplicate_location": "CSV!B2", "reason": "相同工资的重复原件"},
        evidence=(copy_proof,), expected_revision=0, request_id=company.request(),
    )
    issues = assert_incomplete(export, template, code="material_item_unresolved")
    assert any(issue.get("source_id") == "later-duplicate" for issue in issues)


def test_competing_source_cannot_consume_selected_salary_gross_amount_twice(setup):
    company, export, template = setup
    for subject, label, category, period in (
        ("salary-original", "wage", "payroll", "2026-01"),
        ("competing-copy", "copy", "transactions", "2026-02"),
    ):
        source, proof = receive(company, subject,
                                f"name,amount,period\n{label},10000.00,2026-01\n".encode(),
                                category=category, period=period)
        resolve(company, source, proof)
    assert_incomplete(export, template, code="material_business_overallocated")


def withdraw(company, source_id):
    preview = company.engine.preview_delete(source_id)
    return company.engine.delete(
        source_id, preview_digest=preview["digest"], epochs=preview["epochs"],
        request_id=company.request(),
    )


def test_withdrawn_salary_original_own_disposition_does_not_reenter_any_export_scope(setup):
    company, export, template = setup
    for source_id, label in (("first-salary", "first"), ("second-salary", "second")):
        source, proof = receive(company, source_id,
                                f"name,amount,period\n{label},10000.00,2026-01\n".encode())
        resolve(company, source, proof)
    assert_incomplete(export, template, code="material_business_overallocated")
    withdraw(company, "first-salary")
    for source_ids in (["january"], None):
        preview = export.preview("2026-01", template_evidence_digest=template,
                                 source_ids=source_ids)
        assert preview["total_fen"] == 907400
        assert all(summary["source_id"] != "first-salary"
                   for summary in preview["material_coverage"]["2026-01"]["file_summaries"])


def test_active_salary_duplicate_keeps_withdrawn_target_error_in_every_export_scope(setup):
    company, export, template = setup
    original, proof = receive(company, "salary-original",
                              b"name,amount,period\nwage,10000.00,2026-01\n")
    resolve(company, original, proof)
    copy, copy_proof = receive(company, "active-salary-copy",
                              b"name,amount,period\ncopy,10000.00,2026-01\n")
    Materials(company.engine).resolve(
        "salary-copy-resolution", {"period": "2026-01", "source_id": copy["subject_id"],
                                   "source_fact_id": copy["fact_id"], "location": "CSV!B2",
                                   "treatment": "duplicate",
                                   "duplicate_source_id": "salary-original",
                                   "duplicate_location": "CSV!B2", "reason": "重复工资原件"},
        evidence=(copy_proof,), expected_revision=0, request_id=company.request(),
    )
    withdraw(company, "salary-original")
    for source_ids in (["january"], None):
        with pytest.raises(KernelError) as error:
            export.preview("2026-01", template_evidence_digest=template, source_ids=source_ids)
        assert error.value.code == "materials_incomplete"
        assert any(issue.get("code") == "material_source_missing"
                   and issue.get("source_id") == "salary-original"
                   for issue in error.value.details["fact_issues"])


def test_related_control_total_is_checked_even_when_another_row_is_future(setup):
    company, export, template = setup
    book = Workbook()
    sheet = book.active
    sheet.title = "工资"
    for row in (["人员", "金额", "所属期"], ["已选", 10000, "2026-01"],
                ["未来", 20, "2026-02"], ["合计", 10001, "2026-01"]):
        sheet.append(row)
    output = BytesIO()
    book.save(output)
    book.close()
    source, proof = receive(company, "salary-control", output.getvalue(),
                            spec=csv_spec() | {"format": "xlsx", "total_rows": {"工资": [4]}})
    resolve(company, source, proof, location="工资!B2")
    assert_incomplete(export, template, code="material_total_mismatch")


def test_hidden_related_salary_row_is_read_and_blocks_selected_export(setup):
    company, export, template = setup
    book = Workbook()
    sheet = book.active
    sheet.title = "工资"
    sheet.append(["人员", "金额", "所属期"])
    sheet.append(["已选人员", 10000, "2026-01"])
    sheet.append(["隐藏人员", 20, "2026-01"])
    sheet.row_dimensions[3].hidden = True
    output = BytesIO()
    book.save(output)
    book.close()
    source, proof = receive(company, "hidden-salary", output.getvalue(),
                            spec=csv_spec() | {"format": "xlsx"})
    resolve(company, source, proof, location="工资!B2")
    issues = assert_incomplete(export, template, code="material_item_unresolved")
    assert any(issue.get("location") == "工资!B3" for issue in issues)


def test_selected_remuneration_kind_gate_keeps_settled_reimbursement_in_month_scope(setup):
    company, export, template = setup
    company.save(Expense(period="2026-01", counterparty_id="employee", amount_fen=1000,
                         expense_class="administration", creditor_kind="employee"), "reimbursement")
    company.publish("reimbursement")
    company.save(Payment(
        period="2026-01", actual_date="2026-01-20", direction="outflow", bank_account_id="bank",
        counterparty_id="employee", amount_fen=1000,
        allocations=(Allocation(source_kind="expense", source_id="reimbursement",
                                obligation="primary", amount_fen=1000),),
    ), "reimbursement-paid")
    company.publish("reimbursement-paid")
    inventory(company)
    Periods(company.engine).inventory(
        "2026-01", "assets", evidence=[], expected=1, no_business=False,
        confirmation_evidence=company.owner_confirmation, request_id=company.request(),
    )
    assert_incomplete(export, template, sources=["january", "reimbursement"])


@pytest.mark.parametrize("obligation", ["net", "employer_social"])
def test_cross_month_unpublished_salary_or_social_payment_blocks_selected_export(setup, obligation):
    company, export, template = setup
    company.save(Payment(
        period="2026-02", actual_date="2026-02-10", direction="outflow", bank_account_id="bank",
        counterparty_id="employee", amount_fen=10000,
        allocations=(Allocation(source_kind="payroll", source_id="january",
                                obligation=obligation, amount_fen=10000),),
    ), "partial-pay")
    issues = assert_incomplete(export, template)
    assert any(issue["field"] == "partial-pay" for issue in issues)


def test_cross_month_payment_does_not_pull_in_unrelated_future_salary_original(setup):
    company, export, template = setup
    company.save(Payment(
        period="2026-02", actual_date="2026-02-10", direction="outflow", bank_account_id="bank",
        counterparty_id="employee", amount_fen=10000,
        allocations=(Allocation(source_kind="payroll", source_id="january",
                                obligation="net", amount_fen=10000),),
    ), "partial-pay")
    company.publish("partial-pay")
    receive(company, "future-salary", b"name,amount,period\nfebruary,20.00,2026-02\n",
            period="2026-02")
    assert selected(export, template)["total_fen"] == 897400


def test_unpublished_personal_salary_advance_is_a_related_payment_not_an_unrelated_expense(setup):
    company, export, template = setup
    company.save(EmployeeAdvance(
        period="2026-02", payer_id="owner", payer_kind="owner", payment_on_behalf_confirmed=True,
        actual_creditor_payment_date="2026-02-10",
        sources=(Allocation(source_kind="payroll", source_id="january", obligation="net",
                            amount_fen=10000),),
    ), "salary-advance")
    issues = assert_incomplete(export, template)
    assert any(issue["field"] == "salary-advance" for issue in issues)


def test_cross_month_published_salary_payment_keeps_its_original_row_validation(setup):
    company, export, template = setup
    company.save(Payment(
        period="2026-02", actual_date="2026-02-10", direction="outflow", bank_account_id="bank",
        counterparty_id="employee", amount_fen=10000,
        allocations=(Allocation(source_kind="payroll", source_id="january",
                                obligation="net", amount_fen=10000),),
    ), "partial-pay")
    company.publish("partial-pay")
    with company.engine.store.connection(read_only=True) as connection:
        payment = company.engine.store.current_fact(connection, "partial-pay")
    proof = evidence(company, b"name,amount,period\npayment,101.00,2026-02\n", "partial-payment")
    # The original differs from the adopted payment by one yuan. Its February
    # row must be checked even though the remuneration selection is January.
    source = Materials(company.engine).receive(
        "payment-original", {"period": "2026-02", "category": "bank", "purpose": "business",
                             "evidence_digest": proof, "specification": csv_spec()},
        evidence=(proof,), expected_revision=0, request_id=company.request(),
    )
    source = dict(source, period="2026-02")
    resolve(company, source, proof, subject="partial-pay", amount=10000,
            field="fact.amount_fen")
    assert payment.fact.period == "2026-02"
    assert_incomplete(export, template, code="material_split_mismatch")
