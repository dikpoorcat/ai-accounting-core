"""Synthetic staged dialogue seed, preparation and dependency recovery checks.

These tests verify real kernel behavior; conversational question selection is
evaluated by the independent agent through staged_accounting_agent_harness.
"""

from __future__ import annotations

import hashlib
import json
import shutil
from pathlib import Path

import pytest
import staged_accounting_agent_harness as harness
from test_payroll import actual
from test_work_draft_service import business_state

from ai_accounting.kernel.contracts import NeedsInformation
from ai_accounting.kernel.service import LocalService


@pytest.fixture
def seeded(tmp_path):
    root = tmp_path / "data"
    app = LocalService(root)
    app.security.provision(harness.reading.OWNER, harness.reading.PASSWORD)
    token = app.security.login(harness.reading.OWNER, harness.reading.PASSWORD).session_token
    info = harness.seed(app, token, tmp_path)
    yield app, token, info, tmp_path
    app.close()


@pytest.fixture
def recovery_seeded(tmp_path):
    app = LocalService(tmp_path / "data")
    app.security.provision(harness.reading.OWNER, harness.reading.PASSWORD)
    token = app.security.login(harness.reading.OWNER, harness.reading.PASSWORD).session_token
    info = harness.seed(app, token, tmp_path, scenario="stage4_recovery")
    yield app, token, info, tmp_path
    app.close()


@pytest.fixture
def cross_month_seeded(tmp_path):
    if shutil.which("pwsh") is None:
        pytest.skip("PowerShell 7 is required for isolated desktop delivery")
    app = LocalService(tmp_path / "data")
    app.security.provision(harness.reading.OWNER, harness.reading.PASSWORD)
    token = app.security.login(harness.reading.OWNER, harness.reading.PASSWORD).session_token
    try:
        info = harness.seed(app, token, tmp_path, scenario="cross_month_payment")
        yield app, token, info, tmp_path
    finally:
        app.close()


def dispatch(app, token, command, **payload):
    return app.dispatch(command, payload, session_token=token)


def test_seed_keeps_all_kernel_issues_and_four_bank_rows_without_disclosing_answers(seeded):
    app, token, info, directory = seeded
    result = dispatch(app, token, "workflow", company_id=info["company_id"],
                      period=info["period"], as_of=info["as_of"])
    unresolved = [issue for issue in result["issues"]
                  if issue.get("source_id") == "stage-bank-source"]
    assert {issue["location"] for issue in unresolved} == {
        "CSV!B2", "CSV!B3", "CSV!B4", "CSV!B5",
    }
    assert {issue["field"] for issue in result["issues"]} >= {
        "materials.bank", "materials.payroll", "materials.transactions",
        "materials.tax", "materials.assets", "materials.financing", "unpublished_payroll",
    }
    public = json.loads((directory / "seed.json").read_text(encoding="utf-8"))
    driver = json.loads((directory / "evaluator-owner-turns.json").read_text(encoding="utf-8"))
    assert public == info
    assert "bank_answers" not in public and "owner_answer_file" not in public
    assert len(driver["bank_answers"]) == 4
    assert not any(answer in json.dumps(public, ensure_ascii=False)
                   for answer in driver["bank_answers"].values())


def test_future_source_preparation_and_draft_do_not_publish_or_clear_complete_checks(seeded):
    app, token, info, _ = seeded
    company = info["company_id"]
    before = business_state(app, company)
    for area in ("payroll", "transactions"):
        context = dispatch(app, token, "work_context", company_id=company,
                           period=info["period"], work_area=area, include_work_draft=True)
        assert context["materials"] and context["work_draft"]["status"] == "absent"
        source = next(item for item in info["sources"] if area in item["source_id"])
        dispatch(app, token, "inspect_material", company_id=company,
                 evidence_digest=source["evidence_digest"], specification=source["specification"])
        dispatch(app, token, "save_work_draft", company_id=company, period=info["period"],
                 work_area=area, expected_revision=None,
                 draft={"source_refs": [{"source_id": source["source_id"],
                                         "evidence_digest": source["evidence_digest"]}],
                        "resume_note": "资料已提前整理，尚未到所属办理步骤。",
                        "next_step": "到本步骤时核对当前来源与正式状态，再询问必要缺项。"})
    assert business_state(app, company) == before
    engine = app.engine(company)
    with engine.store.connection(read_only=True) as connection:
        assert connection.execute("SELECT count(*) FROM calculation_publication").fetchone()[0] == 0
        assert connection.execute("SELECT count(*) FROM voucher").fetchone()[0] == 0
    workflow = dispatch(app, token, "workflow", company_id=company,
                        period=info["period"], as_of=info["as_of"])
    assert any(issue.get("source_id") == "stage-transactions-source"
               for issue in workflow["issues"])


def test_confirmed_stage2_scheme_waits_for_actual_contributions_before_stage4_net_preview(seeded):
    app, token, info, _ = seeded
    engine = app.engine(info["company_id"])
    from test_payroll import context_for

    from ai_accounting.kernel.payroll_confirmation import resolve_payroll_confirmation

    with engine.store.connection(read_only=True) as connection:
        wage = engine.store.current_fact(connection, "stage-january-wage")
        plan = engine.store.current_fact(connection, "stage-january-plan")
        profile = engine.store.current_fact(connection, "stage-payroll-profile")
        policy = engine.store.current_fact(connection, "stage-contributions-policy")
    confirmation = resolve_payroll_confirmation(
        wage, context_for(wage, facts=(plan, profile, policy)),
        profile_version=profile, contribution_policy_version=policy,
    )
    assert confirmation
    assert plan.fact.payroll_revision.subject_id == wage.subject_id
    assert plan.fact.payroll_revision.revision == wage.revision and plan.evidence
    with pytest.raises(NeedsInformation) as error:
        engine.preview(["stage-january-wage"])
    issues = error.value.response()["fact_issues"]
    assert any(issue["field"] == "contribution_actual.items" for issue in issues)
    dispatch(app, token, "save_work_draft", company_id=info["company_id"],
             period=info["period"], work_area="payroll", expected_revision=None,
             draft={"resume_note": "第2阶段工资方案及个税依据已有确认；第3阶段实际社保尚缺。",
                    "next_step": "第3阶段取得实际社保后，第4阶段核算工资净额并准备薪资代发。"})
    with pytest.raises(NeedsInformation):
        engine.preview(["stage-january-wage"])
    evidence = engine.register_evidence(
        "合成来源：2026年1月该人员养老金个人1000元、单位2000元，已申报。".encode(),
        "text/plain", "合成社保明细", request_id="test-contribution-source",
    )["digest"]
    contribution = actual().model_copy(update={"employee_id": info["employee_id"]})
    dispatch(app, token, "save_fact", company_id=info["company_id"],
             kind=contribution.kind, subject_id="stage-actual-contributions",
             data=contribution.model_dump(mode="json"), evidence=[evidence],
             expected_revision=0, request_id="test-contribution-fact")
    preview = engine.preview(["stage-january-wage"])
    assert preview["status"] == "preview"
    calculated = preview["results"][0]
    assert calculated["values"]["net_fen"] == 887400
    assert calculated["values"]["payroll_confirmation"] == confirmation
    with engine.store.connection(read_only=True) as connection:
        assert connection.execute("SELECT count(*) FROM calculation_publication").fetchone()[0] == 0


def test_eight_stage_checkpoint_restarts_with_company_and_period_isolation(seeded):
    app, token, info, directory = seeded
    scope = {"company_id": info["company_id"], "period": info["period"], "work_area": "bank"}
    source = info["sources"][0]
    driver = json.loads((directory / "evaluator-owner-turns.json").read_text(encoding="utf-8"))
    draft = {
        "questions": [
            {"id": "BANK-1", "question": "1月5日100元是什么用途？",
             "answer": driver["bank_answers"]["BANK-1"],
             "source_refs": [{"source_id": source["source_id"], "location": "CSV!B2"}]},
            {"id": "BANK-4", "question": "1月8日400元是什么用途？"},
        ],
        "resume_note": "八阶段第1阶段流水识别；用途问题分轮核对，本轮已收到第1笔答复。",
        "next_step": "接续第1阶段尚未回答的用途问题；工资方案到第2阶段复用，代发到第4阶段。",
    }
    saved = dispatch(app, token, "save_work_draft", **scope, expected_revision=None, draft=draft)
    before = business_state(app, scope["company_id"])
    app.close()
    restarted = LocalService(directory / "data")
    try:
        new_token = restarted.security.login(
            harness.reading.OWNER, harness.reading.PASSWORD,
        ).session_token
        assert restarted.security.catalog_instance_id == info["catalog_id"]
        assert {(c["id"], c["database_id"]) for c in restarted.catalog.companies()} == {
            (c["id"], c["database_id"]) for c in info["companies"]
        }
        recovered = dispatch(restarted, new_token, "work_context", **scope,
                             include_work_draft=True)["work_draft"]
        assert recovered["revision"] == saved["revision"] and recovered["draft"] == draft
        assert "answer" not in recovered["draft"]["questions"][1]
        for change in ({"company_id": info["companies"][1]["id"]}, {"period": "2026-02"}):
            other = dispatch(restarted, new_token, "work_context", **{**scope, **change},
                             include_work_draft=True)
            assert other["work_draft"]["status"] == "absent"
        assert business_state(restarted, scope["company_id"]) == before
    finally:
        restarted.close()


def test_recovery_keeps_stage4_block_and_stale_stage5_draft(recovery_seeded):
    app, token, info, directory = recovery_seeded
    scope = info["work_context_payload"]
    context = dispatch(app, token, "work_context", **scope)
    draft = context["work_draft"]["draft"]
    assert "进入第5步" in draft["next_step"]
    assert "旧接口不支持" in draft["resume_note"]
    questions = {question["id"]: question for question in draft["questions"]}
    assert questions["prior-stages"]["status"] == "resolved"
    assert "实际养老金个人1000元" in questions["prior-stages"]["answer"]
    pending = {questions[key]["status"] for key in ("labor-net", "stage5-tax", "stage5-office")}
    assert pending == {"pending"}
    template_ref = next(
        ref for ref in draft["source_refs"]
        if ref.get("evidence_digest") == info["template_evidence_digest"]
    )
    assert template_ref["location"] == "Sheet!A1:D1"
    assert "第4步" in template_ref["note"] and "银行四列代发模板" in template_ref["note"]
    assert "preview_export.template_evidence_digest" in template_ref["note"]
    tax_ref = next(ref for ref in draft["source_refs"] if "第5步" in ref.get("note", ""))
    assert "印花税与附加税来源" in tax_ref["note"]
    assert tax_ref["evidence_digest"] != template_ref["evidence_digest"]
    engine = app.engine(info["company_id"])
    with engine.store.connection(read_only=True) as connection:
        opening = engine.store.current_fact(connection, "recovery-opening-bank")
        statement = engine.store.current_fact(connection, "stage-bank-statement")
        template_type = connection.execute(
            "SELECT media_type FROM evidence WHERE digest=?",
            (bytes.fromhex(template_ref["evidence_digest"]),),
        ).fetchone()[0]
        payees = connection.execute(
            "SELECT p.name,p.account,e.digest,e.content FROM payee_revision p "
            "JOIN evidence e ON p.evidence_digest=e.digest ORDER BY p.name",
        ).fetchall()
    assert opening.fact.balance_fen == statement.fact.opening_fen == 1_000_000
    assert template_type == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert {(row["name"], row["account"]) for row in payees} == {
        ("合成员工甲", "001234567890"), ("合成劳务甲", "001234567891"),
    }
    assert any(ref.get("evidence_digest") == info["payee_evidence_digest"]
               for ref in draft["source_refs"])
    for row in payees:
        assert row["digest"].hex() == info["payee_evidence_digest"]
        assert f"收款人：{row['name']}，账号：{row['account']}。" in row["content"].decode("utf-8")
    before = business_state(app, info["company_id"])
    with pytest.raises(NeedsInformation) as missing:
        dispatch(app, token, "preview_export", company_id=info["company_id"], period="2026-01",
                 template_evidence_digest=info["template_evidence_digest"],
                 source_ids=["stage-january-wage", "recovery-gross-labor"])
    issue = missing.value.response()["fact_issues"][0]
    assert issue["field"] == "net_disbursement_tax"
    assert "recovery-gross-labor" in issue["reusable_sources"]
    assert business_state(app, info["company_id"]) == before
    driver = json.loads((directory / "evaluator-owner-turns.json").read_text(encoding="utf-8"))
    assert driver["acceptance"]["expected_question_ids"] == ["labor-net"]
    assert driver["acceptance"]["deferred_question_ids"] == ["stage5-tax", "stage5-office"]
    assert "acceptance" not in json.loads((directory / "seed.json").read_text(encoding="utf-8"))


def test_recovery_supports_labor_tax_without_stage5_writes(recovery_seeded):
    from ai_accounting.kernel.remuneration_tax_basis import RemunerationTaxBasis

    app, token, info, _ = recovery_seeded
    engine = app.engine(info["company_id"])
    basis = RemunerationTaxBasis(
        period="2026-01", source_kind="labor_accrual", source_id="recovery-gross-labor",
        source_fact_id=info["labor_fact"]["fact_id"], tax_period="2026-01", tax_amount_fen=8_000,
        basis_type="final_statement", confirmed=True,
    )
    # A later explicit synthetic owner answer is available only to this kernel
    # capability test. The independent recovery agent must still ask for it.
    proof = engine.register_evidence(
        "合成负责人测试答复：本项劳务最终扣税80元，净额920元。".encode(),
        "text/plain", "合成最终报酬表", request_id="test-final-labor-proof",
    )["digest"]
    dispatch(app, token, "save_fact", company_id=info["company_id"], kind=basis.kind,
             subject_id="test-final-labor-tax", data=basis.model_dump(mode="json"),
             evidence=[proof], expected_revision=0, request_id="test-final-labor-tax")
    preview = engine.preview(["recovery-gross-labor"])
    values = preview["results"][0]["values"]
    assert (values["gross_fen"], values["tax_fen"], values["net_fen"]) == (100_000, 8_000, 92_000)
    with engine.store.connection(read_only=True) as connection:
        assert connection.execute("SELECT count(*) FROM voucher").fetchone()[0] == 2
        assert connection.execute(
            "SELECT count(*) FROM subject WHERE kind IN ('expense','tax_accrual')"
        ).fetchone()[0] == 0
        assert connection.execute("SELECT count(*) FROM jobs").fetchone()[0] == 0


def test_recovery_checkpoint_restart_preserves_identity(recovery_seeded):
    app, token, info, directory = recovery_seeded
    before = business_state(app, info["company_id"])
    original = dispatch(app, token, "work_context", **info["work_context_payload"])["work_draft"]
    app.close()
    restarted = LocalService(directory / "data")
    try:
        new_token = restarted.security.login(
            harness.reading.OWNER, harness.reading.PASSWORD,
        ).session_token
        recovered = dispatch(
            restarted, new_token, "work_context", **info["work_context_payload"],
        )["work_draft"]
        assert restarted.security.catalog_instance_id == info["catalog_id"]
        assert recovered["draft"] == original["draft"]
        assert recovered["revision"] == info["stale_draft_revision"]
        assert business_state(restarted, info["company_id"]) == before
    finally:
        restarted.close()


def test_recovery_scenario_option_is_forwarded_to_isolated_mcp_process(tmp_path):
    params = harness.parameters(tmp_path, scenario="stage4_recovery", resume=True)
    assert params.args[-2:] == ["--scenario", "stage4_recovery"]
    assert "--resume" in params.args


def test_cross_month_recovery_has_complete_verified_delivery_and_no_assumed_payment(
    cross_month_seeded,
):
    from openpyxl import load_workbook

    app, token, info, directory = cross_month_seeded
    assert info["period"] == "2026-09" and info["as_of"] == "2026-10-14"
    context = dispatch(app, token, "work_context", **info["work_context_payload"])
    draft = context["work_draft"]["draft"]
    assert "登记付款后再进入九月第五步" in draft["next_step"]
    questions = {question["id"]: question for question in draft["questions"]}
    assert questions["labor-net"]["status"] == "resolved"
    assert questions["payment-success"]["answer"] == "已全部支付成功"
    for key in ("payment-date", "payment-bank"):
        assert questions[key]["status"] == "pending"
        assert questions[key]["target"].get("period") is None
    assert questions["stage5-office"]["target"]["period"] == "2026-09"
    assert questions["october-office"]["target"]["period"] == "2026-10"
    completed = info["completed_export"]
    assert completed["status"] == "succeeded"
    assert completed["result"]["row_count"] == 2
    assert completed["result"]["total_fen"] == 979_400
    assert {item["subject_id"] for row in info["export_preview"]["rows"]
            for item in row["sources"]} == {info["payroll_subject_id"], "recovery-gross-labor"}
    delivered = Path(info["desktop_delivery"]["file_path"])
    assert delivered.parent == directory / "test-desktop"
    assert hashlib.sha256(delivered.read_bytes()).hexdigest() == completed["result"]["sha256"]
    assert delivered.read_bytes() == Path(completed["result"]["path"]).read_bytes()
    book = load_workbook(delivered)
    try:
        assert {tuple(cell.value for cell in row) for row in book.active.iter_rows(min_row=2)} == {
            ("合成员工甲", "001234567890", "8874.00", "2026-09 工资"),
            ("合成劳务甲", "001234567891", "920.00", "2026-09 劳务"),
        }
    finally:
        book.close()
    engine = app.engine(info["company_id"])
    with engine.store.connection(read_only=True) as connection:
        assert connection.execute(
            "SELECT count(*) FROM subject WHERE kind='payment'"
        ).fetchone()[0] == 0
        pending_net = connection.execute(
            "SELECT sum(amount) FROM balance WHERE category='payable' AND balance_key LIKE '%:net'"
        ).fetchone()[0]
        assert pending_net == 979_400
        declaration = engine.store.current_fact(connection, "month-october-september-declaration")
        assert str(declaration.fact.period) == "2026-10"
        assert str(declaration.fact.tax_period) == "2026-09"
        assert str(declaration.fact.declaration_date) == "2026-10-10"
        assert declaration.fact.declared_tax_fen == 12_600
        assert declaration.evidence == (info["declaration_evidence"],)
        source = engine.store.current_fact(connection, "month-october-office-source")
        assert str(source.fact.period) == "2026-10"
    public = json.loads((directory / "seed.json").read_text(encoding="utf-8"))
    driver = json.loads((directory / "evaluator-owner-turns.json").read_text(encoding="utf-8"))
    assert public == info and "acceptance" not in public
    assert "expected_question_ids" not in json.dumps(public)
    assert driver["acceptance"]["expected_question_ids"] == ["stage5-office"]


def test_cross_month_checkpoint_restarts_with_identity_and_files_preserved(cross_month_seeded):
    app, token, info, directory = cross_month_seeded
    before = business_state(app, info["company_id"])
    original = dispatch(app, token, "work_context", **info["work_context_payload"])["work_draft"]
    app.close()
    restarted = LocalService(directory / "data")
    try:
        new_token = restarted.security.login(
            harness.reading.OWNER, harness.reading.PASSWORD,
        ).session_token
        recovered = dispatch(
            restarted, new_token, "work_context", **info["work_context_payload"],
        )["work_draft"]
        assert restarted.security.catalog_instance_id == info["catalog_id"]
        assert {(item["id"], item["database_id"]) for item in restarted.catalog.companies()} == {
            (item["id"], item["database_id"]) for item in info["companies"]
        }
        assert recovered["revision"] == info["stale_draft_revision"]
        assert recovered["draft"] == original["draft"]
        assert business_state(restarted, info["company_id"]) == before
        delivery_bytes = Path(info["desktop_delivery"]["file_path"]).read_bytes()
        assert hashlib.sha256(delivery_bytes).hexdigest() == (
            info["completed_export"]["result"]["sha256"]
        )
    finally:
        restarted.close()


def test_cross_month_scenario_option_is_forwarded_to_isolated_mcp_process(tmp_path):
    params = harness.parameters(tmp_path, scenario="cross_month_payment", resume=True)
    assert params.args[-2:] == ["--scenario", "cross_month_payment"]
    assert "--resume" in params.args


def test_unknown_staged_scenario_is_rejected(tmp_path):
    with pytest.raises(KeyError, match="unknown"):
        harness.seed(None, None, tmp_path, scenario="unknown")
