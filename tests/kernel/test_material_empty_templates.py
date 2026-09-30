"""Empty formula templates retain physical coverage without inventing amounts."""

import json
from io import BytesIO
from xml.etree import ElementTree as ET

import pytest
from openpyxl import Workbook
from openpyxl.styles import PatternFill
from test_material_capacity import replace_xml

from ai_accounting.kernel import material_inspection_v1
from ai_accounting.kernel.materials import Column, ControlTotal, Specification, inspect_bytes

MAIN = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"


def spec(**changes):
    return Specification(
        format="xlsx", header_rows={"Sheet": 1},
        columns=(Column(column="A", role="context"), Column(column="B", role="context"),
                 Column(column="H", role="context"),
                 *(Column(column=x, role="amount") for x in ("AC", "AD", "AE", "AF"))),
        **changes,
    )


def template(operation=None):
    book = Workbook()
    sheet = book.active
    sheet.title = "Sheet"
    sheet["H1"] = "subtotal"
    for number in range(2, 5):
        sheet[f"H{number}"] = f"=SUM(D{number}:G{number})"
        sheet[f"I{number}"].fill = PatternFill("solid", fgColor="FFFF00")
    output = BytesIO()
    book.save(output)
    book.close()

    def patch(xml):
        for number in range(2, 5):
            cell = xml.find(f".//{{*}}c[@r='H{number}']")
            formula = cell.find("{*}f")
            formula.set("t", "shared")
            formula.set("si", "0")
            if number == 2:
                formula.set("ref", "H2:H4")
            else:
                formula.text = None
            cell.find("{*}v").text = "0"
        if operation is not None:
            operation(xml)

    return replace_xml(output.getvalue(), "xl/worksheets/sheet1.xml", patch)


def missing(result):
    return {x["location"] for x in result["issues"] if x["code"] == "material_amount_missing"}


def add_literal(xml, address, value, *, hidden=False):
    row = xml.find(".//{*}row[@r='2']")
    if hidden:
        row.set("hidden", "1")
    cell = ET.SubElement(row, f"{{{MAIN}}}c", {"r": address})
    if isinstance(value, str):
        cell.set("t", "inlineStr")
        ET.SubElement(ET.SubElement(cell, f"{{{MAIN}}}is"), f"{{{MAIN}}}t").text = value
    else:
        ET.SubElement(cell, f"{{{MAIN}}}v").text = str(value)


def test_shared_sum_templates_keep_original_coverage_and_frozen_v1_behavior():
    raw = template()
    current = inspect_bytes(raw, spec())
    assert {item["code"] for item in current["issues"]} == {"material_no_business_columns"}
    assert current["items"] == []
    assert {(x["location"], x["value"]) for x in current["coverage"]} == {
        ("Sheet!H1", "subtotal"), ("Sheet!H2", "0"), ("Sheet!H3", "0"), ("Sheet!H4", "0"),
    }
    historical_spec = material_inspection_v1.Specification.model_validate_json(
        json.dumps(spec().model_dump(mode="json"))
    )
    historical = material_inspection_v1.inspect_bytes(raw, historical_spec)
    assert len(missing(historical)) == 12
    assert historical["coverage"] != current["coverage"]


@pytest.mark.parametrize("address,value,hidden", [
    ("A2", "actual employee", False), ("B2", "2026-07", False),
    ("D2", 0, False), ("D2", 10, True), ("A2", "hidden identity", True),
])
def test_actual_input_even_zero_or_hidden_still_requires_amount(address, value, hidden):
    parsed = inspect_bytes(
        template(lambda xml: add_literal(xml, address, value, hidden=hidden)), spec()
    )
    assert "Sheet!AC2" in missing(parsed)
    assert any(x["location"] == "Sheet!" + address for x in parsed["coverage"])


@pytest.mark.parametrize("formula", [
    "SUM(D1:G1)", "SUM(Other!D2:G2)", "SUM([external.xlsx]Other!D2:G2)",
    "SUM(D2:H2)", "SUM(H2:H2)", "1-1", "SUM(D2:G3)", "SUM(G2:D2)",
])
def test_unknown_cross_row_external_cyclic_or_other_formula_cannot_hide_missing(formula):
    def patch(xml):
        xml.find(".//{*}c[@r='H2']/{*}f").text = formula

    assert "Sheet!AC2" in missing(inspect_bytes(template(patch), spec()))


@pytest.mark.parametrize("cache,kind", [(None, "n"), ("", "n"), ("1", "n"),
                                       ("#VALUE!", "e"), ("0", "b"),
                                       ("¥0", "n"), ("0,0", "n"), ("0元", "n"),
                                       ("NaN", "n"), ("Infinity", "n")])
def test_formula_cache_must_be_present_exact_numeric_zero(cache, kind):
    def patch(xml):
        cell = xml.find(".//{*}c[@r='H2']")
        cell.set("t", kind)
        cell.find("{*}v").text = cache

    assert "Sheet!AC2" in missing(inspect_bytes(template(patch), spec()))


@pytest.mark.parametrize("broken", [
    "unknown_seed", "missing_ref", "outside_ref", "absolute_future_seed",
])
def test_shared_formula_requires_known_seed_and_exact_declared_range(broken):
    def patch(xml):
        seed = xml.find(".//{*}c[@r='H2']/{*}f")
        if broken == "unknown_seed":
            xml.find(".//{*}c[@r='H3']/{*}f").set("si", "99")
        elif broken == "missing_ref":
            seed.attrib.pop("ref")
        elif broken == "absolute_future_seed":
            # A cross-row seed cannot become eligible just because a later
            # follower happens to occupy the absolute referenced row.
            seed.text = "SUM(D$3:G$3)"
        else:
            seed.set("ref", "H2:H2")

    assert "Sheet!AC3" in missing(inspect_bytes(template(patch), spec()))


@pytest.mark.parametrize("control", ["total_row", "control_location"])
def test_declared_control_row_cannot_be_exempted_as_template(control):
    changes = ({"total_rows": {"Sheet": (2,)}} if control == "total_row" else {
        "controls": (ControlTotal(location="Sheet!H2", locations=("Sheet!H3",)),),
    })
    assert "Sheet!AC2" in missing(inspect_bytes(template(), spec(**changes)))


def test_multiple_same_row_sum_formulas_only_over_empty_inputs_are_safe():
    def patch(xml):
        row = xml.find(".//{*}row[@r='2']")
        cell = ET.SubElement(row, f"{{{MAIN}}}c", {"r": "B2"})
        ET.SubElement(cell, f"{{{MAIN}}}f").text = "SUM(J2:L2)"
        ET.SubElement(cell, f"{{{MAIN}}}v").text = "0.00"

    result = inspect_bytes(template(patch), spec())
    assert not missing(result)
    assert {item["code"] for item in result["issues"]} == {"material_no_business_columns"}


def test_csv_and_xls_literal_zero_rows_do_not_gain_xlsx_template_exemption():
    cols = (Column(column="A", role="context"), Column(column="B", role="amount"))
    csv_result = inspect_bytes(b"0,\n", Specification(
        format="csv", columns=cols, header_rows={"CSV": 0},
    ))
    assert "CSV!B1" in missing(csv_result)
    import xlwt

    book = xlwt.Workbook()
    sheet = book.add_sheet("Sheet")
    sheet.write(0, 0, 0)
    output = BytesIO()
    book.save(output)
    xls_result = inspect_bytes(output.getvalue(), Specification(
        format="xls", columns=cols, header_rows={"Sheet": 0},
    ))
    assert "Sheet!B1" in missing(xls_result)
