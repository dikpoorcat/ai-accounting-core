"""An explicit nonmonetary disposition covers its original cell, not unknown money."""

import csv
import json
from datetime import date
from io import BytesIO, StringIO
from types import SimpleNamespace

import pytest
from openpyxl import Workbook
from test_materials import company as company  # noqa: F401

from ai_accounting.kernel import materials
from ai_accounting.kernel.inspection_cache import InspectionCache
from ai_accounting.kernel.materials import Materials, Specification, inspect_bytes
from ai_accounting.kernel.types import YearMonth


def original(format, values, *, hidden=False):
    rows = [["行号", "申报栏", "业务金额", "日期", "所属期"]]
    rows.extend([index, value, 1, "2026-01-31", "2026-01"]
                for index, value in enumerate(values, 1))
    spec = {
        "format": format,
        "columns": [
            {"column": "A", "role": "context"},
            {"column": "B", "role": "amount"},
            {"column": "C", "role": "amount"},
            {"column": "D", "role": "context"},
            {"column": "E", "role": "recognition_period"},
        ],
    }
    if format == "csv":
        stream = StringIO(newline="")
        csv.writer(stream).writerows(rows)
        return stream.getvalue().encode(), spec, "CSV"
    stream = BytesIO()
    if format == "xls":
        import xlwt

        book = xlwt.Workbook()
        sheet = book.add_sheet("申报")
        for row, cells in enumerate(rows):
            for column, value in enumerate(cells):
                if value is not None:
                    if isinstance(value, str) and value.startswith("="):
                        value = xlwt.Formula(value[1:])
                    style = xlwt.easyxf(num_format_str="YYYY-MM-DD") if isinstance(
                        value, date,
                    ) else xlwt.Style.default_style
                    sheet.write(row, column, value, style)
        if hidden:
            sheet.row(1).hidden = True
    else:
        book = Workbook()
        sheet = book.active
        sheet.title = "申报"
        for cells in rows:
            sheet.append(cells)
        if hidden:
            sheet.row_dimensions[2].hidden = True
    book.save(stream)
    if format == "xlsx":
        book.close()
    return stream.getvalue(), spec, "申报"


def receive(company, format="csv", values=("本年累计金额",), **kwargs):
    raw, spec, sheet = original(format, values, **kwargs)
    source, proof = company.source(raw, spec=spec)
    for row in range(2, len(values) + 2):
        company.resolve(
            source, f"{sheet}!C{row}", treatment="no_accounting",
            non_accounting_reason="not_company_business", reason="合成原件中的对照金额",
        )
    return source, proof, raw, spec, sheet


def dispose(company, source, location, **changes):
    return company.resolve(
        source, location,
        **{
            "treatment": "no_accounting",
            "non_accounting_reason": "non_monetary_content",
            "amount_fen": None,
            "reason": "已核对原件，此格为非金额文字或签章空位",
            **changes,
        },
    )


def found(company):
    return {
        (issue["code"], issue.get("location"))
        for issue in company.materials.check("2026-01")["issues"]
    }


def parse(raw, spec):
    return inspect_bytes(raw, Specification.model_validate_json(json.dumps(spec)))


@pytest.mark.parametrize("format", ["xls", "xlsx", "csv"])
def test_nonmonetary_cells_need_explicit_disposition_even_with_numeric_neighbours(company, format):
    values = ("否", "本年累计金额", "2026年1月31日", "7=4-5", "", None, " ", "负责人签章")
    source, _, raw, spec, sheet = receive(company, format, values)
    targets = {f"{sheet}!B{row}" for row in range(2, len(values) + 2)}
    inspection = parse(raw, spec)
    assert targets <= inspection.non_monetary_candidate_locations
    assert {("material_amount_missing", location) for location in targets} <= found(company)
    for location in sorted(targets):
        dispose(company, source, location)
    complete = company.materials.check("2026-01")
    assert complete["status"] == "complete" and complete["issues"] == []
    with company.engine.store.connection(read_only=True) as connection:
        stored = connection.execute(
            "SELECT r.amount_fen FROM fact_material_resolution_v2 r "
            "JOIN fact_current c ON c.fact_id=r.revision_id WHERE r.source_id=? "
            "AND r.non_accounting_reason='non_monetary_content'",
            (source["subject_id"],),
        ).fetchall()
    assert len(stored) == len(targets) and all(row[0] is None for row in stored)


@pytest.mark.parametrize(
    "changes,other_issue",
    [
        ({"non_accounting_reason": "not_company_business"}, None),
        ({"non_accounting_reason": None}, "material_non_accounting_basis"),
        ({"reason": None}, "material_non_accounting_basis"),
        ({"reason": " \t"}, "material_non_accounting_basis"),
        ({"treatment": "supporting"}, "material_supporting_conflict"),
    ],
)
def test_nonmonetary_provenance_requires_exact_explicit_null_disposition(
    company, changes, other_issue,
):
    source, _, _, _, sheet = receive(company)
    location = f"{sheet}!B2"
    dispose(company, source, location, **changes)
    issues = found(company)
    assert ("material_amount_missing", location) in issues
    if other_issue:
        assert (other_issue, location) in issues


def test_old_not_company_business_cannot_hide_literal_unknown_money(company):
    source, _, _, _, sheet = receive(company, values=("待核对",))
    location = f"{sheet}!B2"
    dispose(company, source, location, non_accounting_reason="not_company_business")
    assert ("material_amount_missing", location) in found(company)


@pytest.mark.parametrize("value,amount", [("本年累计金额", 0), ("12.00", None), ("0", None)])
def test_nonmonetary_reason_cannot_discard_readable_or_supplied_money(company, value, amount):
    source, _, _, _, sheet = receive(company, values=(value,))
    location = f"{sheet}!B2"
    dispose(company, source, location, amount_fen=amount)
    assert ("material_non_monetary_conflict", location) in found(company)


def test_nonmonetary_disposition_cannot_allocate_accounting_links(company):
    source, _, _, _, sheet = receive(company)
    location = f"{sheet}!B2"
    dispose(company, source, location, links=[company.expense("linked", 100)])
    assert {
        ("material_amount_missing", location), ("material_unexpected_links", location),
    } <= found(company)


def test_nonmonetary_disposition_must_be_unique(company):
    source, _, _, _, sheet = receive(company)
    location = f"{sheet}!B2"
    dispose(company, source, location)
    dispose(company, source, location, subject="second-disposition")
    assert {
        ("material_amount_missing", location), ("material_item_unresolved", location),
    } <= found(company)


def test_nonmonetary_disposition_does_not_survive_a_new_source_version(company):
    source, proof, _, spec, sheet = receive(company)
    location = f"{sheet}!B2"
    dispose(company, source, location)
    company.materials.receive(
        source["subject_id"],
        {"period": "2026-01", "evidence_digest": proof, "category": "transactions",
         "purpose": "business", "specification": spec},
        evidence=(proof, company.proof), expected_revision=1, request_id=company.request(),
    )
    assert {
        ("material_amount_missing", location), ("material_source_changed", location),
    } <= found(company)


@pytest.mark.parametrize("format", ["xls", "xlsx", "csv"])
def test_invalid_numeric_amounts_are_not_nonmonetary_candidates(company, format):
    values = ("1.001", "NaN", "Inf", "-Infinity")
    source, _, raw, spec, sheet = receive(company, format, values)
    targets = {f"{sheet}!B{row}" for row in range(2, len(values) + 2)}
    inspection = parse(raw, spec)
    assert not targets & inspection.non_monetary_candidate_locations
    for location in targets:
        dispose(company, source, location)
    assert {("material_amount_missing", location) for location in targets} <= found(company)


@pytest.mark.parametrize("format", ["xls", "xlsx"])
@pytest.mark.parametrize("value", ["本年累计金额", None])
def test_hidden_nonmonetary_cells_still_require_amount_review(company, format, value):
    source, _, raw, spec, sheet = receive(company, format, values=(value,), hidden=True)
    location = f"{sheet}!B2"
    inspection = parse(raw, spec)
    assert location not in inspection.non_monetary_candidate_locations
    dispose(company, source, location)
    assert ("material_amount_missing", location) in found(company)


@pytest.mark.parametrize("format", ["xls", "xlsx"])
def test_legacy_footer_exception_does_not_hide_an_absent_amount_in_hidden_column(company, format):
    stream = BytesIO()
    if format == "xls":
        import xlwt

        book = xlwt.Workbook()
        sheet = book.add_sheet("申报")
        sheet.write(0, 0, "说明")
        sheet.write(0, 1, "金额")
        sheet.write(1, 0, "负责人签章")
        sheet.col(1).hidden = True
    else:
        book = Workbook()
        sheet = book.active
        sheet.title = "申报"
        sheet.append(["说明", "金额"])
        sheet.append(["负责人签章"])
        sheet.column_dimensions["B"].hidden = True
    book.save(stream)
    if format == "xlsx":
        book.close()
    spec = {"format": format, "columns": [
        {"column": "A", "role": "context"}, {"column": "B", "role": "amount"},
    ]}
    source, _ = company.source(stream.getvalue(), spec=spec)
    dispose(company, source, "申报!B2", non_accounting_reason="not_company_business")
    assert ("material_amount_missing", "申报!B2") in found(company)


@pytest.mark.parametrize("format", ["xls", "xlsx"])
def test_formula_blank_cache_is_not_nonmonetary_content(company, format):
    source, _, raw, spec, sheet = receive(company, format, values=("=1+1",))
    location = f"{sheet}!B2"
    inspection = parse(raw, spec)
    assert location not in inspection.non_monetary_candidate_locations
    dispose(company, source, location)
    assert ("material_amount_missing", location) in found(company)


@pytest.mark.parametrize("format", ["xls", "xlsx"])
def test_native_date_and_boolean_are_not_literal_nonmonetary_text(format):
    raw, spec, sheet = original(format, [date(2026, 1, 31), True])
    inspection = parse(raw, spec)
    assert not {f"{sheet}!B2", f"{sheet}!B3"} & inspection.non_monetary_candidate_locations


@pytest.mark.parametrize("grouped", [True, False])
def test_nonmonetary_disposition_never_skips_control_total(company, grouped):
    raw, spec, sheet = original("csv", ["本年累计金额", 1])
    location = f"{sheet}!B2"
    if grouped:
        spec["controls"] = [{"location": location, "locations": [f"{sheet}!B3"]}]
    else:
        spec["total_rows"] = {sheet: [2]}
    assert location not in parse(raw, spec).non_monetary_candidate_locations
    source, _ = company.source(raw, spec=spec)
    dispose(company, source, location)
    for row in (2, 3):
        company.resolve(
            source, f"{sheet}!C{row}", treatment="no_accounting",
            non_accounting_reason="not_company_business", reason="合成对照金额",
        )
    company.resolve(
        source, f"{sheet}!B3", treatment="no_accounting",
        non_accounting_reason="not_company_business", reason="合成对照金额",
    )
    assert {
        ("material_amount_missing", location), ("material_total_mismatch", location),
    } <= found(company)


@pytest.mark.parametrize("suffix", [b"\x06\x00\x04\x00\x00", b""])
def test_xls_formula_provenance_rejects_truncated_worksheet_records(suffix):
    stream = b"\x09\x08\x00\x00" + suffix
    book = SimpleNamespace(mem=stream, _sh_abs_posn=[0], base=0, stream_len=len(stream))
    with pytest.raises(ValueError, match="BIFF"):
        materials._xls_formula_locations(book, 0)


def test_xls_without_retained_formula_stream_does_not_claim_empty_provenance():
    assert materials._xls_formula_locations(SimpleNamespace(mem=None), 0) is None


def test_private_nonmonetary_metadata_is_reparsed_for_cold_completeness(
    company, tmp_path, monkeypatch,
):
    source, proof, _, spec, sheet = receive(company, "xlsx", ["本年累计金额", ""])
    for row in (2, 3):
        dispose(company, source, f"{sheet}!B{row}")
    calls = []
    original_inspect = materials.inspect_bytes

    def counted(raw, specification):
        calls.append(True)
        return original_inspect(raw, specification)

    monkeypatch.setattr(materials, "inspect_bytes", counted)
    reader = Materials(company.engine, inspection_cache=InspectionCache(
        tmp_path, build_id="nonmonetary-test",
    ))
    public = reader.inspect(proof, spec)
    assert "non_monetary_candidate_locations" not in json.dumps(public)
    restarted = Materials(company.engine, inspection_cache=InspectionCache(
        tmp_path, build_id="nonmonetary-test",
    ))
    assert restarted.inspect(proof, spec) == public and len(calls) == 1
    for path in (tmp_path / ".inspection-cache").glob("*.json"):
        assert "non_monetary_candidate_locations" not in path.read_text(encoding="utf-8")
    with company.engine.store.connection(read_only=True) as connection:
        cache = materials._CompletenessInspectionCache(connection)
        month = YearMonth("2026-01").ordinal
        first = materials.check_completeness_many(
            connection, (month,), company.engine.store.registry, _inspection_cache=cache,
        )[month]
        assert len(calls) == 2 and first["issues"] == []
        second = materials.check_completeness_many(
            connection, (month,), company.engine.store.registry, _inspection_cache=cache,
        )[month]
        assert second == first and len(calls) == 2
    assert company.materials.check("2026-01")["issues"] == [] and len(calls) == 3
