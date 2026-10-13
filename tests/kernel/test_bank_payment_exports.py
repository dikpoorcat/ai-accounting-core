"""Bank instructions keep distinct typed debts and exact unpaid amounts."""

import json
from pathlib import Path

import pytest
import test_opening_continuation as opening
from entity_fixture import save_entity_display_profile
from openpyxl import load_workbook
from test_exports import evidence, inventory, queue, template_bytes
from test_integrity_content import damage
from test_payroll_corrections import Company

from ai_accounting.kernel.contracts import KernelError, NeedsInformation
from ai_accounting.kernel.domains.adjustments import EmployeeAdvance, PassThroughReturn
from ai_accounting.kernel.domains.transactions import Allocation, Expense, PassThrough, Payment
from ai_accounting.kernel.exports import MANIFEST_NAME, Exports, run_export_jobs
from ai_accounting.kernel.identity_corrections import IdentityCorrections
from ai_accounting.kernel.periods import Periods

opening_book = opening.book


@pytest.fixture
def bank_company(tmp_path):
    company = Company(tmp_path / "bank-instructions.sqlite")
    export = Exports(company.engine)
    template = evidence(company, template_bytes())
    return company, export, template


def payee(engine, export, party, proof, request):
    export.save_payee(
        party, name=party, account="000123456789", evidence_digest=proof,
        expected_revision=0, request_id=request,
    )


def entrusted(company, amount=30_000):
    save_entity_display_profile(
        company.engine,
        {"kind": "person", "entity_id": "beneficiary", "display_name": "合成最终收款人",
         "source": "本次合成受托付款确认", "evidence_digest": company.owner_confirmation},
        expected_revision=0, request_id=company.request(),
    )
    company.save(PassThrough(
        period="2026-01", payer_id="payer", beneficiary_id="beneficiary",
        amount_fen=amount, rights_and_obligation_confirmed=True,
    ), "entrusted")
    company.publish("entrusted")


def instruction_sources(plan):
    return {source["subject_id"]: source for row in plan["rows"] for source in row["sources"]}


def test_reimbursement_owner_advance_and_entrusted_payment_share_one_file(bank_company, tmp_path):
    company, export, template = bank_company
    entrusted(company)
    for subject, creditor, party, amount in (
        ("reimbursement", "employee", "employee", 12_000),
        ("individual-expense", "individual", "individual", 7_000),
        ("supplier-expense", "supplier", "supplier", 9_000),
    ):
        company.save(Expense(
            period="2026-01", counterparty_id=party, creditor_kind=creditor,
            expense_class="administration", amount_fen=amount,
        ), subject)
        company.publish(subject)
    company.save(EmployeeAdvance(
        period="2026-01", payer_id="owner", payer_kind="owner",
        payment_on_behalf_confirmed=True, actual_creditor_payment_date="2026-01-18",
        sources=(Allocation(source_kind="expense", source_id="supplier-expense",
                            obligation="primary", amount_fen=9_000),),
    ), "owner-advance")
    company.publish("owner-advance")
    for party in ("employee", "individual", "owner", "beneficiary"):
        payee(company.engine, export, party, company.owner_confirmation, company.request())
    selected = ["entrusted", "reimbursement", "individual-expense", "owner-advance"]
    before = {table: company.count(table) for table in ("voucher_version", "fact_revision")}
    plan, job, _ = queue(company, export, template, tmp_path / "mixed", source_ids=selected)
    sources = instruction_sources(plan)
    assert plan["scope"] == "selected" and plan["total_fen"] == 58_000
    assert {subject: (source["category"], source["amount_fen"], source["party_id"])
            for subject, source in sources.items()} == {
        "entrusted": ("代付款", 30_000, "beneficiary"),
        "reimbursement": ("报销", 12_000, "employee"),
        "individual-expense": ("报销", 7_000, "individual"),
        "owner-advance": ("报销", 9_000, "owner"),
    }
    assert sources["entrusted"]["obligation"] == "pass_through:entrusted:remittance"
    completed = run_export_jobs(company.engine)[0]
    assert completed["status"] == "succeeded" and completed["job_id"] == job["job_id"]
    workbook = load_workbook(completed["result"]["path"])
    try:
        assert workbook.active.max_row == 5
        assert {row[0].value: (row[1].value, row[2].value, row[3].value)
                for row in list(workbook.active.rows)[1:]} == {
            "beneficiary": ("000123456789", "300.00", "2026-01 代付款"),
            "employee": ("000123456789", "120.00", "2026-01 报销"),
            "individual": ("000123456789", "70.00", "2026-01 报销"),
            "owner": ("000123456789", "90.00", "2026-01 报销"),
        }
    finally:
        workbook.close()
    assert {table: company.count(table) for table in before} == before
    with pytest.raises(KernelError) as whole_month:
        export.preview("2026-01", template_evidence_digest=template)
    assert whole_month.value.code == "materials_incomplete"


def test_entrusted_refund_and_partial_payment_reduce_only_exact_debts(bank_company):
    company, export, template = bank_company
    entrusted(company)
    company.save(Payment(
        period="2026-01", actual_date="2026-01-10", direction="inflow",
        bank_account_id="bank", counterparty_id="payer", amount_fen=30_000,
        allocations=(Allocation(source_kind="pass_through", source_id="entrusted",
                                obligation="collection", amount_fen=30_000),),
    ), "collection")
    company.publish("collection")
    company.save(PassThroughReturn(
        period="2026-01", source_id="entrusted", amount_fen=8_000, refund_right_confirmed=True,
    ), "return")
    company.publish("return")
    company.save(Payment(
        period="2026-01", actual_date="2026-01-20", direction="outflow",
        bank_account_id="bank", counterparty_id="beneficiary", amount_fen=5_000,
        allocations=(Allocation(source_kind="pass_through", source_id="entrusted",
                                obligation="remittance", amount_fen=5_000),),
    ), "partial")
    company.publish("partial")
    for party in ("payer", "beneficiary"):
        payee(company.engine, export, party, company.owner_confirmation, company.request())
    plan = export.preview("2026-01", template_evidence_digest=template,
                          source_ids=["entrusted", "return"])
    assert {(row["party_id"], row["category"]): row["amount_fen"] for row in plan["rows"]} == {
        ("beneficiary", "代付款"): 17_000,
        ("payer", "代收款退回"): 8_000,
    }
    company.save(Payment(
        period="2026-01", actual_date="2026-01-22", direction="outflow",
        bank_account_id="bank", counterparty_id="payer", amount_fen=8_000,
        allocations=(Allocation(source_kind="pass_through_return", source_id="return",
                                obligation="primary", amount_fen=8_000),),
    ), "returned")
    company.publish("returned")
    after = export.preview("2026-01", template_evidence_digest=template,
                           source_ids=["entrusted", "return"])
    assert after["total_fen"] == 17_000
    assert set(instruction_sources(after)) == {"entrusted"}


def test_custom_memo_and_new_source_do_not_rewrite_a_frozen_file(bank_company, tmp_path):
    company, export, template = bank_company
    entrusted(company)
    payee(company.engine, export, "beneficiary", company.owner_confirmation, company.request())
    first, job, options = queue(company, export, template, tmp_path / "frozen",
                                source_ids=["entrusted"])
    Periods(company.engine).management(
        "entrusted", payment_period="2026-01", payment_category="受托转付款",
        expected_revision=0, request_id=company.request(),
    )
    company.save(Expense(period="2026-01", counterparty_id="beneficiary", amount_fen=2_000,
                         creditor_kind="individual", expense_class="administration"), "new-claim")
    company.publish("new-claim")
    assert export.confirm("2026-01", **options) == job
    new, newer_job, _ = queue(company, export, template, tmp_path / "frozen",
                             source_ids=["entrusted", "new-claim"])
    assert new["total_fen"] == 32_000 and newer_job["job_id"] != job["job_id"]
    assert instruction_sources(new)["entrusted"]["category"] == "受托转付款"
    completed = run_export_jobs(company.engine)
    manifests = {}
    for result in completed:
        assert result["status"] == "succeeded"
        manifests[result["job_id"]] = json.loads(
            (Path(result["result"]["directory"]) / MANIFEST_NAME).read_text(encoding="utf-8")
        )["plan"]
    assert manifests[job["job_id"]] == first
    assert manifests[newer_job["job_id"]] == new


def test_entrusted_export_requires_actual_payee_not_collection_payer(bank_company):
    company, export, template = bank_company
    entrusted(company)
    payee(company.engine, export, "payer", company.owner_confirmation, company.request())
    with pytest.raises(NeedsInformation) as missing:
        export.preview("2026-01", template_evidence_digest=template, source_ids=["entrusted"])
    assert missing.value.issues[0]["field"] == "payee"
    assert missing.value.issues[0]["reusable_sources"] == ["beneficiary"]


def test_opening_reimbursements_and_net_wage_export_but_tax_and_supplier_do_not(opening_book):
    engine, _, _, package, proof = opening_book
    members = [
        ("opening_obligation", "employee-claim", {
            "counterparty_id": "employee", "nature": "employee_reimbursement",
            "outstanding_fen": 11_000, "business_reference": "prior-employee-claim",
        }),
        ("opening_obligation", "owner-claim", {
            "counterparty_id": "owner", "nature": "owner_reimbursement",
            "outstanding_fen": 12_000, "business_reference": "prior-owner-claim",
        }),
        ("opening_obligation", "supplier-claim", {
            "counterparty_id": "supplier", "nature": "supplier_administration_payable",
            "outstanding_fen": 5_000, "business_reference": "prior-supplier-claim",
        }),
        ("opening_payroll_payable", "net-wage", {
            "employee_id": "employee", "recipient_id": "employee", "payroll_period": "2025-12",
            "component": "net", "outstanding_fen": 13_000,
        }),
        ("opening_payroll_payable", "tax", {
            "employee_id": "employee", "recipient_id": "tax-authority",
            "payroll_period": "2025-12", "component": "withheld_tax", "outstanding_fen": 1_400,
        }),
        ("opening_bank", "opening-cash", {"bank_account_id": "bank", "balance_fen": 42_400}),
    ]
    package(members)
    export = Exports(engine)
    for index, party in enumerate(("employee", "owner")):
        payee(engine, export, party, proof, "payee-" + str(index))
    template = engine.register_evidence(template_bytes(), "application/octet-stream", "template",
                                       request_id="template")["digest"]
    plan = export.preview("2026-01", template_evidence_digest=template,
                          source_ids=["employee-claim", "owner-claim", "net-wage"])
    assert plan["total_fen"] == 36_000
    assert {subject: (source["category"], source["amount_fen"])
            for subject, source in instruction_sources(plan).items()} == {
        "employee-claim": ("报销", 11_000), "owner-claim": ("报销", 12_000),
        "net-wage": ("工资", 13_000),
    }
    for unsupported in ("supplier-claim", "tax"):
        with pytest.raises(KernelError) as rejected:
            export.preview("2026-01", template_evidence_digest=template, source_ids=[unsupported])
        assert rejected.value.code == "invalid_export_sources"


def test_complete_scope_includes_entrusted_payment_once_month_is_complete(bank_company):
    company, export, template = bank_company
    entrusted(company)
    payee(company.engine, export, "beneficiary", company.owner_confirmation, company.request())
    inventory(company)
    plan = export.preview("2026-01", template_evidence_digest=template)
    assert plan["scope"] == "complete" and plan["total_fen"] == 30_000
    assert instruction_sources(plan)["entrusted"]["category"] == "代付款"


def test_export_rejects_a_current_head_pointing_to_another_valid_source(bank_company):
    company, export, template = bank_company
    for subject, amount in (("first", 1000), ("second", 2000)):
        company.save(Expense(period="2026-01", counterparty_id="employee", amount_fen=amount,
                             creditor_kind="employee", expense_class="administration"), subject)
        company.publish(subject)
    damage(company.engine, "calculation_current",
           "DELETE FROM calculation_current WHERE subject_id='second'")
    damage(company.engine, "calculation_current",
           "UPDATE calculation_current SET calculation_id=(SELECT id FROM calculation "
           "WHERE subject_id='second') WHERE subject_id='first'")
    with pytest.raises(KernelError) as invalid:
        export.preview("2026-01", template_evidence_digest=template, source_ids=["first"])
    assert invalid.value.code == "content_integrity_failed"


def test_identity_superseded_source_cannot_reenter_selected_or_complete_export(bank_company):
    company, export, template = bank_company
    for subject, amount in (("duplicate", 1000), ("retained", 1200)):
        company.save(Expense(period="2026-01", counterparty_id="employee", amount_fen=amount,
                             creditor_kind="employee", expense_class="administration"), subject)
        company.publish(subject)
    payee(company.engine, export, "employee", company.owner_confirmation, company.request())
    corrections = IdentityCorrections(company.engine)
    options = {
        "changes": [{"subject_id": "duplicate", "expected_revision": 1,
                     "action": "supersede", "replacement_subject_id": "retained"}],
        "evidence": [company.owner_confirmation], "reason": "合成确认两笔是同一事项，保留准确来源",
    }
    preview = corrections.preview_identity_correction(**options)
    corrections.confirm_identity_correction(
        **options, preview_digest=preview["digest"], epochs=preview["epochs"],
        request_id=company.request(),
    )
    with pytest.raises(KernelError) as retired:
        export.preview("2026-01", template_evidence_digest=template, source_ids=["duplicate"])
    assert retired.value.code == "invalid_export_sources"
    inventory(company)
    plan = export.preview("2026-01", template_evidence_digest=template)
    assert plan["total_fen"] == 1200 and set(instruction_sources(plan)) == {"retained"}
