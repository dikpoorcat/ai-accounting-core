"""Frozen exports use real published balances and recover file-job interruptions."""

from __future__ import annotations

from io import BytesIO

import pytest
from openpyxl import Workbook, load_workbook
from test_payroll import (
    actual,
    contribution_policy,
    income_tax_policy,
    labor,
    labor_policy,
    opening,
    payroll,
    profile,
)
from test_payroll_corrections import Company

from ai_accounting.kernel.contracts import KernelError, NeedsInformation
from ai_accounting.kernel.domains.adjustments import EmployeeAdvance
from ai_accounting.kernel.domains.transactions import Allocation, Expense, Payment
from ai_accounting.kernel.exports import WORKBOOK_NAME, Exports, run_export_jobs
from ai_accounting.kernel.periods import MATERIAL_CATEGORIES, Periods


def template_bytes(*, formula=False):
    book = Workbook()
    sheet = book.active
    sheet.append(["收款方名称", "收款方账号", "金额", "附言/用途"])
    if formula:
        sheet.append(["=SUM(1,2)", "00123456", "1.00", "工资"])
    buffer = BytesIO()
    book.save(buffer)
    book.close()
    return buffer.getvalue()


def inventory(company, period="2026-01"):
    for category in MATERIAL_CATEGORIES:
        evidence = sorted({ev for month, ev in company.materials[category] if month <= period})
        Periods(company.engine).inventory(
            period,
            category,
            evidence=evidence,
            expected=len(evidence),
            no_business=not evidence,
            confirmation_evidence=company.owner_confirmation,
            request_id=company.request(),
        )


def evidence(company, content, name="template"):
    return company.engine.register_evidence(
        content,
        "application/octet-stream",
        name,
        request_id=company.request(),
    )["digest"]


@pytest.fixture
def setup(tmp_path):
    company = Company(tmp_path / "exports.sqlite")
    for fact, subject in (
        (profile(), "profile"),
        (contribution_policy(), "contributions"),
        (income_tax_policy(), "income-tax"),
        (opening(), "opening"),
        (payroll(), "january"),
    ):
        company.save(fact, subject)
    company.publish("january")
    export = Exports(company.engine)
    export.save_payee(
        "employee",
        name="张三",
        account="001234567890",
        evidence_digest=company.owner_confirmation,
        expected_revision=0,
        request_id=company.request(),
    )
    template = evidence(company, template_bytes())
    inventory(company)
    return company, export, template


def queue(
    company, export, template, directory, *, source_ids=None, period="2026-01", request_id=None
):
    preview = export.preview(period, template_evidence_digest=template, source_ids=source_ids)
    options = {
        "template_evidence_digest": template,
        "source_ids": source_ids,
        "preview_digest": preview["digest"],
        "epochs": preview["epochs"],
        "output_directory": str(directory),
        "request_id": request_id or company.request(),
    }
    return preview, export.confirm(period, **options), options


def test_personal_advances_export_only_explicit_employee_reimbursement(setup):
    company, export, template = setup
    for party, payer_kind, amount in (("employee", "employee", 12000), ("owner", "owner", 9000)):
        source = "source-" + party
        company.save(
            Expense(
                period="2026-01",
                counterparty_id="supplier",
                amount_fen=amount,
                expense_class="administration",
                creditor_kind="supplier",
            ),
            source,
        )
        company.publish(source)
        company.save(
            EmployeeAdvance(
                period="2026-01",
                payer_id=party,
                payer_kind=payer_kind,
                payment_on_behalf_confirmed=True,
                actual_creditor_payment_date="2026-01-20",
                sources=(
                    Allocation(
                        source_kind="expense",
                        source_id=source,
                        obligation="primary",
                        amount_fen=amount,
                    ),
                ),
            ),
            "advance-" + party,
        )
        company.publish("advance-" + party)
    inventory(company)
    preview = export.preview("2026-01", template_evidence_digest=template)
    sources = {source["subject_id"]: source for row in preview["rows"] for source in row["sources"]}
    assert "advance-employee" in sources
    assert "advance-owner" not in sources
    assert sources["advance-employee"]["amount_fen"] == 12000
    assert sources["advance-employee"]["party_id"] == "employee"


def test_complete_export_uses_salary_labor_and_explicit_employee_reimbursement(setup):
    company, export, template = setup
    company.save(labor_policy(), "labor-policy")
    company.save(labor(), "labor")
    company.save(
        Expense(
            period="2026-01",
            counterparty_id="employee",
            amount_fen=12_000,
            expense_class="administration",
            creditor_kind="employee",
        ),
        "reimbursement",
    )
    company.save(
        Expense(
            period="2026-01",
            counterparty_id="supplier",
            amount_fen=50_000,
            expense_class="administration",
            creditor_kind="supplier",
        ),
        "supplier-expense",
    )
    company.publish("labor", "reimbursement", "supplier-expense")
    export.save_payee(
        "contractor",
        name="李四",
        account="000987654321",
        evidence_digest=company.owner_confirmation,
        expected_revision=0,
        request_id=company.request(),
    )
    inventory(company)
    preview = export.preview("2026-01", template_evidence_digest=template)
    assert preview["scope"] == "complete"
    assert {row["category"]: row["amount_fen"] for row in preview["rows"]} == {
        "工资": 907_400,
        "劳务": 840_000,
        "报销": 12_000,
    }
    assert preview["total_fen"] == 1_759_400
    assert {source["subject_id"] for row in preview["rows"] for source in row["sources"]} == {
        "january",
        "labor",
        "reimbursement",
    }
    selected = export.preview("2026-01", template_evidence_digest=template, source_ids=["january"])
    assert selected["scope"] == "selected" and selected["total_fen"] == 907_400
    with pytest.raises(KernelError) as invalid:
        export.preview(
            "2026-01", template_evidence_digest=template, source_ids=["supplier-expense"]
        )
    assert invalid.value.code == "invalid_export_sources"


def test_partial_actual_payment_reduces_amount_and_generating_file_records_no_payment(
    setup, tmp_path
):
    company, export, template = setup
    company.save(
        Payment(
            period="2026-02",
            actual_date="2026-02-10",
            direction="outflow",
            bank_account_id="bank",
            counterparty_id="employee",
            amount_fen=400_000,
            allocations=(
                Allocation(
                    source_kind="payroll", source_id="january", obligation="net", amount_fen=400_000
                ),
            ),
        ),
        "partial-payment",
    )
    company.publish("partial-payment")
    before = {
        table: company.count(table)
        for table in ("fact_revision", "calculation", "voucher_version", "voucher_line")
    }
    period = company.engine.overview("2026-02")
    preview, job, _ = queue(company, export, template, tmp_path / "out")
    assert preview["total_fen"] == 507_400
    assert not (tmp_path / "out").exists()  # Confirmation only freezes a durable job.
    outcomes = run_export_jobs(company.engine)
    assert outcomes[0]["job_id"] == job["job_id"]
    assert outcomes[0]["status"] == "succeeded"
    assert {table: company.count(table) for table in before} == before
    assert company.engine.overview("2026-02") == period
    book = load_workbook(tmp_path / "out" / WORKBOOK_NAME, data_only=False)
    try:
        assert book.active.max_column == 4
        assert [cell.value for cell in book.active[2]] == [
            "张三",
            "001234567890",
            "5074.00",
            "2026-01 工资",
        ]
        assert all(cell.data_type == "s" for cell in book.active[2])
    finally:
        book.close()


@pytest.mark.parametrize("lane", ["accounting", "material", "management"])
def test_any_lane_change_expires_unconfirmed_export(setup, tmp_path, lane):
    company, export, template = setup
    preview = export.preview("2026-01", template_evidence_digest=template)
    if lane == "accounting":
        company.save(actual(), "actual")
        company.publish("actual")
    elif lane == "material":
        inventory(company)
    else:
        export.save_payee(
            "employee",
            name="张三",
            account="000000000002",
            evidence_digest=company.owner_confirmation,
            expected_revision=1,
            request_id=company.request(),
        )
    with pytest.raises(KernelError) as error:
        export.confirm(
            "2026-01",
            template_evidence_digest=template,
            preview_digest=preview["digest"],
            epochs=preview["epochs"],
            output_directory=str(tmp_path / "expired"),
            request_id=company.request(),
        )
    assert error.value.code == "preview_expired"
    assert company.count("jobs") == 0
    assert not (tmp_path / "expired").exists()


def test_frozen_job_uses_approved_payee_and_does_not_invalidate_other_previews(setup, tmp_path):
    company, export, template = setup
    with company.engine.store.connection(read_only=True) as connection:
        epochs = company.engine.store.epochs(connection)
    preview, job, arguments = queue(company, export, template, tmp_path / "frozen")
    assert export.confirm("2026-01", **arguments) == job
    with company.engine.store.connection(read_only=True) as connection:
        assert company.engine.store.epochs(connection) == epochs
    export.save_payee(
        "employee",
        name="张三",
        account="000000000002",
        evidence_digest=company.owner_confirmation,
        expected_revision=1,
        request_id=company.request(),
    )
    run_export_jobs(company.engine)
    book = load_workbook(tmp_path / "frozen" / WORKBOOK_NAME)
    try:
        assert book.active.cell(2, 2).value == preview["rows"][0]["account"] == "001234567890"
    finally:
        book.close()
    assert export.confirm("2026-01", **arguments) == job
    assert company.count("jobs") == 1


def test_files_published_before_crash_are_verified_and_reused(setup, tmp_path):
    company, export, template = setup
    _, job, _ = queue(company, export, template, tmp_path / "crash")

    class ProcessCrash(BaseException):
        pass

    def crash(stage, identifier):
        if stage == "files_published":
            assert identifier == job["job_id"]
            raise ProcessCrash()

    with pytest.raises(ProcessCrash):
        run_export_jobs(company.engine, fault=crash)
    target = tmp_path / "crash" / WORKBOOK_NAME
    original = target.read_bytes()
    with company.engine.store.connection(read_only=True) as connection:
        assert (
            connection.execute("SELECT status FROM jobs WHERE id=?", (job["job_id"],)).fetchone()[0]
            == "running"
        )
    retried = run_export_jobs(company.engine)
    assert retried[0]["status"] == "succeeded"
    assert target.read_bytes() == original
    assert run_export_jobs(company.engine) == []
    with company.engine.store.connection(read_only=True) as connection:
        assert (
            connection.execute("SELECT attempts FROM jobs WHERE id=?", (job["job_id"],)).fetchone()[
                0
            ]
            == 2
        )


def test_file_failure_is_retryable_and_work_runs_outside_sql_write_transaction(setup, tmp_path):
    company, export, template = setup
    queue(company, export, template, tmp_path / "retry")

    def fail(stage, identifier):
        if stage == "before_files":
            # A separate writer succeeds while the slow worker owns only its OS job lock.
            company.engine.register_evidence(
                b"independent write", "text/plain", "other", request_id=company.request()
            )
            raise OSError("simulated file write failure")

    outcomes = run_export_jobs(company.engine, fault=fail)
    assert outcomes[0]["status"] == "failed"
    assert not (tmp_path / "retry").exists()
    assert run_export_jobs(company.engine)[0]["status"] == "succeeded"


def test_existing_unrelated_directory_is_never_overwritten(setup, tmp_path):
    company, export, template = setup
    destination = tmp_path / "occupied"
    destination.mkdir()
    sentinel = destination / "existing.txt"
    sentinel.write_text("owner document")
    queue(company, export, template, destination)
    result = run_export_jobs(company.engine)[0]
    assert result["status"] == "failed"
    assert sentinel.read_text() == "owner document"
    assert not (destination / WORKBOOK_NAME).exists()


def test_complete_export_rejects_missing_materials_pending_and_unpublished_facts(setup):
    company, export, template = setup
    Periods(company.engine).inventory(
        "2026-01",
        "bank",
        evidence=[],
        expected=1,
        no_business=False,
        confirmation_evidence=company.owner_confirmation,
        request_id=company.request(),
    )
    with pytest.raises(KernelError) as material:
        export.preview("2026-01", template_evidence_digest=template)
    assert material.value.code == "materials_incomplete"
    inventory(company)
    company.save(actual(), "actual")
    with pytest.raises(KernelError) as pending:
        export.preview("2026-01", template_evidence_digest=template)
    assert pending.value.code == "materials_incomplete"
    company.publish("actual")
    company.save(
        Expense(
            period="2026-01",
            counterparty_id="employee",
            amount_fen=100,
            expense_class="administration",
            creditor_kind="employee",
        ),
        "unpublished",
    )
    with pytest.raises(KernelError) as unpublished:
        export.preview("2026-01", template_evidence_digest=template)
    assert unpublished.value.code == "materials_incomplete"


def test_management_grouping_does_not_change_accounting_period_or_amount(setup):
    company, export, template = setup
    before = company.current("january")
    Periods(company.engine).management(
        "january",
        note=None,
        payment_period="2026-02",
        payment_category="次月工资",
        expected_revision=0,
        request_id=company.request(),
    )
    inventory(company, "2026-02")
    preview = export.preview("2026-02", template_evidence_digest=template)
    assert preview["rows"][0]["category"] == "次月工资"
    assert preview["rows"][0]["sources"][0]["source_period"] == "2026-01"
    assert preview["total_fen"] == 907_400
    assert company.current("january") == before
    assert set(preview["inventories"]) == {"2026-01", "2026-02"}


def test_template_and_payee_reject_formula_injection_and_caller_cannot_supply_amount(
    setup, tmp_path
):
    company, export, template = setup
    with pytest.raises(ValueError):
        export.save_payee(
            "employee",
            name="=WEBSERVICE(1)",
            account="00123",
            evidence_digest=company.owner_confirmation,
            expected_revision=1,
            request_id=company.request(),
        )
    malicious_template = evidence(company, template_bytes(formula=True))
    with pytest.raises(ValueError):
        export.preview("2026-01", template_evidence_digest=malicious_template)
    with pytest.raises(TypeError):
        export.preview("2026-01", template_evidence_digest=template, amount_fen=1)
    with pytest.raises(ValueError):
        export.preview("2026-01", template_evidence_digest=template, source_ids="january")


def test_missing_payee_is_a_management_information_gap_not_zero_payment(setup):
    company, export, template = setup
    company.save(labor_policy(), "labor-policy")
    company.save(labor(), "labor")
    company.publish("labor")
    inventory(company)
    with pytest.raises(NeedsInformation) as error:
        export.preview("2026-01", template_evidence_digest=template)
    assert error.value.response()["fact_issues"][0]["field"] == "payee"


def test_settled_source_does_not_block_other_unpaid_sources(setup):
    company, export, template = setup
    company.save(
        Expense(
            period="2026-01",
            counterparty_id="employee",
            amount_fen=12_000,
            expense_class="administration",
            creditor_kind="employee",
        ),
        "expense",
    )
    company.publish("expense")
    company.save(
        Payment(
            period="2026-02",
            actual_date="2026-02-10",
            direction="outflow",
            bank_account_id="bank",
            counterparty_id="employee",
            amount_fen=907_400,
            allocations=(
                Allocation(
                    source_kind="payroll", source_id="january", obligation="net", amount_fen=907_400
                ),
            ),
        ),
        "paid",
    )
    company.publish("paid")
    inventory(company)
    preview = export.preview("2026-01", template_evidence_digest=template)
    assert preview["total_fen"] == 12_000
    assert {source["subject_id"] for row in preview["rows"] for source in row["sources"]} == {
        "expense"
    }
