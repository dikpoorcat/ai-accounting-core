"""Row capacity does not bypass original hidden content or decimal values."""

from datetime import date
from decimal import localcontext
from io import BytesIO
from xml.etree import ElementTree
from zipfile import ZipFile

import pytest
from openpyxl import Workbook, load_workbook

from ai_accounting.kernel import materials
from ai_accounting.kernel.contracts import KernelError
from ai_accounting.kernel.materials import Column, Specification, inspect_bytes


def workbook_bytes(rows, *, hidden_column=False):
    book = Workbook()
    sheet = book.active
    sheet.title = "Sheet"
    for row in rows:
        sheet.append(row)
    if hidden_column:
        sheet.column_dimensions["B"].hidden = True
    output = BytesIO()
    book.save(output)
    book.close()
    return output.getvalue()


def replace_xml(raw, path, operation):
    result = BytesIO()
    with ZipFile(BytesIO(raw)) as source, ZipFile(result, "w") as target:
        for name in source.namelist():
            content = source.read(name)
            if name == path:
                xml = ElementTree.fromstring(content)
                operation(xml)
                content = ElementTree.tostring(xml)
            target.writestr(name, content)
    return result.getvalue()


def test_xlsx_exact_decimal_cached_formula_dates_and_hidden_column():
    raw = workbook_bytes([["date", "amount"], [date(2026, 3, 1), "=1+2"]], hidden_column=True)

    def patch(document):
        cell = document.find(".//{*}c[@r='B2']")
        cell.find("{*}v").text = "90071992547409.93"

    raw = replace_xml(raw, "xl/worksheets/sheet1.xml", patch)
    result = inspect_bytes(
        raw,
        Specification(
            format="xlsx",
            columns=(
                Column(column="A", role="recognition_period"),
                Column(column="B", role="amount"),
            ),
        ),
    )
    assert not result["issues"]
    assert result["items"][0]["amount_fen"] == 9_007_199_254_740_993
    assert result["items"][0]["recognition_period"] == "2026-03"
    assert result["items"][0]["hidden"] is True


@pytest.mark.parametrize("kind", ["csv", "xlsx", "xls"])
def test_data_rows_exclude_declared_headers_and_control_totals(kind, monkeypatch):
    monkeypatch.setattr(materials, "MAX_DATA_ROWS", 2)
    rows = [["reference", "amount"], ["one", "1.23"], ["two", "1.23"], ["total", "2.46"]]
    if kind == "csv":
        raw, sheet = "\n".join(",".join(row) for row in rows).encode(), "CSV"
    elif kind == "xlsx":
        raw, sheet = workbook_bytes(rows), "Sheet"
    else:
        import xlwt

        book = xlwt.Workbook()
        worksheet = book.add_sheet("Sheet")
        for index, row in enumerate(rows):
            for column, value in enumerate(row):
                worksheet.write(index, column, value)
        output = BytesIO()
        book.save(output)
        raw, sheet = output.getvalue(), "Sheet"
    spec = Specification(
        format=kind,
        columns=(Column(column="A", role="context"), Column(column="B", role="amount")),
        total_rows={sheet: (4,)},
    )
    result = inspect_bytes(raw, spec)
    assert len(result["items"]) == 2 and not result["issues"]
    assert result["control_totals"][0]["actual_fen"] == 246
    with pytest.raises(KernelError, match="十万条数据行"):
        inspect_bytes(raw, spec.model_copy(update={"total_rows": {}}))


@pytest.mark.parametrize("reference", ["A1:XFD1", "A1:A1048576"])
def test_xlsx_formatted_used_range_does_not_expand_or_limit_actual_data(reference):
    raw = workbook_bytes([["amount"], ["1.23"]])
    raw = replace_xml(
        raw,
        "xl/worksheets/sheet1.xml",
        lambda document: document.find("{*}dimension").set("ref", reference),
    )
    result = inspect_bytes(
        raw, Specification(format="xlsx", columns=(Column(column="A", role="amount"),))
    )
    assert len(result["items"]) == 1 and result["items"][0]["amount_fen"] == 123


def test_xlsx_checks_actual_column_even_if_dimension_claims_small_sheet():
    raw = workbook_bytes([["amount"], ["1.23"]])
    raw = replace_xml(
        raw,
        "xl/worksheets/sheet1.xml",
        lambda document: document.find(".//{*}c[@r='A2']").set("r", "CM2"),
    )
    with pytest.raises(KernelError) as error:
        inspect_bytes(
            raw, Specification(format="xlsx", columns=(Column(column="A", role="amount"),))
        )
    assert error.value.code == "material_too_large"


def test_global_row_budget_applies_across_hidden_sheets(monkeypatch):
    monkeypatch.setattr(materials, "MAX_DATA_ROWS", 2)
    book = Workbook()
    for index in range(2):
        sheet = book.active if index == 0 else book.create_sheet("Hidden")
        sheet.append(["amount"])
        sheet.append(["1.23"])
        sheet.append(["1.23"])
        if index:
            sheet.sheet_state = "veryHidden"
    output = BytesIO()
    book.save(output)
    book.close()
    with pytest.raises(KernelError) as error:
        inspect_bytes(
            output.getvalue(),
            Specification(format="xlsx", columns=(Column(column="A", role="amount"),)),
        )
    assert error.value.code == "material_too_large"


def test_original_byte_and_expanded_zip_bounds(monkeypatch):
    with pytest.raises(KernelError) as error:
        inspect_bytes(b"x" * (20 * 1024 * 1024 + 1), Specification(format="csv"))
    assert error.value.code == "material_too_large"
    monkeypatch.setattr(materials, "MAX_EXPANDED_BYTES", 1)
    with pytest.raises(KernelError) as error:
        inspect_bytes(workbook_bytes([["amount"], ["1.23"]]), Specification(format="xlsx"))
    assert error.value.code == "material_too_large"


@pytest.mark.parametrize(
    "changes",
    [{"header_rows": {"CSV": 101}}, {"total_rows": {"CSV": tuple(range(1, 1002))}}],
)
def test_header_and_control_exemptions_have_independent_bounds(changes):
    spec = Specification(format="csv", columns=(Column(column="A", role="amount"),), **changes)
    with pytest.raises(KernelError) as error:
        inspect_bytes(b"1.23\n", spec)
    assert error.value.code == "material_too_large"


def test_boolean_cell_is_not_implicitly_treated_as_money():
    result = inspect_bytes(
        workbook_bytes([["amount"], [True]]),
        Specification(format="xlsx", columns=(Column(column="A", role="amount"),)),
    )
    assert result["items"][0]["amount_fen"] is None
    assert {item["code"] for item in result["issues"]} == {"material_amount_missing"}


def test_mapping_missing_amounts_cannot_amplify_beyond_cell_budget(monkeypatch):
    monkeypatch.setattr(materials, "MAX_NONEMPTY_CELLS", 2)
    spec = Specification(
        format="csv",
        columns=(
            Column(column="A", role="context"),
            Column(column="B", role="amount"),
            Column(column="C", role="amount"),
        ),
        header_rows={"CSV": 0},
    )
    with pytest.raises(KernelError) as error:
        inspect_bytes(b"business row\n", spec)
    assert error.value.code == "material_too_large"


def test_duplicate_xml_cell_location_is_a_structured_parse_failure():
    raw = workbook_bytes([["amount"], ["1.23", "2.00"]])
    raw = replace_xml(
        raw,
        "xl/worksheets/sheet1.xml",
        lambda document: document.find(".//{*}c[@r='B2']").set("r", "A2"),
    )
    with pytest.raises(KernelError) as error:
        inspect_bytes(raw, Specification(format="xlsx"))
    assert error.value.code == "material_unreadable"


@pytest.mark.parametrize(
    "value,expected",
    [
        ("1.23", 123),
        ("1.234", None),
        ("1e999999999", None),
        ("1e-999999999", None),
        ("92233720368547758.07", 2**63 - 1),
        ("-92233720368547758.08", -(2**63)),
        ("92233720368547758.08", None),
    ],
)
def test_material_money_is_exact_under_low_ambient_decimal_precision(value, expected):
    with localcontext() as context:
        context.prec = 2
        result = inspect_bytes(
            f"amount\n{value}\n".encode(),
            Specification(format="csv", columns=(Column(column="A", role="amount"),)),
        )
    assert result["items"][0]["amount_fen"] == expected
    assert bool(result["issues"]) is (expected is None)


@pytest.mark.parametrize("storage", ["inlineStr", "sharedStrings"])
@pytest.mark.parametrize("rich_text", [False, True])
def test_xlsx_phonetic_guides_are_not_part_of_plain_or_rich_amount(storage, rich_text):
    namespace = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    text = ElementTree.Element(
        f"{{{namespace}}}is" if storage == "inlineStr" else f"{{{namespace}}}si"
    )
    if rich_text:
        for part in ("1", "23"):
            run = ElementTree.SubElement(text, f"{{{namespace}}}r")
            ElementTree.SubElement(run, f"{{{namespace}}}t").text = part
    else:
        ElementTree.SubElement(text, f"{{{namespace}}}t").text = "123"
    guide = ElementTree.SubElement(text, f"{{{namespace}}}rPh", sb="0", eb="3")
    ElementTree.SubElement(guide, f"{{{namespace}}}t").text = "4"

    def patch_cell(document):
        cell = document.find(".//{*}c[@r='A2']")
        for child in list(cell):
            cell.remove(child)
        if storage == "inlineStr":
            cell.set("t", "inlineStr")
            cell.append(text)
        else:
            cell.set("t", "s")
            ElementTree.SubElement(cell, f"{{{namespace}}}v").text = "0"

    raw = replace_xml(
        workbook_bytes([["amount"], ["placeholder"]]), "xl/worksheets/sheet1.xml", patch_cell
    )
    if storage == "sharedStrings":
        output = BytesIO()
        with ZipFile(BytesIO(raw)) as original, ZipFile(output, "w") as archive:
            for name in original.namelist():
                content = original.read(name)
                if name == "[Content_Types].xml":
                    document = ElementTree.fromstring(content)
                    tag = document.tag.replace("Types", "Override")
                    ElementTree.SubElement(
                        document,
                        tag,
                        PartName="/xl/sharedStrings.xml",
                        ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml",
                    )
                    content = ElementTree.tostring(document)
                elif name == "xl/_rels/workbook.xml.rels":
                    document = ElementTree.fromstring(content)
                    ElementTree.SubElement(
                        document,
                        document.tag.replace("Relationships", "Relationship"),
                        Id="rIdSyntheticStrings",
                        Target="sharedStrings.xml",
                        Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/sharedStrings",
                    )
                    content = ElementTree.tostring(document)
                archive.writestr(name, content)
            table = ElementTree.Element(f"{{{namespace}}}sst", count="1", uniqueCount="1")
            table.append(text)
            archive.writestr("xl/sharedStrings.xml", ElementTree.tostring(table))
        raw = output.getvalue()
    book = load_workbook(BytesIO(raw), read_only=True)
    try:
        assert book.active["A2"].value == "123"
    finally:
        book.close()
    result = inspect_bytes(
        raw, Specification(format="xlsx", columns=(Column(column="A", role="amount"),))
    )
    assert not result["issues"]
    assert result["items"][0]["amount_fen"] == 12300
    assert (
        next(cell for cell in result["coverage"] if cell["location"] == "Sheet!A2")["value"]
        == "123"
    )
