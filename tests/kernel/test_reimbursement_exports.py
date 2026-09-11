"""Payment exports consume published reimbursement debts, including multi-party batches."""

from pathlib import Path

import pytest
from openpyxl import load_workbook
from test_exports import evidence, inventory, queue, template_bytes
from test_payroll_corrections import Company
from test_reimbursement_assets import accepted_batch, asset, batch_card, pay

from ai_accounting.kernel.contracts import KernelError, NeedsInformation
from ai_accounting.kernel.domains.assets import ReimbursedAsset, ReimbursedAssetBatch
from ai_accounting.kernel.domains.transactions import Payment, ReimbursedDeposit
from ai_accounting.kernel.exports import WORKBOOK_NAME, Exports, run_export_jobs
from ai_accounting.kernel.periods import Periods
from ai_accounting.kernel.types import canonical


def parse(model, data):
    return model.model_validate_json(canonical(data))


@pytest.fixture
def reimbursement_company(tmp_path):
    company = Company(tmp_path / "reimbursements.sqlite")
    for fact, subject in (
        (parse(ReimbursedAsset, asset()), "direct"),
        (parse(ReimbursedAssetBatch, accepted_batch()), "batch"),
        (parse(ReimbursedAsset, batch_card()), "computer"),
        (parse(ReimbursedAsset, batch_card(30000)), "chair"),
        (
            ReimbursedDeposit(
                period="2026-02",
                counterparty_id="landlord",
                employee_id="alice",
                amount_fen=50000,
                company_acceptance_confirmed=True,
                refund_right_confirmed=True,
            ),
            "deposit",
        ),
    ):
        company.save(fact, subject)
    company.publish("direct", "batch", "computer", "chair", "deposit")
    export = Exports(company.engine)
    template = evidence(company, template_bytes())
    inventory(company, "2026-02")
    return company, export, template


def payees(company, export, *, parties=("alice", "bob")):
    for index, party in enumerate(parties):
        export.save_payee(
            party,
            name=party,
            account=f"00123456789{index}",
            evidence_digest=company.owner_confirmation,
            expected_revision=0,
            request_id=company.request(),
        )


def sources(plan):
    return [source for row in plan["rows"] for source in row["sources"]]


def test_multi_employee_and_same_person_cross_sources_have_distinct_obligations(
    reimbursement_company, tmp_path
):
    company, export, template = reimbursement_company
    payees(company, export)
    plan, job, options = queue(
        company, export, template, tmp_path / "instructions", period="2026-02"
    )
    assert {row["party_id"]: row["amount_fen"] for row in plan["rows"]} == {
        "alice": 220000,
        "bob": 100000,
    }
    assert len(sources(plan)) == 5
    assert {source["kind"] for source in sources(plan)} == {
        "reimbursed_asset",
        "reimbursed_asset_batch",
        "reimbursed_deposit",
    }
    assert {source["subject_id"] for source in sources(plan)} == {"direct", "batch", "deposit"}
    assert (
        len(
            {
                (source["subject_id"], source["calculation_id"], source["obligation"])
                for source in sources(plan)
            }
        )
        == 5
    )
    assert {source["category"] for source in sources(plan)} == {"报销"}
    assert not any(source["obligation"].endswith(":refund") for source in sources(plan))
    assert export.confirm("2026-02", **options) == job
    assert run_export_jobs(company.engine)[0]["status"] == "succeeded"
    book = load_workbook(Path(options["output_directory"]) / WORKBOOK_NAME, data_only=True)
    try:
        rows = list(book.active.iter_rows(values_only=True))
        assert len(rows) == 3  # header + the two original party/category grouping keys
        assert {row[0] for row in rows[1:]} == {"alice", "bob"}
    finally:
        book.close()
    # File generation is not payment, and does not reduce any published debt.
    again = export.preview("2026-02", template_evidence_digest=template)
    assert again["total_fen"] == plan["total_fen"] == 320000


def test_actual_partial_and_full_payments_reduce_only_the_named_debt(reimbursement_company):
    company, export, template = reimbursement_company
    payees(company, export)
    for subject, kind, source, party, amount in (
        ("alice-partial", "reimbursed_asset_batch", "batch", "alice", 25000),
        ("bob-complete", "reimbursed_asset", "direct", "bob", 40000),
    ):
        company.save(
            parse(
                Payment,
                pay(kind, source, party, party, amount, period="2026-02", actual_date="2026-02-15"),
            ),
            subject,
        )
    company.publish("alice-partial", "bob-complete")
    inventory(company, "2026-02")
    plan = export.preview("2026-02", template_evidence_digest=template)
    assert {row["party_id"]: row["amount_fen"] for row in plan["rows"]} == {
        "alice": 195000,
        "bob": 60000,
    }
    actual = {row["obligation"]: row["amount_fen"] for row in sources(plan)}
    assert actual["reimbursed_asset_batch:batch:alice"] == 65000
    assert "reimbursed_asset:direct:bob" not in actual
    assert plan["total_fen"] == 255000


def test_deposit_refund_is_not_exported_or_required_to_have_a_payee(reimbursement_company):
    company, export, template = reimbursement_company
    payees(company, export, parties=("alice",))
    plan = export.preview("2026-02", template_evidence_digest=template, source_ids=["deposit"])
    assert len(sources(plan)) == 1
    assert sources(plan)[0]["obligation"] == "reimbursed_deposit:deposit:reimbursement"
    assert plan["total_fen"] == 50000
    company.save(
        parse(
            Payment,
            pay(
                "reimbursed_deposit",
                "deposit",
                "reimbursement",
                "alice",
                50000,
                period="2026-02",
                actual_date="2026-02-15",
            ),
        ),
        "deposit-repaid",
    )
    company.publish("deposit-repaid")
    inventory(company, "2026-02")
    with pytest.raises(KernelError) as completed:
        export.preview("2026-02", template_evidence_digest=template, source_ids=["deposit"])
    assert completed.value.code == "no_payables"
    with company.engine.store.connection(read_only=True) as connection:
        assert (
            connection.execute(
                "SELECT amount FROM balance WHERE balance_key=?",
                ("reimbursed_deposit:deposit:refund",),
            ).fetchone()[0]
            == 50000
        )


def test_missing_payee_only_blocks_affected_export_and_complete_material_gate_stays(
    reimbursement_company,
):
    company, export, template = reimbursement_company
    payees(company, export, parties=("alice",))
    # Both acquisition and month readiness are valid before Bob's management account exists.
    assert (
        Periods(company.engine).preview_close(
            "2026-02", owner_confirmation=company.owner_confirmation
        )["status"]
        == "preview"
    )
    with pytest.raises(NeedsInformation, match="收款姓名和账号"):
        export.preview("2026-02", template_evidence_digest=template)
    assert (
        export.preview("2026-02", template_evidence_digest=template, source_ids=["deposit"])[
            "total_fen"
        ]
        == 50000
    )
    payees(company, export, parties=("bob",))
    known = sorted({digest for _, digest in company.materials["assets"]})
    Periods(company.engine).inventory(
        "2026-02",
        "assets",
        evidence=known,
        expected=len(known) + 1,
        no_business=False,
        confirmation_evidence=company.owner_confirmation,
        request_id=company.request(),
    )
    for selected in (None, ["deposit"]):
        with pytest.raises(KernelError) as incomplete:
            export.preview("2026-02", template_evidence_digest=template, source_ids=selected)
        assert incomplete.value.code == "materials_incomplete"


def test_batch_card_cannot_duplicate_parent_debt_and_repeated_selection_is_idempotent(
    reimbursement_company,
):
    company, export, template = reimbursement_company
    payees(company, export)
    first = export.preview("2026-02", template_evidence_digest=template, source_ids=["batch"])
    repeat = export.preview(
        "2026-02", template_evidence_digest=template, source_ids=["batch", "batch"]
    )
    assert repeat == first and first["total_fen"] == 150000
    assert len(sources(first)) == 2
    for selected in (["computer"], ["batch", "computer"]):
        with pytest.raises(KernelError) as duplicate:
            export.preview("2026-02", template_evidence_digest=template, source_ids=selected)
        assert duplicate.value.code == "invalid_export_sources"
