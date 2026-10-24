"""Different bank layouts preserve account text, exact cents, styles and fixed instructions."""

from io import BytesIO

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

from ai_accounting_modular.contracts.public import ContractError
from ai_accounting_modular.payments import _workbook


def template(*, shifted=False, footer=False):
    book = Workbook()
    sheet = book.active
    sheet.title = "银行批量付款"
    header = 3 if shifted else 1
    columns = (2, 4, 6, 8) if shifted else (1, 2, 3, 4)
    for col, value in zip(
        columns, ("收款方名称(*)", "收款方账号(*)", "金额(*)", "附言/用途"), strict=True
    ):
        sheet.cell(header, col, value)
        sheet.cell(header + 1, col).fill = PatternFill("solid", fgColor="ABCDEF")
        sheet.cell(header + 1, col).font = Font(bold=True)
        sheet.column_dimensions[sheet.cell(1, col).column_letter].width = 28
    sheet.cell(2, 10, "银行说明：账号必须为文本，不是付款记录")
    sheet.merge_cells("J3:L3")
    sheet["J3"] = "模板固定说明"
    if footer:
        sheet.cell(8, columns[0], "下方固定说明，不可覆盖")
    stream = BytesIO()
    book.save(stream)
    return stream.getvalue()


ROWS = [
    {
        "name": "合成职员",
        "account": "00012345678901234567890",
        "amount_fen": "123456789012345",
        "purpose": "2026-09 工资",
    }
]


@pytest.mark.parametrize("shifted", [False, True])
def test_exact_text_and_styles_roundtrip_in_different_bank_layouts(shifted):
    raw = template(shifted=shifted)
    layout = _workbook.resolve_layout(raw)
    assert layout["first_data_row"] == (4 if shifted else 2)
    generated = _workbook.render(raw, layout, ROWS)
    _workbook.verify(generated, raw, layout, ROWS)
    book = load_workbook(BytesIO(generated), data_only=False)
    sheet = book[layout["sheet_name"]]
    cells = [
        sheet.cell(layout["first_data_row"], layout["columns"][key]) for key in _workbook.FIELDS
    ]
    assert [cell.value for cell in cells] == [
        "合成职员",
        "00012345678901234567890",
        "1234567890123.45",
        "2026-09 工资",
    ]
    assert all(cell.data_type == "s" and cell.number_format == "@" for cell in cells)
    assert all(cell.font.bold and cell.fill.fgColor.rgb == "00ABCDEF" for cell in cells)
    assert sheet["J2"].value == "银行说明：账号必须为文本，不是付款记录"
    assert sheet["J3"].value == "模板固定说明"
    sheet.cell(layout["first_data_row"], layout["columns"]["amount"], "1234567890123.46")
    altered = BytesIO()
    book.save(altered)
    book.close()
    with pytest.raises(ContractError, match="逐格"):
        _workbook.verify(altered.getvalue(), raw, layout, ROWS)


def test_explicit_writable_region_keeps_footer_and_rejects_overflow():
    raw = template(shifted=True, footer=True)
    with pytest.raises(ContractError, match="last_data_row"):
        _workbook.resolve_layout(raw)
    layout = {
        "sheet_name": "银行批量付款",
        "first_data_row": 4,
        "last_data_row": 6,
        "columns": dict(zip(_workbook.FIELDS, (2, 4, 6, 8), strict=True)),
    }
    assert _workbook.resolve_layout(raw, layout) == layout
    generated = _workbook.render(raw, layout, ROWS)
    _workbook.verify(generated, raw, layout, ROWS)
    book = load_workbook(BytesIO(generated))
    assert book.active["B8"].value == "下方固定说明，不可覆盖"
    book.close()
    with pytest.raises(ContractError, match="可写区域"):
        _workbook.render(raw, layout, ROWS * 4)


def test_explicit_layout_can_start_at_first_row_without_assuming_a_bank_header():
    book = Workbook()
    book.active.title = "导入"
    stream = BytesIO()
    book.save(stream)
    book.close()
    raw = stream.getvalue()
    layout = {
        "sheet_name": "导入",
        "first_data_row": 1,
        "columns": dict(zip(_workbook.FIELDS, (1, 2, 3, 4), strict=True)),
    }
    generated = _workbook.render(raw, layout, ROWS)
    _workbook.verify(generated, raw, layout, ROWS)
    actual = load_workbook(BytesIO(generated))
    assert actual.active["A1"].value == ROWS[0]["name"]
    actual.close()


def test_sparse_fixed_instructions_do_not_expand_a_blank_payment_rectangle():
    book = load_workbook(BytesIO(template()))
    book.active.cell(10000, 100, "银行保留说明")
    stream = BytesIO()
    book.save(stream)
    book.close()
    raw = stream.getvalue()
    layout = {
        "sheet_name": "银行批量付款",
        "first_data_row": 2,
        "columns": dict(zip(_workbook.FIELDS, (1, 2, 3, 4), strict=True)),
    }
    generated = _workbook.render(raw, layout, ROWS)
    _workbook.verify(generated, raw, layout, ROWS)
    actual = load_workbook(BytesIO(generated))
    assert actual.active.cell(10000, 100).value == "银行保留说明"
    assert len(actual.active._cells) < 40
    actual.close()


def test_ambiguous_headers_and_merged_payment_regions_are_explicitly_unsupported():
    book = Workbook()
    for title in ("银行甲", "银行乙"):
        sheet = book.create_sheet(title)
        sheet.append(["收款方名称", "收款方账号", "金额", "用途"])
    stream = BytesIO()
    book.save(stream)
    with pytest.raises(ContractError, match="唯一识别"):
        _workbook.resolve_layout(stream.getvalue())
    sheet = book["银行甲"]
    layout = {
        "sheet_name": sheet.title,
        "first_data_row": 2,
        "columns": dict(zip(_workbook.FIELDS, (1, 2, 3, 4), strict=True)),
    }
    sheet.merge_cells("B2:B3")
    stream = BytesIO()
    book.save(stream)
    with pytest.raises(ContractError, match="合并"):
        _workbook.resolve_layout(stream.getvalue(), layout)
    book.close()


@pytest.mark.parametrize("value", ["=1+1", "+123", "@malicious", "-123", "含\n换行"])
def test_formula_markers_are_preserved_as_literal_text(value):
    raw = template()
    layout = _workbook.resolve_layout(raw)
    rows = [{**ROWS[0], "name": value, "account": "+000123", "purpose": "=工资用途"}]
    generated = _workbook.render(raw, layout, rows)
    _workbook.verify(generated, raw, layout, rows)
    book = load_workbook(BytesIO(generated), data_only=False)
    assert [book.active.cell(2, col).value for col in (1, 2, 4)] == [value, "+000123", "=工资用途"]
    assert all(book.active.cell(2, col).data_type == "s" for col in (1, 2, 3, 4))
    book.close()


def test_template_formulas_outside_payment_area_are_kept_and_sample_formulas_are_replaced():
    book = load_workbook(BytesIO(template()))
    sheet = book.active
    sheet.append(["原模板示例", "原模板账号", "=1+1", "示例用途"])
    sheet["J2"] = "=SUM(C2:C9)"
    notes = book.create_sheet("银行说明")
    notes["A1"] = "=1+2"
    stream = BytesIO()
    book.save(stream)
    book.close()
    raw = stream.getvalue()
    layout = _workbook.resolve_layout(raw)
    generated = _workbook.render(raw, layout, ROWS)
    _workbook.verify(generated, raw, layout, ROWS)
    actual = load_workbook(BytesIO(generated), data_only=False)
    assert actual.active["J2"].value == "=SUM(C2:C9)"
    assert actual.active["J2"].data_type == "f"
    assert actual["银行说明"]["A1"].value == "=1+2"
    assert actual.active["C2"].data_type == "s"
    assert actual.active["C4"].value is None
    actual.close()
