"""Pinned workbook compatibility and complete saved-plan verification, without a service."""

import io
import json
import zipfile
from copy import deepcopy
from datetime import date
from decimal import Decimal, localcontext
from pathlib import Path
from types import SimpleNamespace
from uuid import uuid4
from xml.etree import ElementTree

import pytest
from openpyxl import load_workbook

from ai_accounting_modular.contracts.public import ContractError, content_digest
from ai_accounting_modular.reading import _report_plans, _report_workbook
from ai_accounting_modular.reading.public import create_report_producer

ROOT = Path(__file__).resolve().parents[3]


@pytest.fixture
def plan():
    value = json.loads((ROOT / "docs/redesign/examples/file-cases.json").read_text("utf-8"))[
        "contract_examples"
    ]["report_export_plan"]
    value["variant"] = "tax_bureau"
    value["snapshot"]["scope"].update(period_to="2026-03", as_of="2026-03")
    value["organization"]["legal_name"] = "合成 & <核验>公司"
    value["scope_head"] = _report_plans.scope_head(value)
    return value


def test_retained_template_all_cells_formula_caches_and_exact_unchanged_parts(plan):
    _report_plans.validate(plan)
    original = _report_workbook._template_bytes()
    raw = _report_workbook.render(plan, str(uuid4()))
    source = load_workbook(io.BytesIO(original), data_only=False)
    book = load_workbook(io.BytesIO(raw), data_only=False)
    cached = load_workbook(io.BytesIO(raw), data_only=True)
    try:
        assert book.sheetnames == ["资产负债表", "利润表_月季报", "现金流量表_月季报"]
        for before, after in zip(source.worksheets, book.worksheets, strict=True):
            assert tuple(before.merged_cells.ranges) == tuple(after.merged_cells.ranges)
            assert before.protection == after.protection
            assert before.data_validations == after.data_validations
            for row in before:
                for cell in row:
                    assert cell.style_id == after[cell.coordinate].style_id
                    if cell.data_type == "f":
                        assert cell.value == after[cell.coordinate].value
        assert cached.worksheets[0]["D3"].value == plan["organization"]["credit_code"]
        assert cached.worksheets[0]["H3"].value == "合成 & <核验>公司"
        assert cached.worksheets[1]["F3"].value == "合成 & <核验>公司"
        assert cached.worksheets[1]["D4"].value.date() == date(2026, 1, 1)
        assert cached.worksheets[2]["F4"].value.date() == date(2026, 3, 31)
        rows = [*range(7, 22), *range(23, 38)]
        right_rows = [*range(7, 18), *range(19, 25), *range(32, 38)]
        for line, row in enumerate(rows, 1):
            for column, name in (("D", "ending"), ("E", "beginning")):
                expected = plan["snapshot"]["balance_sheet"][str(line)][name]["amount_fen"]
                assert Decimal(str(cached.worksheets[0][f"{column}{row}"].value)) * 100 == int(
                    expected
                )
        for line, row in enumerate(right_rows, 31):
            for column, name in (("H", "ending"), ("I", "beginning")):
                expected = plan["snapshot"]["balance_sheet"][str(line)][name]["amount_fen"]
                assert Decimal(str(cached.worksheets[0][f"{column}{row}"].value)) * 100 == int(
                    expected
                )
        for sheet, table, coordinates in (
            (1, "profit_statement", range(6, 38)),
            (2, "cash_flow_statement", [*range(7, 14), *range(15, 21), *range(22, 31)]),
        ):
            for line, row in enumerate(coordinates, 1):
                for column, name in (("D", "current"), ("E", "year_to_date")):
                    expected = plan["snapshot"][table][str(line)][name]["amount_fen"]
                    assert Decimal(
                        str(cached.worksheets[sheet][f"{column}{row}"].value)
                    ) * 100 == int(expected)
    finally:
        source.close()
        book.close()
        cached.close()
    with zipfile.ZipFile(io.BytesIO(original)) as before, zipfile.ZipFile(io.BytesIO(raw)) as after:
        assert before.namelist() == after.namelist()
        changed = {name for name in before.namelist() if before.read(name) != after.read(name)}
        assert changed == {
            "xl/workbook.xml",
            "docProps/custom.xml",
            *(f"xl/worksheets/sheet{number}.xml" for number in (1, 2, 3)),
        }


@pytest.mark.parametrize("damage", ["amount", "cache", "formula", "identity", "extra"])
def test_complete_verifier_refuses_changed_workbook_even_with_valid_zip(plan, damage):
    context = SimpleNamespace(task_id=str(uuid4()))
    original = _report_workbook.render(plan, context.task_id)
    with zipfile.ZipFile(io.BytesIO(original)) as source:
        entries = {info.filename: source.read(info.filename) for info in source.infolist()}
    if damage == "identity":
        entries["docProps/custom.xml"] = entries["docProps/custom.xml"].replace(
            context.task_id.encode(), str(uuid4()).encode()
        )
    elif damage == "extra":
        entries["hidden.xml"] = b"unexpected"
    else:
        name = "xl/worksheets/sheet1.xml"
        xml = entries[name].decode()
        if damage == "formula":
            xml = xml.replace("<f>", "<f>1+", 1)
        elif damage == "cache":
            xml = xml.replace("<v>", "<v>9", 1)
        else:
            xml = _report_workbook._set_numeric(xml, "D37", "9999.99")
        entries[name] = xml.encode()
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w") as package:
        for name, raw in entries.items():
            package.writestr(name, raw)
    output.seek(0)
    with pytest.raises(ContractError):
        create_report_producer(None).verify(plan, output, context)


def test_identity_in_existing_custom_properties_and_signed_cents_are_literal(plan):
    plan["snapshot"]["profit_statement"]["32"]["current"]["amount_fen"] = "-1"
    plan["scope_head"] = _report_plans.scope_head(plan)
    context = SimpleNamespace(task_id=str(uuid4()))
    raw = _report_workbook.render(plan, context.task_id)
    with zipfile.ZipFile(io.BytesIO(raw)) as workbook:
        properties = ElementTree.fromstring(workbook.read("docProps/custom.xml"))
        binding = next(item for item in properties if item.attrib["name"] == "AccountingReport")
        assert json.loads(binding[0].text) == {
            **{
                key: plan[key]
                for key in (
                    "company_id",
                    "database_id",
                    "scope_head",
                    "variant",
                    "template_version",
                )
            },
            "task_id": context.task_id,
        }
    book = load_workbook(io.BytesIO(raw), data_only=True)
    try:
        assert book.worksheets[1]["D37"].value == -0.01
    finally:
        book.close()


@pytest.mark.parametrize(
    "fault", ["default", "version", "partial", "unknown", "large", "month", "digest"]
)
def test_unsupported_or_incomplete_plan_is_not_a_fake_complete_export(plan, fault):
    if fault == "default":
        plan["variant"] = "default"
    elif fault == "version":
        plan["template_version"] = 2
    elif fault == "partial":
        plan["snapshot"]["status"] = "partial"
    elif fault == "unknown":
        plan["snapshot"]["balance_sheet"]["1"]["ending"] = {
            "state": "unestablished",
            "issue_indexes": [0],
        }
    elif fault == "large":
        plan["snapshot"]["balance_sheet"]["1"]["ending"]["amount_fen"] = "999999999999999999"
    elif fault == "month":
        plan["snapshot"]["scope"].update(period_to="2026-01", as_of="2026-01")
    if fault != "digest":
        plan["scope_head"] = _report_plans.scope_head(plan)
    else:
        plan["scope_head"] = "0" * 64
    with pytest.raises(ContractError):
        create_report_producer(None).describe(plan)


def test_scope_ignores_unrelated_global_or_working_head_but_binds_rendered_inputs(plan):
    original = plan["scope_head"]
    plan["snapshot"]["head"] += 100
    plan["organization"]["working_period_to"] = "2027-12"
    assert _report_plans.scope_head(plan) == original
    changed = deepcopy(plan)
    changed["organization"]["legal_name"] = "另一公司名称"
    assert _report_plans.scope_head(changed) != original
    changed = deepcopy(plan)
    changed["snapshot"]["balance_sheet"]["1"]["ending"]["amount_fen"] = "1"
    assert _report_plans.scope_head(changed) != original
    assert content_digest(plan["snapshot"]) != content_digest(changed["snapshot"])


@pytest.mark.parametrize(
    "text", ["公司\x01名称", "公司\x00名称", "公司\ud800名称", "公司\ufffe名称"]
)
def test_unrepresentable_company_text_is_refused_without_changing_the_identity(plan, text):
    from ai_accounting_modular.reading import _models

    plan["organization"]["legal_name"] = text
    metadata = SimpleNamespace(**plan["organization"], profile_seq=1)
    state = {
        "content_digest": content_digest({"snapshot": _models.semantic(plan["snapshot"])}),
        "ready_head": 1,
    }
    preview = _report_plans.preview(plan["snapshot"], state, metadata)
    assert preview["available"] is False
    assert "Excel不支持的字符" in preview["reason"]
    assert metadata.legal_name == text
    # A lone surrogate is already rejected by the shared JSON contract; normal
    # JSON control characters reach the workbook's more specific text check.
    with pytest.raises(ContractError, match="Excel不支持的字符|无效 JSON"):
        _report_plans.validate(plan)
    with pytest.raises(ContractError, match="Excel不支持的字符"):
        _report_workbook.render(plan, str(uuid4()))


def test_xml_safe_whitespace_and_escaped_characters_survive_independent_parsing(plan):
    text = "合成\t公司\n分部\r&<核验>😀"
    plan["organization"]["legal_name"] = text
    plan["scope_head"] = _report_plans.scope_head(plan)
    raw = _report_workbook.render(plan, str(uuid4()))
    with zipfile.ZipFile(io.BytesIO(raw)) as workbook:
        for name in workbook.namelist():
            if name.endswith((".xml", ".rels")):
                ElementTree.fromstring(workbook.read(name))
    book = load_workbook(io.BytesIO(raw), data_only=True)
    try:
        assert book.worksheets[0]["H3"].value == text
        assert book.worksheets[1]["F3"].value == text
    finally:
        book.close()


def test_fast_preview_uses_verified_component_metadata_without_revalidating_schema(
    plan, monkeypatch
):
    from ai_accounting_modular.reading import _models

    metadata = SimpleNamespace(**plan["organization"], profile_seq=7)
    state = {
        "content_digest": content_digest({"snapshot": _models.semantic(plan["snapshot"])}),
        "ready_head": 5,
    }
    expected = _report_plans.scope_head(plan)
    monkeypatch.setattr(
        _report_plans, "validate_contract", lambda *_: pytest.fail("重复校验完整Schema")
    )
    monkeypatch.setattr(_report_plans, "plan", lambda *_: pytest.fail("热读重新创建完整计划"))
    preview = _report_plans.preview(plan["snapshot"], state, metadata)
    assert preview["preview_digest"] == expected
    assert preview["epochs"] == {"accounting": 5, "material": 0, "management": 7}


def test_fen_conversion_is_exact_independent_of_decimal_context(plan):
    plan["snapshot"]["balance_sheet"]["1"]["ending"]["amount_fen"] = "123456"
    with localcontext() as context:
        context.prec = 2
        raw = _report_workbook.render(plan, str(uuid4()))
    book = load_workbook(io.BytesIO(raw), data_only=True)
    try:
        assert book.worksheets[0]["D7"].value == 1234.56
    finally:
        book.close()


def test_wrong_member_size_is_rejected_before_inflating_its_body(plan, monkeypatch):
    context = SimpleNamespace(task_id=str(uuid4()))
    raw = _report_workbook.render(plan, context.task_id)
    name = "xl/worksheets/sheet1.xml"
    with zipfile.ZipFile(io.BytesIO(raw)) as original:
        entries = {item: original.read(item) for item in original.namelist()}
    entries[name] = b"x" * 1_000_000
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for item, body in entries.items():
            archive.writestr(item, body)
    output.seek(0)
    read = zipfile.ZipFile.read

    def bounded(archive, member, *args, **kwargs):
        if archive.fp is output and member == name:
            pytest.fail("inflated a member whose size cannot belong to this workbook")
        return read(archive, member, *args, **kwargs)

    monkeypatch.setattr(zipfile.ZipFile, "read", bounded)
    with pytest.raises(ContractError):
        create_report_producer(None).verify(plan, output, context)
