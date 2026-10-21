"""Hidden spreadsheet content and PDF extraction gaps are visible, without recalculating facts."""

import io
import json
from datetime import datetime, timedelta
from xml.etree import ElementTree
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
import xlwt
from openpyxl import Workbook
from openpyxl.utils.datetime import CALENDAR_MAC_1904
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject, NumberObject
from test_source_api import incoming, query, store

from ai_accounting_modular.contracts.public import ContractError


def xlsx_bytes():
    book = Workbook()
    sheet = book.active
    sheet.title = "正常页"
    sheet["A1"] = "单元格原文" * 4500
    sheet["B2"] = "=SUM(1,2)"
    sheet.row_dimensions[2].hidden = True
    hidden = book.create_sheet("隐藏页")
    hidden.sheet_state = "hidden"
    hidden["C9"] = "隐藏行原文"
    hidden.column_dimensions["C"].hidden = True
    raw = io.BytesIO()
    book.save(raw)
    # A declared used range is not proof of actual rows; preserve the actual stored cells only.
    output = io.BytesIO()
    with ZipFile(raw) as archive, ZipFile(output, "w", ZIP_DEFLATED) as rewritten:
        for name in archive.namelist():
            data = archive.read(name)
            if name == "xl/worksheets/sheet1.xml":
                data = data.replace(b'<dimension ref="A1:B2"', b'<dimension ref="A1:XFD1048576"')
            rewritten.writestr(name, data)
    return output.getvalue()


def test_xlsx_preserves_hidden_cells_long_values_formula_and_empty_cache(records_case):
    storage, (company, _), records = records_case
    source = store(storage, company, records, [incoming(company, "原表.xlsx", xlsx_bytes())])[0]
    result = query(
        records, storage, company, action="inspect", source_id=source["source_id"], limit=500
    )["inspection"]
    assert not result["unsupported_media"]
    long = "".join(
        item["text"] for item in result["blocks"] if item["location"].startswith("sheet:正常页!A1")
    )
    assert json.loads(long)["value"] == "单元格原文" * 4500
    hidden = [item for item in result["blocks"] if item["location"].startswith("sheet:隐藏页!C9")]
    assert len(hidden) == 1 and "[hidden]" in hidden[0]["location"]
    assert json.loads(hidden[0]["text"])["value"] == "隐藏行原文"
    formula = next(
        item for item in result["blocks"] if item["location"].startswith("sheet:正常页!B2")
    )
    assert "[hidden]" in formula["location"]
    assert json.loads(formula["text"])["formula"]["text"] == "SUM(1,2)"
    assert json.loads(formula["text"])["value"] is None
    scoped = query(
        records,
        storage,
        company,
        action="inspect",
        source_id=source["source_id"],
        section="sheet:隐藏页",
        limit=20,
    )["inspection"]
    assert scoped["page"]["total_count"] == result["page"]["total_count"]
    assert scoped["page"]["filtered_count"] == 1
    with pytest.raises(ContractError) as failed:
        query(
            records,
            storage,
            company,
            action="inspect",
            source_id=source["source_id"],
            section="sheet:不存在",
            limit=20,
        )
    assert failed.value.code == "invalid_input"


def test_xls_includes_hidden_sheet_row_and_column_values(records_case):
    storage, (company, _), records = records_case
    book = xlwt.Workbook()
    normal = book.add_sheet("正常")
    normal.write(0, 0, "真实名称")
    hidden = book.add_sheet("隐藏")
    hidden.visibility = 1
    hidden.row(1).hidden = True
    hidden.col(2).hidden = True
    hidden.write(1, 2, 100.25)
    raw = io.BytesIO()
    book.save(raw)
    source = store(storage, company, records, [incoming(company, "原表.xls", raw.getvalue())])[0]
    result = query(
        records, storage, company, action="inspect", source_id=source["source_id"], limit=20
    )["inspection"]
    assert result["page"]["total_count"] == 2
    cell = next(item for item in result["blocks"] if item["location"].startswith("sheet:隐藏!C2"))
    assert "[hidden]" in cell["location"]
    assert json.loads(cell["text"])["value"] == "100.25"


def test_xls_formula_limitations_stay_visible_on_every_page_and_actual_sheet(records_case):
    storage, (company, _), records = records_case
    book = xlwt.Workbook()
    normal = book.add_sheet("正常")
    normal.write(0, 0, "真实原文")
    normal.write(0, 1, xlwt.Formula("100+200"))
    hidden = book.add_sheet("隐藏")
    hidden.visibility = 1
    hidden.write(1, 2, xlwt.Formula("SUM(1,2)"))
    raw = io.BytesIO()
    book.save(raw)
    source = store(storage, company, records, [incoming(company, "公式原表.xls", raw.getvalue())])[
        0
    ]
    cursor, blocks, warnings = None, [], None
    while True:
        scope = {"cursor": cursor} if cursor is not None else {}
        result = query(
            records,
            storage,
            company,
            action="inspect",
            source_id=source["source_id"],
            limit=1,
            **scope,
        )["inspection"]
        assert result["unsupported_media"]
        if warnings is None:
            warnings = result["unsupported_sections"]
        assert result["unsupported_sections"] == warnings
        assert len(warnings) == 2
        assert all(
            "公式表达式不能提取，仅提供缓存值；空缓存不代表业务事实为空" in value
            for value in warnings
        )
        assert {value.split("：", 1)[0] for value in warnings} == {"sheet:正常", "sheet:隐藏"}
        blocks.extend(result["blocks"])
        cursor = result["page"]["next_cursor"]
        if cursor is None:
            break
    assert len(blocks) == result["page"]["total_count"] == 3
    formula = next(block for block in blocks if block["location"] == "sheet:正常!B1")
    assert json.loads(formula["text"]) == {"type": 1, "value": ""}
    assert any(block["location"] == "sheet:隐藏!C2 [hidden]" for block in blocks)
    sheet = query(
        records,
        storage,
        company,
        action="inspect",
        source_id=source["source_id"],
        section="sheet:正常",
        limit=20,
    )["inspection"]
    assert sheet["unsupported_media"] and sheet["unsupported_sections"] == warnings
    assert sheet["page"]["filtered_count"] == 2


def pdf_bytes():
    writer = PdfWriter()
    page = writer.add_blank_page(width=200, height=200)
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    page[NameObject("/Resources")] = DictionaryObject(
        {NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})}
    )
    content = DecodedStreamObject()
    content.set_data(b"BT /F1 12 Tf 10 100 Td (Original statement 100.00) Tj ET")
    page[NameObject("/Contents")] = writer._add_object(content)
    image_page = writer.add_blank_page(width=200, height=200)
    image = DecodedStreamObject()
    image.set_data(b"\xff\xff\xff")
    image.update(
        {
            NameObject("/Type"): NameObject("/XObject"),
            NameObject("/Subtype"): NameObject("/Image"),
            NameObject("/Width"): NumberObject(1),
            NameObject("/Height"): NumberObject(1),
            NameObject("/ColorSpace"): NameObject("/DeviceRGB"),
            NameObject("/BitsPerComponent"): NumberObject(8),
        }
    )
    image_page[NameObject("/Resources")] = DictionaryObject(
        {NameObject("/XObject"): DictionaryObject({NameObject("/Im0"): writer._add_object(image)})}
    )
    content = DecodedStreamObject()
    content.set_data(b"q 100 0 0 100 0 0 cm /Im0 Do Q")
    image_page[NameObject("/Contents")] = writer._add_object(content)
    raw = io.BytesIO()
    writer.write(raw)
    return raw.getvalue()


def test_mixed_pdf_preserves_real_text_and_identifies_nonextractable_page(records_case):
    storage, (company, _), records = records_case
    source = store(storage, company, records, [incoming(company, "混合PDF.pdf", pdf_bytes())])[0]
    result = query(
        records, storage, company, action="inspect", source_id=source["source_id"], limit=20
    )["inspection"]
    assert result["unsupported_media"]
    assert result["unsupported_sections"] == ["page:2"]
    assert len(result["blocks"]) == 1
    assert result["blocks"][0]["location"] == "page:1"
    assert "Original statement 100.00" in result["blocks"][0]["text"]
    scanned = query(
        records,
        storage,
        company,
        action="inspect",
        source_id=source["source_id"],
        section="page:2",
        limit=20,
    )["inspection"]
    assert scanned["unsupported_media"] and scanned["blocks"] == []


def test_inspect_cursor_cannot_be_reused_for_different_section_or_content(records_case):
    storage, (company, _), records = records_case
    source = store(
        storage, company, records, [incoming(company, "原文.txt", b"one\ntwo\nthree\n")]
    )[0]
    first = query(
        records, storage, company, action="inspect", source_id=source["source_id"], limit=1
    )["inspection"]
    with pytest.raises(ContractError) as failed:
        query(
            records,
            storage,
            company,
            action="inspect",
            source_id=source["source_id"],
            section="text",
            limit=1,
            cursor=first["page"]["next_cursor"],
        )
    assert failed.value.code == "snapshot_changed"


@pytest.mark.parametrize("encoding", ["utf-8-sig", "utf-16", "utf-32", "gb18030"])
def test_text_encodings_preserve_complete_chinese_originals(records_case, encoding):
    storage, (company, _), records = records_case
    text = "中文原始说明\r\n实际金额100.00\r末行"
    source = store(
        storage, company, records, [incoming(company, "编码原文.txt", text.encode(encoding))]
    )[0]
    result = query(
        records, storage, company, action="inspect", source_id=source["source_id"], limit=20
    )["inspection"]
    assert "".join(block["text"] for block in result["blocks"]) == text.replace(
        "\r\n", "\n"
    ).replace("\r", "\n")


def test_csv_positions_include_columns_beyond_excel_width(records_case):
    storage, (company, _), records = records_case
    raw = "第一列," + "," * 18277 + "最后列\n"
    source = store(storage, company, records, [incoming(company, "宽表.csv", raw.encode())])[0]
    result = query(
        records, storage, company, action="inspect", source_id=source["source_id"], limit=20
    )["inspection"]
    assert [(block["location"], block["text"]) for block in result["blocks"]] == [
        ("CSV!A1", "第一列"),
        ("CSV!row:1:column:18279", "最后列"),
    ]


@pytest.mark.parametrize(
    ("name", "raw"),
    [
        ("未闭合.csv", b'header,amount\n"unclosed,100.00\n'),
        ("无效.xls", bytes.fromhex("d0cf11e0a1b11ae1") + b"not an OLE workbook"),
    ],
)
def test_unreadable_known_formats_fail_without_returning_a_partial_complete_view(
    records_case, name, raw
):
    storage, (company, _), records = records_case
    source = store(storage, company, records, [incoming(company, name, raw)])[0]
    with pytest.raises(ContractError) as failed:
        query(records, storage, company, action="inspect", source_id=source["source_id"], limit=20)
    assert failed.value.code == "source_unreadable"


def test_xlsx_1904_dates_durations_and_original_amount_lexemes(records_case):
    storage, (company, _), records = records_case
    book = Workbook()
    book.epoch = CALENDAR_MAC_1904
    sheet = book.active
    sheet["A1"] = datetime(2026, 1, 12, 13, 14)
    sheet["B1"] = timedelta(days=1, hours=1)
    sheet["C1"] = 100.25
    raw = io.BytesIO()
    book.save(raw)
    source = store(storage, company, records, [incoming(company, "日期.xlsx", raw.getvalue())])[0]
    result = query(
        records, storage, company, action="inspect", source_id=source["source_id"], limit=20
    )["inspection"]
    values = {block["location"]: json.loads(block["text"]) for block in result["blocks"]}
    assert values["sheet:Sheet!A1"]["display_value"] == "2026-01-12T13:14:00"
    assert values["sheet:Sheet!B1"]["display_value"] == "1 day, 1:00:00"
    assert values["sheet:Sheet!C1"]["value"] == "100.25"
    assert "display_value" not in values["sheet:Sheet!C1"]


@pytest.mark.parametrize("damage", ["negative_style", "negative_string", "wrong_row"])
def test_xlsx_invalid_positions_and_indices_do_not_select_other_cells(records_case, damage):
    storage, (company, _), records = records_case
    book = Workbook()
    book.active["A1"] = "真实原文"
    raw, output = io.BytesIO(), io.BytesIO()
    book.save(raw)
    with ZipFile(raw) as archive, ZipFile(output, "w", ZIP_DEFLATED) as rewritten:
        for name in archive.namelist():
            data = archive.read(name)
            if name == "xl/worksheets/sheet1.xml":
                root = ElementTree.fromstring(data)
                cell = root.find("{*}sheetData/{*}row/{*}c")
                if damage == "negative_style":
                    cell.set("s", "-1")
                elif damage == "wrong_row":
                    cell.set("r", "A2")
                else:
                    cell.set("t", "s")
                    for child in list(cell):
                        cell.remove(child)
                    ElementTree.SubElement(
                        cell, "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}v"
                    ).text = "-1"
                data = ElementTree.tostring(root)
            rewritten.writestr(name, data)
        if damage == "negative_string":
            rewritten.writestr(
                "xl/sharedStrings.xml",
                '<sst xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
                "<si><t>不能因负数索引读到这条</t></si></sst>",
            )
    imported = [incoming(company, "异常.xlsx", output.getvalue())]
    source = store(storage, company, records, imported)[0]
    with pytest.raises(ContractError) as failed:
        query(records, storage, company, action="inspect", source_id=source["source_id"], limit=20)
    assert failed.value.code == "source_unreadable"
