"""Synthetic file producers for task tests, not product backup/report owners."""

from __future__ import annotations

import hashlib
import json
import shutil
from collections import Counter
from contextlib import contextmanager
from pathlib import Path
from uuid import uuid4
from zipfile import ZipFile

from openpyxl import Workbook, load_workbook

from ai_accounting_modular.contracts.public import canonical_json, content_digest
from ai_accounting_modular.files.public import (
    FileArtifact,
    FileScope,
    ProductionError,
    create_files,
    repository_factory,
)
from ai_accounting_modular.identity.public import IdentityContext, create_identity
from ai_accounting_modular.identity.public import repository_factory as identity_factory
from ai_accounting_modular.storage.public import (
    CompanySpec,
    OwnerProfile,
    Storage,
    create_catalog,
)

ROOT = Path(__file__).resolve().parents[3]


def make_storage(root):
    catalog = root / "catalog.sqlite3"
    create_catalog(catalog, OwnerProfile(str(uuid4()), "文件任务合成老板", "test-hash"))
    storage = reopen_storage(catalog)
    companies = tuple(
        storage.create_company(
            CompanySpec(
                str(uuid4()),
                str(uuid4()),
                "91111111111111111" + str(index),
                f"文件任务合成公司{index}",
                "2026-01",
                root / f"company-{index}",
            )
        )
        for index in (1, 2)
    )
    return storage, companies


def reopen_storage(catalog):
    return Storage(catalog, repositories=(repository_factory(), identity_factory()))


def save_business_object(storage, company):
    identity = create_identity()
    with storage.write_unit(company.company_id) as uow:
        plan = identity.plan_change(
            [
                {
                    "object_id": str(uuid4()),
                    "kind": "person",
                    "name": "实际并发提交",
                    "expected_seq": 0,
                    "data": {},
                }
            ],
            IdentityContext(uow, str(uuid4())),
        )
        identity.apply_current(plan, uow)


def identity(company):
    return {
        "credit_code": company.credit_code,
        "legal_name": company.legal_name,
        "accounting_standard": "small_enterprise",
        "currency": "CNY",
        "start_period": company.start_period,
    }


def backup_plan(storage, company, *, capture_id=None, label="first"):
    capture_id = capture_id or str(uuid4())
    capture = company.company_root / "runtime/tasks" / capture_id
    database = capture / "runtime/company.sqlite3"
    database.parent.mkdir(parents=True)
    with storage.read_unit(company.company_id) as read:
        snapshot = storage.capture_database(read, database)
    data = database.read_bytes()
    digest = hashlib.sha256(data).hexdigest()
    manifest = {
        "format": "ai-accounting-company-backup/1",
        "company_identity": identity(company),
        "company_id": company.company_id,
        "database_id": company.database_id,
        "database_head": snapshot.head,
        "structure_fingerprint": snapshot.structure_fingerprint,
        "document_set_id": content_digest(label),
        "files": [
            {
                "path": "runtime/company.sqlite3",
                "kind": "database",
                "sha256": digest,
                "byte_size": len(data),
            }
        ],
    }
    return {
        "kind": "backup",
        "company_id": company.company_id,
        "database_id": company.database_id,
        "scope_head": content_digest({"label": label, "database_id": company.database_id}),
        "capture_id": capture_id,
        "manifest": manifest,
        "inputs": [
            {
                "source_root": "capture",
                "relative_path": "runtime/company.sqlite3",
                "archive_path": "runtime/company.sqlite3",
                "sha256": digest,
                "byte_size": len(data),
            }
        ],
    }


def report_plan(company, *, label="first", period="2026-01"):
    examples = json.loads(
        (ROOT / "docs/design/examples/file-cases.json").read_text(encoding="utf-8")
    )
    plan = examples["contract_examples"]["report_export_plan"]
    plan.update(company_id=company.company_id, database_id=company.database_id)
    plan["scope_head"] = content_digest({"label": label, "database": company.database_id})
    plan["organization"] = identity(company)
    plan["snapshot"].update(company_id=company.company_id, database_id=company.database_id)
    plan["snapshot"]["scope"]["company_id"] = company.company_id
    plan["snapshot"]["scope"]["period_from"] = period
    plan["snapshot"]["scope"]["period_to"] = period
    plan["snapshot"]["scope"]["as_of"] = period
    return plan


def payment_plan(company, *, label="first", period="2026-01"):
    business_id, party_id = str(uuid4()), str(uuid4())
    refs = [{"business_id": business_id, "slot": "primary"}]
    return {
        "kind": "payment_export",
        "plan_version": 1,
        "company_id": company.company_id,
        "database_id": company.database_id,
        "scope_head": content_digest({"label": label, "database": company.database_id}),
        "scope": {"period": period, "selection": "period"},
        "balance_period": period,
        "candidates": [
            {
                "business_id": business_id,
                "kind": "operations.expense",
                "party_id": party_id,
                "purpose": "合成报销",
                "remaining_fen": "30000",
                "payable_fen": "30000",
                "obligation_refs": refs,
                "basis_digest": content_digest("synthetic payable"),
            }
        ],
        "rows": [
            {
                "party_id": party_id,
                "name": "合成员工",
                "account": "000000000000000001",
                "purpose": "合成报销",
                "amount_fen": "30000",
                "obligation_refs": refs,
            }
        ],
        "total_fen": "30000",
        "template": {
            "source_id": str(uuid4()),
            "sha256": content_digest("synthetic template"),
            "layout": {
                "sheet_name": "代发",
                "first_data_row": 2,
                "columns": {"payee_name": 1, "payee_account": 2, "amount": 3, "purpose": 4},
            },
        },
    }


class SyntheticProducer:
    """Produces actual ZIP/XLSX bytes; only the task boundary is under test."""

    def __init__(self):
        self.events = []
        self.pins = Counter()
        self.current_heads = None
        self.failures = 0
        self.cleanup_failures = 0
        self.produce_hook = None
        self.publish_hook = None
        self.verify_hook = None
        self.cleanup_hook = None
        self.discard_hook = None

    def describe(self, plan):
        if plan["kind"] == "backup":
            return FileScope("company", plan["scope_head"], "synthetic.zip")
        if plan["kind"] == "payment_export":
            return FileScope(
                "payment:" + plan["scope"]["period"], plan["scope_head"], "合成代发.xlsx"
            )
        scope = plan["snapshot"]["scope"]
        key = "report:" + scope["period_from"] + ":" + plan["variant"]
        return FileScope(key, plan["scope_head"], "合成报表.xlsx")

    @contextmanager
    def recover_lease(self, plan, context):
        capture_id = plan.get("capture_id")
        key = (context.company.company_id, capture_id)
        self.events.append(("lease", context.task_id, capture_id))
        self.pins[key] += 1
        try:
            yield
        finally:
            self.pins[key] -= 1
            self.events.append(("release", context.task_id, capture_id))

    def is_current(self, plan, context):
        return self.current_heads is None or plan["scope_head"] in self.current_heads

    def produce(self, plan, context):
        self.events.append(("produce", context.task_id, plan["scope_head"]))
        if self.produce_hook:
            self.produce_hook(plan, context)
        if self.failures:
            self.failures -= 1
            raise ProductionError("synthetic_failure", "合成的暂时生产失败。")
        artifact = FileArtifact(
            "work/result.zip" if plan["kind"] == "backup" else "work/result.xlsx"
        )
        path = context.task_root / artifact.relative_path
        if plan["kind"] == "backup":
            with ZipFile(path, "w") as bundle:
                bundle.writestr(
                    "test-identity.json",
                    canonical_json(
                        {
                            "company_id": plan["company_id"],
                            "head": plan["scope_head"],
                        }
                    ),
                )
                for item in plan["inputs"]:
                    root = (
                        context.capture_root
                        if item["source_root"] == "capture"
                        else context.company.company_root
                    )
                    bundle.write(root / item["relative_path"], item["archive_path"])
        else:
            workbook = Workbook()
            workbook.active.append([plan["company_id"], plan["scope_head"], "synthetic task test"])
            workbook.save(path)
            workbook.close()
        return artifact

    def verify(self, plan, stream, context):
        self.events.append(("verify", context.task_id, plan["scope_head"]))
        if self.verify_hook:
            self.verify_hook(plan, stream, context)
        stream.seek(0)
        if plan["kind"] == "backup":
            with ZipFile(stream) as bundle:
                header = json.loads(bundle.read("test-identity.json"))
                if header != {"company_id": plan["company_id"], "head": plan["scope_head"]}:
                    raise ProductionError("wrong_file", "合成文件身份不匹配。", retryable=False)
                for item in plan["inputs"]:
                    data = bundle.read(item["archive_path"])
                    if (
                        len(data) != item["byte_size"]
                        or hashlib.sha256(data).hexdigest() != item["sha256"]
                    ):
                        raise ProductionError("wrong_file", "合成包内容不匹配。", retryable=False)
        else:
            workbook = load_workbook(stream, read_only=True)
            try:
                if (workbook.active["A1"].value, workbook.active["B1"].value) != (
                    plan["company_id"],
                    plan["scope_head"],
                ):
                    raise ProductionError("wrong_file", "合成报表身份不匹配。", retryable=False)
            finally:
                workbook.close()

    def publish(self, plan, artifact, context):
        self.events.append(("publish", context.task_id, plan["scope_head"]))
        if self.publish_hook:
            return self.publish_hook(plan, artifact, context)
        return context.task_root / artifact.relative_path

    def cleanup(self, context, *, keep_capture):
        self.events.append(("cleanup", context.task_id, keep_capture))
        if self.cleanup_hook:
            self.cleanup_hook(context, keep_capture)
        if self.cleanup_failures:
            self.cleanup_failures -= 1
            raise OSError("synthetic cleanup lock")
        if context.capture_root is not None and not keep_capture and context.capture_root.exists():
            assert not self.pins[(context.company.company_id, context.capture_root.name)]
            shutil.rmtree(context.capture_root)

    def discard_capture(self, context):
        self.events.append(("discard", context.capture_id))
        if self.discard_hook:
            self.discard_hook(context)
        if self.cleanup_failures:
            self.cleanup_failures -= 1
            raise OSError("synthetic capture cleanup lock")
        assert not self.pins[(context.company.company_id, context.capture_id)]
        if context.capture_root.exists():
            shutil.rmtree(context.capture_root)


def service(storage, producer):
    return create_files(
        storage,
        {
            "backup": producer,
            "report_export": producer,
            "payment_export": producer,
        },
    )
