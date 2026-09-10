import json
from collections import defaultdict
from hashlib import sha256
from io import BytesIO

import pytest
import xlwt
from sqlalchemy import func, select
from test_mybank_export import export_case as export_case
from test_mybank_export import generation, reimbursement, workbook

from ai_accounting.config import Settings
from ai_accounting.models import Evidence, MybankPaymentSourceVersion, PayrollLine, Voucher
from ai_accounting.mybank_profiles import company_export_profile
from ai_accounting.mybank_schemas import ImportMybankPaymentSourceRequest
from ai_accounting.mybank_sources import MybankPaymentSourceService, read_payment_source


def tax_file(name="张三", income="40000.00", tax="900.00"):
    return (
        "姓名,税款所属期起,税款所属期止,本期收入,已缴税额,应补(退)税额\n"
        f"{name},2026-03-01,2026-03-31,{income},12345.67,{tax}\n"
    ).encode()


@pytest.fixture
def source_case(session, organization, tmp_path, monkeypatch):
    settings = Settings(finance_storage_dir=tmp_path, finance_evidence_dir=tmp_path / "evidence")
    settings.finance_evidence_dir.mkdir()
    monkeypatch.setattr("ai_accounting.material_reader.get_settings", lambda: settings)
    monkeypatch.setattr("ai_accounting.company_notes.get_settings", lambda: settings)

    def record(raw, kind="actual_tax", revision=0, key="first"):
        digest = sha256(raw).hexdigest()
        path = settings.finance_evidence_dir / digest
        path.write_bytes(raw)
        evidence = session.scalar(
            select(Evidence).where(Evidence.org_id == organization.id, Evidence.sha256 == digest)
        )
        if evidence is None:
            evidence = Evidence(
                org_id=organization.id,
                original_name="原表.csv",
                source="test",
                storage_path=str(path),
                sha256=digest,
                size_bytes=len(raw),
            )
            session.add(evidence)
            session.flush()
        request = ImportMybankPaymentSourceRequest(
            org_id=organization.id,
            payroll_period="2026-03",
            source_kind=kind,
            evidence_id=evidence.id,
            expected_revision=revision,
            idempotency_key=key,
        )
        return MybankPaymentSourceService(session).import_source(request), request, path

    return record


def actual_profile(organization):
    from ai_accounting.company_notes import EMPTY_HASH, update_company_notes

    profile = company_export_profile(organization)
    profile["salary"]["source"] = "reported_salary_minus_actual_tax_and_employee_contributions"
    update_company_notes(
        organization, EMPTY_HASH, "```mybank-export\n" + json.dumps(profile) + "\n```"
    )


def test_tax_bureau_xls_uses_current_tax_not_cumulative_paid():
    book = xlwt.Workbook()
    sheet = book.add_sheet("税额")
    for r, row in enumerate(
        [
            ["姓名", "税款所属期起", "税款所属期止", "本期收入", "已缴税额", "应补(退)税额"],
            ["张三", "2026-03-01", "2026-03-31", 40000, 12345.67, 900],
        ]
    ):
        for c, value in enumerate(row):
            sheet.write(r, c, value)
    out = BytesIO()
    book.save(out)
    rows = read_payment_source(out.getvalue(), ".xls", "actual_tax", "2026-03")
    assert rows[0]["actual_individual_income_tax_fen"] == 90000


@pytest.mark.parametrize(
    "raw,code",
    [
        (tax_file(tax=""), "AMOUNT_MISSING"),
        (tax_file(tax="0.001"), "AMOUNT_PRECISION"),
        (tax_file().replace(b"2026-03-31", b"2026-04-30"), "PERIOD_CONFLICT"),
        (tax_file() + tax_file().splitlines(True)[1], "DUPLICATE_PERSON"),
    ],
)
def test_invalid_original_is_atomic(session, source_case, raw, code):
    with pytest.raises(ValueError, match=code):
        source_case(raw)
    assert session.scalar(select(func.count()).select_from(MybankPaymentSourceVersion)) == 0


def test_payment_register_dash_is_not_assumed_to_be_zero(session, source_case):
    raw = (
        "姓名,月份,报税工资,未报税劳务,回扣报销1,回扣报销2,发票报销\n张三,2026-03,100,-,0,0,0\n"
    ).encode()
    with pytest.raises(ValueError, match="AMOUNT_UNREADABLE"):
        source_case(raw, "payment_register")
    assert session.scalar(select(func.count()).select_from(MybankPaymentSourceVersion)) == 0


def test_import_version_idempotency_and_immutable_history(session, source_case):
    result, request, _ = source_case(tax_file())
    service = MybankPaymentSourceService(session)
    assert service.import_source(request)["source_id"] == result["source_id"]
    updated, _, _ = source_case(tax_file(tax="901.00"), revision=1, key="second")
    assert updated["revision"] == 2
    with pytest.raises(ValueError, match="REVISION_CHANGED"):
        source_case(tax_file(tax="902.00"), revision=1, key="stale")
    row = session.scalar(
        select(MybankPaymentSourceVersion).order_by(MybankPaymentSourceVersion.revision)
    )
    row.content = {"rows": []}
    with pytest.raises(ValueError, match="IMMUTABLE"):
        session.flush()
    session.rollback()


def test_actual_tax_salary_still_deducts_social_and_fund(
    session, organization, export_case, source_case, tmp_path
):
    service, request, employee, batch = export_case
    actual_profile(organization)
    line = session.scalar(select(PayrollLine).where(PayrollLine.payroll_batch_id == batch.batch_id))
    income = f"{line.tax_reported_salary_fen // 100}.{line.tax_reported_salary_fen % 100:02d}"
    source_case(tax_file(employee.name, income, "900.00"))
    result = service.preview(request)
    assert result["status"] == "ready", result
    assert result["totals"]["salary_fen"] == (
        line.tax_reported_salary_fen
        - 90000
        - line.employee_social_insurance_fen
        - line.employee_housing_fund_fen
    )
    before = session.scalar(select(func.count()).select_from(Voucher))
    generated = service.generate(generation(request, result, tmp_path))
    assert len(generated["files"]) == 1
    assert session.scalar(select(func.count()).select_from(Voucher)) == before


def test_missing_actual_tax_does_not_fall_back_to_calculation(
    organization, export_case, source_case
):
    service, request, _, _ = export_case
    actual_profile(organization)
    result = service.preview(request)
    assert any(x["code"] == "ACTUAL_INCOME_TAX_MISSING" for x in result["missing_information"])


@pytest.mark.parametrize("export_case", ["with_material_period"], indirect=True)
def test_complete_export_discovers_reimbursement_and_reports_unposted_labor(
    session, organization, export_case, source_case
):
    from ai_accounting.models import BusinessMetadataVersion

    service, request, employee, batch = export_case
    line = session.scalar(select(PayrollLine).where(PayrollLine.payroll_batch_id == batch.batch_id))
    income = f"{line.tax_reported_salary_fen // 100}.{line.tax_reported_salary_fen % 100:02d}"
    raw = (
        "姓名,月份,报税工资,未报税劳务,回扣报销1,回扣报销2,发票报销\n"
        f"{employee.name},2026-03,{income},0,1.01,2.02,3.03\n"
        "劳务甲,2026-03,0,10443.72,0,0,0\n"
    ).encode()
    source_case(raw, "payment_register")
    for purpose, amount in [("回扣报销1", 101), ("回扣报销2", 202), ("发票报销", 303)]:
        item = reimbursement(
            session,
            organization.id,
            employee,
            purpose=purpose,
            amount=amount,
            invoice=purpose == "发票报销",
        )
        # Append metadata instead of changing existing management history.
        session.add(
            BusinessMetadataVersion(
                org_id=organization.id,
                event_id=item.source_event_id,
                component_key="source",
                version=2,
                idempotency_key=f"period:{purpose}",
                request_hash="b" * 64,
                metadata_values={
                    "purpose": purpose,
                    "payment_period": "2026-03",
                    "beneficiary": {"kind": "employee", "name": employee.name},
                },
            )
        )
    session.flush()
    result = service.preview(request.model_copy(update={"scope": "complete"}))
    assert result["totals"]["reimbursement_fen"] == 606
    assert result["totals"]["labor_fen"] == 1044372
    codes = {x["code"] for x in result["missing_information"]}
    assert "UNTAXED_LABOR_PAYABLE_MISSING" in codes
    assert "MYBANK_MATERIAL_INCOMPLETE" in codes
    assert "REGISTER_REIMBURSEMENT_SOURCE_CONFLICT" not in codes


def test_single_sheet_contains_three_distinct_purposes():
    from openpyxl import load_workbook

    from ai_accounting.mybank_export import render_import

    template = workbook([["收款方名称", "收款方账号", "金额", "附言/用途"]], template=True)
    rows = [
        {"name": "张三", "account": "001234567890", "amount_fen": 101, "category": label}
        for label in ("薪资", "劳务", "报销")
    ]
    result = load_workbook(BytesIO(render_import(template, rows, "2026-08")))
    assert len(result.sheetnames) == 1
    assert [result.active.cell(r, 4).value for r in range(2, 5)] == [
        "2026-08 薪资",
        "2026-08 劳务",
        "2026-08 报销",
    ]


def test_source_change_is_detected(session, organization, source_case):
    _, request, path = source_case(tax_file())
    path.write_bytes(b"changed")
    with pytest.raises(ValueError, match="MATERIAL_SOURCE_CHANGED"):
        MybankPaymentSourceService(session).latest(organization.id, "2026-03", "actual_tax")
    with pytest.raises(ValueError, match="MATERIAL_SOURCE_CHANGED"):
        MybankPaymentSourceService(session).import_source(request)


def test_currency_accounting_zero_does_not_fill_unknown_cells():
    header = "姓名,月份,报税工资,未报税劳务,回扣报销1,回扣报销2,发票报销\n"
    row = "张三,2026-03,40000,¥ -,￥\u00a0-,0,0\n"
    result = read_payment_source((header + row).encode(), ".csv", "payment_register", "2026-03")
    assert result[0]["untaxed_labor_fen"] == result[0]["rebate_1_fen"] == 0
    for unknown in ("", "-", "待定"):
        with pytest.raises(ValueError):
            read_payment_source(
                (header + row.replace("¥ -", unknown)).encode(),
                ".csv",
                "payment_register",
                "2026-03",
            )


def test_replay_exports_latest_source_as_original_evidence_reference(
    session, organization, source_case, tmp_path
):
    from ai_accounting import replay_cli
    from ai_accounting.material_replay import export_material_operations

    source_case(tax_file())
    latest, request, _ = source_case(tax_file(tax="901.00"), revision=1, key="second")
    maps = defaultdict(dict)
    maps["evidence"][str(request.evidence_id)] = {
        "$ref": "evidence",
        "sha256": latest["evidence_sha256"],
    }
    operations = export_material_operations(session, organization.id, maps, tmp_path)
    assert len(operations) == 1
    operation = operations[0]
    assert operation["kind"] == "tool"
    assert operation["tool"] == "finance_import_mybank_payment_source"
    assert operation["request"]["org_id"] == "${ORG_ID}"
    assert operation["request"]["evidence_id"] == maps["evidence"][str(request.evidence_id)]
    assert operation["request"]["expected_revision"] == 0
    assert "rows" not in operation["request"]
    replay_cli._verify_operation_references(
        operations, evidence_hashes={latest["evidence_sha256"]}, org_id=str(organization.id)
    )


@pytest.mark.parametrize("register_amount", [None, "0", "1.01"])
@pytest.mark.parametrize("export_case", ["with_material_period"], indirect=True)
def test_complete_export_keeps_known_labor_without_an_extra_register(
    session, organization, export_case, source_case, register_amount
):
    from ai_accounting.models import BusinessMetadataVersion

    service, request, employee, batch = export_case
    item = reimbursement(session, organization.id, employee, purpose="劳务", amount=101)
    session.add(
        BusinessMetadataVersion(
            org_id=organization.id,
            event_id=item.source_event_id,
            component_key="source",
            version=2,
            idempotency_key="labor-payment-scope",
            request_hash="d" * 64,
            metadata_values={
                "payment_period": "2026-03",
                "payment_category": "labor",
                "beneficiary": {"kind": "employee", "name": employee.name},
            },
        )
    )
    session.flush()
    if register_amount is not None:
        line = session.scalar(
            select(PayrollLine).where(PayrollLine.payroll_batch_id == batch.batch_id)
        )
        income = f"{line.tax_reported_salary_fen // 100}.{line.tax_reported_salary_fen % 100:02d}"
        source_case(
            (
                "姓名,月份,报税工资,未报税劳务,回扣报销1,回扣报销2,发票报销\n"
                f"{employee.name},2026-03,{income},{register_amount},0,0,0\n"
            ).encode(),
            "payment_register",
        )
    result = service.preview(request.model_copy(update={"scope": "complete"}))
    assert result["totals"]["labor_fen"] == 101
    assert result["totals"]["reimbursement_fen"] == 0
    codes = {issue["code"] for issue in result["missing_information"]}
    assert "UNTAXED_LABOR_PAYABLE_MISSING" not in codes
    assert "LABOR_BENEFICIARY_MISSING" not in codes
    assert ("UNTAXED_LABOR_AMOUNT_CONFLICT" in codes) == (register_amount == "0")


@pytest.mark.parametrize(
    "raw_amount,expected", [("90071992547409.93", 9007199254740993), ("3076.8700000000001", None)]
)
def test_tax_xlsx_retains_original_numeric_precision(raw_amount, expected):
    from xml.etree import ElementTree
    from zipfile import ZipFile

    raw = workbook(
        [
            ["姓名", "税款所属期起", "税款所属期止", "本期收入", "应补(退)税额"],
            ["张三", "2026-03-01", "2026-03-31", 100, 1],
        ]
    )
    output = BytesIO()
    with ZipFile(BytesIO(raw)) as source, ZipFile(output, "w") as target:
        for entry in source.infolist():
            data = source.read(entry.filename)
            if entry.filename == "xl/worksheets/sheet1.xml":
                document = ElementTree.fromstring(data)
                document.find(".//{*}c[@r='E2']/{*}v").text = raw_amount
                data = ElementTree.tostring(document)
            target.writestr(entry, data)
    if expected is None:
        with pytest.raises(ValueError, match="AMOUNT_PRECISION"):
            read_payment_source(output.getvalue(), ".xlsx", "actual_tax", "2026-03")
    else:
        rows = read_payment_source(output.getvalue(), ".xlsx", "actual_tax", "2026-03")
        assert rows[0]["actual_individual_income_tax_fen"] == expected
