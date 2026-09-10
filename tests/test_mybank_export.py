from __future__ import annotations

import json
import uuid
from datetime import date
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

import pytest
from openpyxl import Workbook, load_workbook
from pydantic import ValidationError
from sqlalchemy import func, select
from test_payroll_service import preview_and_confirm

from ai_accounting import mcp_server
from ai_accounting.coa import seed_organization
from ai_accounting.models import (
    BusinessEvent,
    BusinessEventComponent,
    BusinessMetadataVersion,
    Employee,
    OpenItem,
    PayrollLine,
    Voucher,
)
from ai_accounting.mybank_export import MybankExportError, main, read_recipients, render_import
from ai_accounting.mybank_profiles import company_export_profile
from ai_accounting.mybank_schemas import GenerateMybankExportRequest, PreviewMybankExportRequest
from ai_accounting.mybank_service import MybankExportService


def workbook(rows, *, template=False):
    book = Workbook()
    book.template = template
    for row in rows:
        book.active.append(row)
    stream = BytesIO()
    book.save(stream)
    return stream.getvalue()


@pytest.fixture
def export_case(session, organization, tmp_path, request):
    if getattr(request, "param", None) == "with_material_period":
        from ai_accounting.accounting_period_schemas import GenerateAccountingPeriodRequest
        from ai_accounting.accounting_period_service import AccountingPeriodService

        generated = AccountingPeriodService(session).generate_accounting_period(
            GenerateAccountingPeriodRequest(
                org_id=organization.id, period_month="2026-03", idempotency_key="march-materials"
            )
        )
        assert generated.status == "posted", generated
    _, batch = preview_and_confirm(session, organization)
    employee = session.scalar(select(Employee).where(Employee.org_id == organization.id))
    template = tmp_path / "bank.xltx"
    template.write_bytes(
        workbook(
            [
                [
                    "收款方名称（必输，文本格式）",
                    "收款方账号（必输，本文格式）",
                    "金额（必输，文本格式）",
                    "附言/用途（必输，文本格式，最多40字）",
                    None,
                    "填写说明",
                ],
                ["旧名字", "001111111111111111", "9999999.99", "旧用途"],
            ],
            template=True,
        )
    )
    recipients = tmp_path / "recipients.xlsx"
    # Deliberately poison ALL external monetary fields; none may reach an export.
    recipients.write_bytes(
        workbook(
            [
                ["编号(*)", "账户名称(*)", "账号(*)", "交易金额(*)", "实发工资", "发票报销"],
                [1, employee.name, "001234567890123456789", "99999999.99", "=1/0", 999999999],
            ]
        )
    )
    request = PreviewMybankExportRequest(
        org_id=organization.id,
        payroll_period="2026-03",
        scope="selected",
        template_path=str(template),
        recipients_path=str(recipients),
    )
    return MybankExportService(session), request, employee, batch


def reimbursement(
    session,
    org_id,
    employee,
    *,
    purpose="回扣报销1",
    amount=12345,
    settled=0,
    invoice=False,
    anonymous=False,
):
    """Model-only source fixture, independent of the export implementation."""
    key = str(uuid.uuid4())
    event = BusinessEvent(
        org_id=org_id,
        idempotency_key=key,
        event_type="composite",
        status="posted",
        facts={},
        business_date=date(2026, 4, 2),
        posting_date=date(2026, 4, 2),
    )
    session.add(event)
    session.flush()
    component = BusinessEventComponent(
        org_id=org_id,
        event_id=event.id,
        key="source",
        ordinal=1,
        kind="expense" if invoice else "pass_through",
        facts={"amount_fen": amount, "payment_basis": "person_advance"}
        if invoice
        else {"amount_fen": amount},
    )
    session.add(component)
    session.flush()
    item = OpenItem(
        org_id=org_id,
        source_event_id=event.id,
        source_component_id=component.id,
        component_key="primary",
        item_type="payable",
        original_amount_fen=amount,
        settled_amount_fen=settled,
        status="settled" if settled == amount else "partial" if settled else "open",
        payable_category=None if invoice else "pass_through",
        pass_through_key=None if invoice else key,
        counterparty_id=employee.counterparty_id if invoice else None,
    )
    metadata = {"purpose": purpose} if purpose else {}
    if not anonymous:
        metadata["beneficiary"] = {"kind": "employee", "name": employee.name}
    session.add_all(
        [
            item,
            BusinessMetadataVersion(
                org_id=org_id,
                event_id=event.id,
                component_key="source",
                version=1,
                metadata_values=metadata,
                idempotency_key=key,
                request_hash="a" * 64,
            ),
        ]
    )
    session.flush()
    return item


def generation(request, preview, tmp_path):
    return GenerateMybankExportRequest(
        **request.model_dump(),
        expected_source_hash=preview["source_hash"],
        output_dir=str(tmp_path / "out"),
    )


@pytest.mark.parametrize("export_case", ["with_material_period"], indirect=True)
def test_salary_inventory_people_cannot_be_silently_excluded(
    session, export_case, tmp_path, monkeypatch
):
    from hashlib import sha256

    from ai_accounting.accounting_periods import canonical_sha256
    from ai_accounting.company_notes import read_company_notes
    from ai_accounting.config import Settings
    from ai_accounting.material_schemas import (
        RegisterPeriodMaterialsRequest,
        UpdatePeriodMaterialInventoryRequest,
    )
    from ai_accounting.material_service import MaterialService
    from ai_accounting.models import AccountingPeriod, Evidence, Organization, PayrollBatch

    service, request, employee, batch_result = export_case
    settings = Settings(finance_storage_dir=tmp_path / "storage", finance_evidence_dir=tmp_path)
    monkeypatch.setattr("ai_accounting.company_notes.get_settings", lambda: settings)
    monkeypatch.setattr("ai_accounting.material_reader.get_settings", lambda: settings)
    batch = session.get(PayrollBatch, batch_result.batch_id)
    component = session.scalar(
        select(BusinessEventComponent).where(
            BusinessEventComponent.event_id == batch.business_event_id,
            BusinessEventComponent.kind == "payroll_accrual",
        )
    )
    line = session.scalar(select(PayrollLine).where(PayrollLine.payroll_batch_id == batch.id))
    period = session.scalar(
        select(AccountingPeriod).where(
            AccountingPeriod.org_id == request.org_id,
            AccountingPeriod.calendar_year == 2026,
            AccountingPeriod.calendar_month == 3,
        )
    )
    path = tmp_path / "工资资料.csv"
    gross = line.gross_salary_fen
    path.write_text(
        f"姓名,应发\n{employee.name},{gross // 100}.{gross % 100:02d}\n", encoding="utf-8"
    )
    raw = path.read_bytes()
    evidence = Evidence(
        org_id=request.org_id,
        original_name=path.name,
        storage_path=str(path),
        sha256=sha256(raw).hexdigest(),
        size_bytes=len(raw),
        source="test",
    )
    session.add(evidence)
    session.flush()
    material_service = MaterialService(session)
    registered = material_service.register(
        RegisterPeriodMaterialsRequest(
            org_id=request.org_id,
            period_id=period.id,
            expected_revision=0,
            idempotency_key="salary-scope",
            sources=[
                {
                    "evidence_id": evidence.id,
                    "columns": [
                        {"column": "A", "role": "context"},
                        {"column": "B", "role": "amount"},
                    ],
                }
            ],
        )
    )
    item = registered["completeness"]["items"][0]
    org = session.get(Organization, request.org_id)
    material_service.update(
        UpdatePeriodMaterialInventoryRequest(
            org_id=request.org_id,
            period_id=period.id,
            expected_revision=1,
            idempotency_key="salary-scope-reviewed",
            reviewed_notes_hash=read_company_notes(org)["sha256"],
            resolutions=[
                {
                    "item_key": item["key"],
                    "treatment": "recognize",
                    "business_kind": "payroll_accrual",
                    "recognition_period": "2026-03",
                    "export_category": "salary",
                    "employee_id": employee.id,
                    "links": [
                        {
                            "source": {"component_id": component.id},
                            "amount_fen": gross,
                            "expected_facts_hash": canonical_sha256(component.facts),
                        }
                    ],
                }
            ],
        )
    )
    preview = service.preview(
        request.model_copy(update={"scope": "complete", "employee_ids": [uuid.uuid4()]})
    )
    issue = next(
        item
        for item in preview["missing_information"]
        if item["code"] == "MYBANK_MATERIAL_SALARY_EMPLOYEES_MISSING"
    )
    assert issue["employee_ids"] == [str(employee.id)]


def test_uses_kernel_net_and_all_three_reimbursements_only(session, export_case, tmp_path):
    service, request, employee, batch = export_case
    items = [
        reimbursement(
            session,
            request.org_id,
            employee,
            purpose=purpose,
            amount=amount,
            invoice=purpose == "发票报销",
        )
        for purpose, amount in [("回扣报销1", 101), ("回扣报销2", 202), ("发票报销", 303)]
    ]
    request = request.model_copy(
        update={"reimbursement_open_item_ids": [item.id for item in items]}
    )
    before = session.scalar(select(func.count()).select_from(Voucher))
    preview = service.preview(request)
    line = session.scalar(select(PayrollLine).where(PayrollLine.payroll_batch_id == batch.batch_id))
    assert preview["status"] == "ready", preview
    assert preview["totals"]["salary_fen"] == line.net_salary_fen < line.gross_salary_fen
    assert preview["totals"]["reimbursement_fen"] == 606
    assert preview["totals"]["invoice_fen"] == 303
    generated = service.generate(generation(request, preview, tmp_path))
    assert generated["status"] == "generated"
    assert session.scalar(select(func.count()).select_from(Voucher)) == before
    assert all(item.settled_amount_fen == 0 for item in items)
    assert "001234567890123456789" not in json.dumps(generated)
    path = tmp_path / "out" / "网商银行_2026-03_代发.xlsx"
    book = load_workbook(path)
    assert [book.active.cell(3, c).value for c in range(1, 5)] == [
        employee.name,
        "001234567890123456789",
        "6.06",
        "2026-03 报销",
    ]
    assert all(book.active.cell(2, c).data_type == "s" for c in range(1, 5))
    assert book.active["F1"].value == "填写说明"
    assert book.template is False
    with ZipFile(path) as archive:
        assert b"spreadsheetml.sheet.main+xml" in archive.read("[Content_Types].xml")
    with pytest.raises(MybankExportError, match="已存在"):
        service.generate(generation(request, preview, tmp_path))


def test_merged_payable_uses_remaining_total_without_inventing_split(session, export_case):
    service, request, employee, _ = export_case
    item = reimbursement(
        session, request.org_id, employee, purpose=None, amount=123456, settled=12345
    )
    request = request.model_copy(update={"reimbursement_open_item_ids": [item.id]})
    result = service.preview(request)
    assert result["status"] == "ready"
    assert result["totals"]["combined_reimbursement_fen"] == 111111
    assert result["totals"]["invoice_fen"] == 0
    assert result["totals"]["reimbursement_fen"] == 111111


def test_settled_anonymous_combined_payment_is_never_reexported(session, export_case):
    service, request, employee, _ = export_case
    item = reimbursement(
        session,
        request.org_id,
        employee,
        purpose=None,
        amount=2342650,
        settled=2342650,
        anonymous=True,
    )
    result = service.preview(request.model_copy(update={"reimbursement_open_item_ids": [item.id]}))
    assert result["status"] == "ready"
    assert result["totals"]["reimbursement_fen"] == 0
    assert any(
        row["source"]["open_item_id"] == str(item.id)
        for row in result["skipped"]
        if "source" in row
    )


@pytest.mark.parametrize("change", ["partial", "settled", "reversed", "deleted"])
def test_salary_payment_and_lifecycle_prevent_duplicate_export(
    session, export_case, tmp_path, change
):
    service, request, _, batch = export_case
    preview = service.preview(request)
    item = session.scalar(
        select(OpenItem).where(
            OpenItem.source_event_id == batch.event_id, OpenItem.payable_category == "salary"
        )
    )
    if change in {"partial", "settled"}:
        item.status = change
        item.settled_amount_fen = 1 if change == "partial" else item.original_amount_fen
    else:
        session.get(BusinessEvent, batch.event_id).status = change
    session.flush()
    result = service.generate(generation(request, preview, tmp_path))
    assert result["status"] == "needs_information"
    assert not (tmp_path / "out").exists()
    assert result["totals"]["salary_fen"] == 0


def test_foreign_company_source_fails_without_outputs(session, export_case, tmp_path):
    service, request, employee, _ = export_case
    foreign = seed_organization(
        session,
        name="其他公司",
        taxpayer_identification_number="91330106MAK342U989",
        accounting_period_control_enabled=False,
    )
    item = reimbursement(session, foreign.id, employee, anonymous=True)
    selected = request.model_copy(update={"reimbursement_open_item_ids": [item.id]})
    result = service.generate(generation(selected, service.preview(selected), tmp_path))
    assert result["status"] == "needs_information"
    assert any(i["code"] == "REIMBURSEMENT_SOURCE_NOT_FOUND" for i in result["missing_information"])
    assert not (tmp_path / "out").exists()


def test_missing_account_does_not_silently_define_employee_scope(export_case, tmp_path):
    service, request, employee, _ = export_case
    Path(request.recipients_path).write_bytes(workbook([["姓名", "账号"]]))
    preview = service.preview(request)
    assert preview["status"] == "needs_information"
    assert preview["payments"][0]["name"] == employee.name
    assert preview["totals"]["salary_fen"] > 0
    result = service.generate(generation(request, preview, tmp_path))
    assert result["status"] == "needs_information"
    assert not (tmp_path / "out").exists()


def test_expected_reimbursement_person_without_kernel_source_blocks_all_files(
    export_case, tmp_path
):
    service, request, employee, _ = export_case
    request = request.model_copy(update={"reimbursement_employee_ids": [employee.id]})
    preview = service.preview(request)
    assert preview["totals"]["salary_fen"] > 0
    assert preview["status"] == "needs_information"
    assert preview["missing_information"][0]["code"] == "EXPECTED_REIMBURSEMENT_SOURCE_MISSING"
    result = service.generate(generation(request, preview, tmp_path))
    assert result["status"] == "needs_information"
    assert not (tmp_path / "out").exists()


def test_changed_account_or_kernel_balance_invalidates_preview(session, export_case, tmp_path):
    service, request, employee, _ = export_case
    item = reimbursement(session, request.org_id, employee)
    request = request.model_copy(update={"reimbursement_open_item_ids": [item.id]})
    preview = service.preview(request)
    item.settled_amount_fen = 1
    item.status = "partial"
    session.flush()
    result = service.generate(generation(request, preview, tmp_path))
    assert result["missing_information"][0]["code"] == "SOURCE_CHANGED"
    assert not (tmp_path / "out").exists()
    preview = service.preview(request)
    Path(request.recipients_path).write_bytes(
        workbook([["姓名", "账号"], [employee.name, "012345678999999"]])
    )
    result = service.generate(generation(request, preview, tmp_path))
    assert result["missing_information"][0]["code"] == "SOURCE_CHANGED"


@pytest.mark.parametrize("account", [123456789012345678, "1.234567E+18", "=1+1"])
def test_unsafe_account_encoding_rejected(account):
    with pytest.raises(MybankExportError):
        read_recipients(workbook([["姓名", "账号"], ["张三", account]]))


def test_template_instructions_preserved_and_bank_row_limit(export_case):
    _, request, _, _ = export_case
    book = load_workbook(request.template_path)
    book.active.merge_cells("F2:P2")
    book.active["F2"] = "数字须用文本；最多2000条"
    stream = BytesIO()
    book.save(stream)
    rows = [{"name": "张三", "account": "000012345678900", "amount_fen": 1}]
    result = render_import(stream.getvalue(), rows, "2026-03", "薪资")
    output = load_workbook(BytesIO(result))
    assert str(output.active.merged_cells) == "F2:P2"
    assert output.active["F2"].value == "数字须用文本；最多2000条"
    assert output.active["C2"].value == "0.01"
    with pytest.raises(MybankExportError, match="2000"):
        render_import(stream.getvalue(), rows * 2001, "2026-03", "薪资")


def test_request_rejects_amounts_and_duplicate_sources(export_case):
    _, request, _, _ = export_case
    with pytest.raises(ValidationError):
        PreviewMybankExportRequest.model_validate({**request.model_dump(), "salary_fen": 1})
    source = uuid.uuid4()
    with pytest.raises(ValidationError):
        PreviewMybankExportRequest.model_validate(
            {**request.model_dump(), "reimbursement_open_item_ids": [source, source]}
        )


def test_cli_uses_registered_authenticated_tool(export_case, tmp_path, monkeypatch, capsys):
    _, request, _, _ = export_case
    path = tmp_path / "request.json"
    path.write_text(request.model_dump_json(), encoding="utf-8")
    registered = mcp_server.mcp._tool_manager.get_tool("finance_preview_mybank_export")
    called = []
    monkeypatch.setattr(
        registered, "fn", lambda request: called.append(request.org_id) or {"status": "ready"}
    )
    monkeypatch.setattr(
        mcp_server, "finance_preview_mybank_export", lambda **kw: pytest.fail("unsecured handler")
    )
    monkeypatch.setattr(mcp_server, "_initialize_mcp_credential_store", lambda **kw: None)
    monkeypatch.setattr(mcp_server, "_clear_mcp_credential_store", lambda: None)
    assert main(["--request", str(path)]) == 0
    assert called == [request.org_id]
    assert json.loads(capsys.readouterr().out)["status"] == "ready"


def test_owner_confirmed_defaults_are_company_scoped_and_not_mutable_by_export(
    export_case, organization, tmp_path, monkeypatch
):
    from ai_accounting.company_notes import EMPTY_HASH, update_company_notes
    from ai_accounting.config import Settings

    monkeypatch.setattr(
        "ai_accounting.company_notes.get_settings",
        lambda: Settings(finance_storage_dir=tmp_path, finance_evidence_dir=tmp_path / "evidence"),
    )
    profile = company_export_profile(organization)
    profile["salary"]["source"] = "reported_salary_minus_actual_tax_and_employee_contributions"
    profile["change_policy"] = "owner_explicit_request_only"
    content = "```mybank-export\n" + json.dumps(profile, ensure_ascii=False) + "\n```"
    update_company_notes(organization, EMPTY_HASH, content)
    saved = company_export_profile(organization)
    assert (
        saved["salary"]["source"] == "reported_salary_minus_actual_tax_and_employee_contributions"
    )
    assert saved["reimbursement"]["categories"] == ["回扣报销1", "回扣报销2", "发票报销"]
    profile["reimbursement"]["categories"].clear()
    assert len(company_export_profile(organization)["reimbursement"]["categories"]) == 3
    _, request, _, _ = export_case
    with pytest.raises(ValidationError):
        PreviewMybankExportRequest.model_validate(
            {**request.model_dump(), "company_export_profile": profile}
        )
